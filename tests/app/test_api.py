"""FR-INGEST-003/004, FR-REVIEW-001/002, FR-EXPORT-001, PRD-PROTOTYPE-001"""
import hashlib,io,json,subprocess,threading
from email.message import EmailMessage
from dataclasses import replace
from decimal import Decimal
from pathlib import Path
from copy import deepcopy
import pytest
import httpx
from openpyxl import load_workbook
from PIL import Image
from app.gateway import Gateway,ModelResult,mock_extract
from app.parsers import PARSER_VERSION
from app.runner import adjacent_extraction_batches
from .conftest import upload,run_demo


def test_health_and_no_key_exposed(client):
    assert client.get('/').status_code==200
    s=client.get('/api/settings').json()
    assert s['provider']=='mock' and not s['live_ready']
    assert 'api_key' not in s and s['thinking']=='disabled'

@pytest.mark.parametrize('name',['','   '])
def test_invalid_name(client,name):assert client.post('/api/projects',json={'name':name}).status_code in (400,422)

def test_extra_fields_rejected(client):assert client.post('/api/projects',json={'name':'x','admin':True}).status_code==422


def test_project_budget_is_user_selected_audited_and_never_below_committed(client):
    created=client.post('/api/projects',json={'name':'Custom budget','budget_cny':'425.50'})
    assert created.status_code==201 and created.json()['budget_limit_cny']=='425.500000'
    pid=created.json()['id'];db=client.app.state.db
    updated=client.put(f'/api/projects/{pid}/budget',json={'limit_cny':'500.25'})
    assert updated.status_code==200 and updated.json()['limit_cny']=='500.250000'
    event=db.one('SELECT previous_units,new_units FROM budget_limit_events WHERE project_id=?',(pid,))
    assert event=={'previous_units':425500000,'new_units':500250000}
    db.execute('UPDATE budget_accounts SET spent_units=? WHERE project_id=?',(200000000,pid))
    blocked=client.put(f'/api/projects/{pid}/budget',json={'limit_cny':'199.99'})
    assert blocked.status_code==409 and db.cost(pid)['limit_cny']=='500.250000'
    assert client.put(f'/api/projects/{pid}/budget',json={'limit_cny':'0'}).status_code==422
    assert client.put(f'/api/projects/{pid}/budget',json={'limit_cny':'0.001'}).status_code==422

def test_cross_origin_blocked(client):
    assert client.post('/api/projects',json={'name':'x'},headers={'Origin':'https://evil.invalid'}).status_code==403
    assert client.get('/api/projects',headers={'Host':'evil.invalid'}).status_code==400

def test_header_required(client):
    assert client.post('/api/projects',json={'name':'x'},headers={'X-CIRP-Client':''}).status_code==403

def test_resumable_duplicate_and_corrupt_chunk(client,project):
    u=client.post(f'/api/projects/{project["id"]}/uploads',json={'name':'a.txt','size':6}).json()
    base=f'/api/uploads/{u["id"]}'
    assert client.put(base+'/chunk?offset=0',content=b'abc').json()['offset']==3
    assert client.put(base+'/chunk?offset=0',content=b'abc').json()['replayed']
    assert client.put(base+'/chunk?offset=0',content=b'xxx').status_code==409
    assert client.post(base+'/complete').status_code==409
    assert client.put(base+'/chunk?offset=4',content=b'de').status_code==409
    assert client.put(base+'/chunk?offset=3',content=b'def').status_code==200
    result=client.post(base+'/complete').json()
    assert result['state']=='COMPLETE'
    assert len(client.get(base).json()['chunks'])==2
    assert client.post(base+'/complete').json()['document_id']==result['document_id']
    duplicate=upload(client,project['id'],'another.txt',b'abcdef')
    assert duplicate['state']=='DUPLICATE' and duplicate['document_id']==result['document_id']

def test_upload_limit_and_basename(client,project):
    assert client.post(f'/api/projects/{project["id"]}/uploads',json={'name':'huge.txt','size':10000000001}).status_code==413
    u=upload(client,project['id'],'../../outside.txt',b'abc')
    assert u['name']=='outside.txt'


def test_eml_upload_reaches_canonical_evidence_without_attachment_content(client,project):
    message=EmailMessage();message['Subject']='RFI 101';message['From']='contractor@example.test'
    message['To']='engineer@example.test';message.set_content('   \n')
    message.add_alternative(
        '<html><body><p>Response:</p><p>Provide Type L copper pipe.</p>'
        '<p>From: Architect &lt;a@example.test&gt; Sent: Monday To: Contractor Subject: RFI 101</p>'
        '<p>Question: May PVC be used?</p></body></html>',subtype='html')
    message.add_attachment(b'ATTACHMENT-ONLY MATERIAL',maintype='application',subtype='pdf',filename='detail.pdf')
    document=upload(client,project['id'],'rfi-response.eml',message.as_bytes())
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id']
    db=client.app.state.db;runner=client.app.state.runner
    db.execute("UPDATE runs SET status='RUNNING' WHERE id=?",(rid,))
    run=runner.get(rid);run['model']='mock';runner.parse_one(run,document['document_id'])

    rows=db.all('SELECT payload FROM evidence WHERE run_id=? ORDER BY id',(rid,))
    evidence=[json.loads(row['payload']) for row in rows]
    summary=json.loads(db.one('SELECT summary FROM document_results WHERE run_id=?',(rid,))['summary'])

    assert any('Type L copper pipe' in item['raw_text'] for item in evidence)
    assert all('ATTACHMENT-ONLY MATERIAL' not in item['raw_text'] for item in evidence)
    assert any('EMAIL > BODY > RFI 101 > RESPONSE' in (item['locator']['section'] or '') for item in evidence)
    assert any('May PVC be used?' in item['raw_text'] and
               'EMAIL > QUOTED HISTORY' in (item['locator']['section'] or '') for item in evidence)
    assert all('May PVC be used?' not in item['raw_text'] for item in evidence
               if 'EMAIL > BODY' in (item['locator']['section'] or ''))
    assert summary['workflow_contexts'][0]['role']=='RESPONSE'
    assert summary['attachments']==[{'file_name':'detail.pdf','content_type':'application/pdf',
                                     'status':'NOT_PROCESSED'}]


def test_attached_email_body_cannot_reach_parent_canonical_evidence(client,project):
    nested=EmailMessage();nested['Subject']='RFI 901';nested.set_content(
        'Question:\nATTACHMENT-ONLY: May PVC be used?')
    nested.add_attachment(b'INNER-ONLY',maintype='application',subtype='pdf',filename='inner.pdf')
    outer=EmailMessage();outer['Subject']='RFI 901';outer.set_content('Response:\nUse Type L copper.')
    outer.add_attachment(nested,filename='forwarded.eml')
    document=upload(client,project['id'],'outer.eml',outer.as_bytes())
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id']
    db=client.app.state.db;runner=client.app.state.runner
    db.execute("UPDATE runs SET status='RUNNING' WHERE id=?",(rid,))
    run=runner.get(rid);run['model']='mock';runner.parse_one(run,document['document_id'])

    evidence=[json.loads(row['payload']) for row in
              db.all('SELECT payload FROM evidence WHERE run_id=? ORDER BY id',(rid,))]
    summary=json.loads(db.one('SELECT summary FROM document_results WHERE run_id=?',(rid,))['summary'])

    assert any('Use Type L copper' in item['raw_text'] for item in evidence)
    assert all('ATTACHMENT-ONLY' not in item['raw_text'] and 'INNER-ONLY' not in item['raw_text']
               for item in evidence)
    assert summary['workflow_contexts'][0]['role']=='RESPONSE'
    assert summary['attachments']==[{'file_name':'forwarded.eml','content_type':'message/rfc822',
                                     'status':'NOT_PROCESSED'}]


def test_email_subject_routes_primary_workflow_while_body_identifier_stays_reference(client,project):
    message=EmailMessage();message['Subject']='Re: [External Email] Submittal 23 05 00-01';message.set_content(
        'RFI 42\nStatus: Approved as noted\nResponse package attached separately.')
    document=upload(client,project['id'],'RFI-42.eml',message.as_bytes())
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id']
    db=client.app.state.db;runner=client.app.state.runner
    db.execute("UPDATE runs SET status='RUNNING' WHERE id=?",(rid,))
    run=runner.get(rid);run['model']='mock';runner.parse_one(run,document['document_id'])

    items=client.get(f'/api/analysis-runs/{rid}/workflows').json()['items']
    groups={(item['kind'],item['identifier']):item for item in items}

    assert groups[('SUBMITTAL','23 05 00-01')]['members'][0]['source']=='PRIMARY'
    assert groups[('SUBMITTAL','23 05 00-01')]['members'][0]['status']=='APPROVED AS NOTED'
    assert groups[('RFI','42')]['members'][0]['source']=='REFERENCE'


def test_email_submittal_status_does_not_cross_to_different_subject_identifier(client,project):
    message=EmailMessage();message['Subject']='Submittal 23-01';message.set_content(
        'Submittal 23-02\nStatus: Rejected\nPump P-2 does not comply.')
    document=upload(client,project['id'],'submittal-scope.eml',message.as_bytes())
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id']
    db=client.app.state.db;runner=client.app.state.runner
    db.execute("UPDATE runs SET status='RUNNING' WHERE id=?",(rid,))
    run=runner.get(rid);run['model']='mock';before=db.one('SELECT COUNT(*) AS n FROM model_calls')['n']
    runner.parse_one(run,document['document_id'])

    groups={(item['kind'],item['identifier']):item for item in
            client.get(f'/api/analysis-runs/{rid}/workflows').json()['items']}
    evidence=[json.loads(row['payload']) for row in
              db.all('SELECT payload FROM evidence WHERE run_id=? ORDER BY id',(rid,))]
    pump=next(item for item in evidence if 'Pump P-2' in item['raw_text'])

    assert groups[('SUBMITTAL','23-01')]['members'][0]['source']=='PRIMARY'
    assert groups[('SUBMITTAL','23-01')]['members'][0]['status'] is None
    assert groups[('SUBMITTAL','23-02')]['members'][0]['source']=='REFERENCE'
    assert 'SUBMITTAL 23-02 > STATUS: REJECTED' in pump['locator']['section']
    assert db.one('SELECT COUNT(*) AS n FROM model_calls')['n']==before


def test_generic_email_routes_by_first_explicit_body_workflow_heading(client,project):
    message=EmailMessage();message['Subject']='Coordination';message.set_content(
        'Submittal 23-01\nStatus: Pending\nPump P-1 data.\n'
        'RFI 42\nQuestion:\nConfirm clearance.')
    document=upload(client,project['id'],'coordination.eml',message.as_bytes())
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id']
    db=client.app.state.db;runner=client.app.state.runner
    db.execute("UPDATE runs SET status='RUNNING' WHERE id=?",(rid,))
    run=runner.get(rid);run['model']='mock';before=db.one('SELECT COUNT(*) AS n FROM model_calls')['n']
    runner.parse_one(run,document['document_id'])

    groups={(item['kind'],item['identifier']):item for item in
            client.get(f'/api/analysis-runs/{rid}/workflows').json()['items']}

    assert groups[('SUBMITTAL','23-01')]['members'][0]['source']=='PRIMARY'
    assert groups[('SUBMITTAL','23-01')]['members'][0]['status']=='PENDING'
    assert groups[('RFI','42')]['members'][0]['source']=='REFERENCE'
    assert db.one('SELECT COUNT(*) AS n FROM model_calls')['n']==before


def test_full_request_for_information_name_reaches_workflow_index(client,project):
    message=EmailMessage();message['Subject']='Submittal 23-01';message.set_content(
        'Status: Pending\nSee Request for Information No. 0042 before release.')
    document=upload(client,project['id'],'submittal-rfi-reference.eml',message.as_bytes())
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id']
    db=client.app.state.db;runner=client.app.state.runner
    db.execute("UPDATE runs SET status='RUNNING' WHERE id=?",(rid,))
    run=runner.get(rid);run['model']='mock';before=db.one('SELECT COUNT(*) AS n FROM model_calls')['n']
    runner.parse_one(run,document['document_id'])

    groups={(item['kind'],item['identifier']):item for item in
            client.get(f'/api/analysis-runs/{rid}/workflows').json()['items']}

    assert groups[('SUBMITTAL','23-01')]['members'][0]['source']=='PRIMARY'
    assert groups[('RFI','42')]['members'][0]['source']=='REFERENCE'
    assert db.one('SELECT COUNT(*) AS n FROM model_calls')['n']==before


def test_conflicting_message_ids_reach_workflow_review_as_ambiguous_hash_only_metadata(client,project):
    raw=(b'Subject: Coordination\r\n'
         b'Message-ID: <first@example.test>\r\n'
         b'Message-ID: <second@example.test>\r\n'
         b'Content-Type: text/plain; charset=utf-8\r\n\r\nCurrent coordination text.')
    document=upload(client,project['id'],'conflicting-message-id.eml',raw)
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id']
    db=client.app.state.db;runner=client.app.state.runner
    db.execute("UPDATE runs SET status='RUNNING' WHERE id=?",(rid,))
    run=runner.get(rid);run['model']='mock';before=db.one('SELECT COUNT(*) AS n FROM model_calls')['n']
    runner.parse_one(run,document['document_id'])

    response=client.get(f'/api/analysis-runs/{rid}/workflows')
    summary=db.one('SELECT summary FROM document_results WHERE run_id=?',(rid,))['summary']

    assert response.status_code==200 and response.json()['items'][0]['state']=='AMBIGUOUS'
    assert 'multiple distinct Message-ID' in response.json()['items'][0]['warnings'][0]
    assert 'message_id_conflict' in summary
    assert 'first@example.test' not in summary and 'second@example.test' not in summary
    assert db.one('SELECT COUNT(*) AS n FROM model_calls')['n']==before


def test_workflow_relationship_endpoint_is_bounded_and_never_calls_model(client,project,monkeypatch):
    for name,text in [('RFI-042-question.txt','RFI 042\nQuestion:\nMay PVC be used?'),
                      ('RFI-42-response.txt','RFI 42\nResponse:\nProvide Type L copper.'),
                      ('Submittal-23-01.txt','Submittal 23-01\nStatus: Pending')]:
        upload(client,project['id'],name,text.encode())
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id']
    db=client.app.state.db;runner=client.app.state.runner
    db.execute("UPDATE runs SET status='RUNNING' WHERE id=?",(rid,))
    run=runner.get(rid);run['model']='mock'
    for did in run['document_ids']:runner.parse_one(run,did)

    stored=db.one('SELECT document_id,summary FROM document_results WHERE run_id=? ORDER BY document_id',(rid,))
    summary=json.loads(stored['summary']);summary['pages']=[{'payload':'x'*250_000}]
    full_summary=json.dumps(summary)
    db.execute('UPDATE document_results SET summary=? WHERE run_id=? AND document_id=?',
               (full_summary,rid,stored['document_id']))
    projected_bytes=[];original_all=db.all
    def measured_all(query,args=()):
        rows=original_all(query,args)
        if 'workflow_values' in query:
            projected_bytes.append(sum(len(json.dumps(row)) for row in rows))
        return rows
    monkeypatch.setattr(db,'all',measured_all)

    before=db.one('SELECT COUNT(*) AS n FROM model_calls')['n']
    response=client.get(f'/api/analysis-runs/{rid}/workflows?offset=0&limit=1')
    repeated=client.get(f'/api/analysis-runs/{rid}/workflows?offset=0&limit=1')
    second=client.get(f'/api/analysis-runs/{rid}/workflows?offset=1&limit=1')
    after=db.one('SELECT COUNT(*) AS n FROM model_calls')['n']

    assert response.status_code==200 and response.json()==repeated.json()
    assert second.status_code==200
    result=response.json();second_result=second.json()
    assert result['pagination']=={'offset':0,'limit':1,'total':2,'next_offset':1}
    assert second_result['pagination']=={'offset':1,'limit':1,'total':2,'next_offset':None}
    assert result['items'][0]['group_id']!=second_result['items'][0]['group_id']
    assert result['items'][0]['kind']=='RFI' and result['items'][0]['state']=='LINKED'
    assert result['summary']['rfi_groups']==1 and before==after
    assert len(full_summary)>250_000 and projected_bytes and max(projected_bytes)<10_000

def test_chunk_limit(client,project):
    u=client.post(f'/api/projects/{project["id"]}/uploads',json={'name':'x.txt','size':6000000}).json()
    assert client.put(f'/api/uploads/{u["id"]}/chunk',content=b'x'*(4*1024*1024+1)).status_code==413

def test_abort(client,project):
    u=client.post(f'/api/projects/{project["id"]}/uploads',json={'name':'x.txt','size':100}).json()
    assert client.post(f'/api/uploads/{u["id"]}/abort').json()['state']=='ABORTED'
    assert client.put(f'/api/uploads/{u["id"]}/chunk',content=b'x').status_code==409

def test_empty_project_no_run(client,project):assert client.post(f'/api/projects/{project["id"]}/analysis-runs').status_code==409

def test_record_summaries_are_paginated_and_keep_legacy_records(client,project):
    """Large runs list compact rows while the established full endpoint stays intact."""
    rid=run_demo(client,project['id']);rows=client.get(f'/api/analysis-runs/{rid}/records').json();db=client.app.state.db
    source=rows[0];base=source['record'];kind=base['kind']
    for number in range(505):
        envelope=deepcopy(base);record_id=f'REC-LARGE-{number:03d}'
        envelope['meta']['record_id']=record_id
        db.execute('INSERT INTO records(id,run_id,project_id,kind,envelope,logical_key) VALUES(?,?,?,?,?,?)',
                   (record_id,rid,project['id'],kind,json.dumps(envelope),f'large-{number:03d}'))
    legacy=client.get(f'/api/analysis-runs/{rid}/records').json()
    page=client.get(f'/api/analysis-runs/{rid}/record-summaries?kind={kind}&offset=300&limit=10').json()
    first=client.get(f'/api/analysis-runs/{rid}/record-summaries?kind={kind}&offset=0&limit=500').json()
    assert isinstance(legacy,list) and len(legacy)==len(rows)+505
    assert first['pagination']['next_offset']==500
    assert page['pagination']=={'offset':300,'limit':10,'total':505+sum(r['record']['kind']==kind for r in rows),'next_offset':310}
    assert len(page['items'])==min(10,page['pagination']['total']-300)
    assert page['counts'][kind]>=505 and 'candidate_key' not in page['items'][0]['record']['candidate']
    detail=client.get('/api/records/'+page['items'][0]['record']['meta']['record_id']).json()
    assert detail['record']['candidate']['candidate_key']

def test_one_active_project_and_pause(client,project):
    upload(client,project['id'],'a.txt',b'hello')
    a=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()
    assert client.post(f'/api/projects/{project["id"]}/analysis-runs').status_code==409
    assert client.post(f'/api/analysis-runs/{a["id"]}/pause').json()['status']=='PAUSED'
    assert client.post(f'/api/analysis-runs/{a["id"]}/resume').json()['status']=='QUEUED'
    assert client.post(f'/api/analysis-runs/{a["id"]}/cancel').json()['status']=='CANCELLED'


def test_selected_local_workers_parse_two_documents_concurrently(client,project,monkeypatch):
    upload(client,project['id'],'one.txt',b'one');upload(client,project['id'],'two.txt',b'two')
    run=client.post(f'/api/projects/{project["id"]}/analysis-runs',json={'local_workers':2}).json()
    assert run['capabilities']['local_workers']==2
    db=client.app.state.db;runner=client.app.state.runner
    db.execute("UPDATE runs SET status='RUNNING' WHERE id=?",(run['id'],));run=runner.get(run['id'])
    barrier=threading.Barrier(2);active=0;peak=0;lock=threading.Lock()
    def parse(run_arg,did):
        nonlocal active,peak
        with lock:active+=1;peak=max(peak,active)
        barrier.wait(timeout=2)
        with lock:active-=1
        summary={'status':'SUCCESS','fragments':[],'pages':[{'page':1,'status':'TEXT_EXTRACTED'}],
                 'warnings':[],'visual_tasks':[],'geometry_summaries':[],'takeoffs':[]}
        db.execute('INSERT INTO document_results VALUES(?,?,?,?)',(run_arg['id'],did,'SUCCESS',json.dumps(summary)))
    monkeypatch.setattr(runner,'parse_one',parse)
    runner.parse_documents(run)
    assert peak==2 and runner.get(run['id'])['coverage']['pages_processed']==2


def test_adjacent_extraction_batching_is_bounded_and_respects_legacy_single_tasks():
    def evidence(number,page,text='x'):
        return {'evidence_id':f'EV-{number}','document_id':'DOC-1','raw_text':text,
                'locator':{'page_number':page}}
    items=[(None,evidence(number,page)) for number,page in enumerate((1,1,2,3,4),1)]

    assert [len(batch) for batch in adjacent_extraction_batches(items)]==[4,1]
    assert [len(batch) for batch in adjacent_extraction_batches(items,{'EV-2'})]==[1,1,3]
    assert [len(batch) for batch in adjacent_extraction_batches([
        (None,evidence(1,1,'x'*5000)),(None,evidence(2,1,'y'*5000))])]==[1,1]


def test_mock_run_combines_adjacent_text_fragments_once(client,project,monkeypatch):
    lines=[]
    for index in range(4):
        lines.append(f'DEMO_MATERIAL|PIPE-{index}|Copper Water Pipe {index}|-|diameter=2|note='+('x'*650))
    upload(client,project['id'],'batch.txt','\n'.join(lines).encode())
    runner=client.app.state.runner;calls=[];original=runner.gateway.extract_many
    def counted(run,evidences):
        calls.append(len(evidences));return original(run,evidences)
    monkeypatch.setattr(runner.gateway,'extract_many',counted)
    run=runner.create(project['id'],local_workers=2)

    runner.process(run['id'])

    current=runner.get(run['id'])
    assert current['coverage']['fragments_total']>=2
    assert calls==[current['coverage']['fragments_total']]
    assert len(client.get(f'/api/analysis-runs/{run["id"]}/records').json())==4


def test_reconciled_adjacent_batch_recovers_as_the_same_pending_family(client,project,monkeypatch):
    from app.db import dumps
    lines=[f'DEMO_MATERIAL|PIPE-{index}|Copper Water Pipe {index}|-|note='+('x'*650) for index in range(4)]
    upload(client,project['id'],'resume-batch.txt','\n'.join(lines).encode())
    runner=client.app.state.runner;db=client.app.state.db;run=runner.create(project['id'])
    db.execute("UPDATE runs SET status='RUNNING' WHERE id=?",(run['id'],));run=runner.get(run['id'])
    runner.parse_one(run,run['document_ids'][0])
    evidence=[json.loads(row['payload']) for row in db.all(
        'SELECT payload FROM evidence WHERE run_id=? ORDER BY rowid',(run['id'],))]
    assert len(evidence)==2
    ids=[item['evidence_id'] for item in evidence]
    family='extract-batch:'+ids[0]+':'+hashlib.sha256(dumps(ids).encode()).hexdigest()[:16]
    call=db.reserve(project['id'],run['id'],family,Decimal('0.1'),'model','hash',Decimal('1'),Decimal('2'))
    db.unknown(call,dumps({'kind':'NETWORK_ERROR','class':'TIMEOUT'}))
    db.execute("UPDATE runs SET status='PAUSED_PROVIDER' WHERE id=?",(run['id'],))
    client.post(f'/api/model-calls/{call}/reconcile',json={
        'resolution':'NOT_BILLED','actual_cny':'0','confirmation':'PROVIDER_BILLING_CHECKED'}).raise_for_status()

    resumed=client.post(f'/api/analysis-runs/{run["id"]}/resume')

    assert resumed.status_code==200 and resumed.json()['status']=='QUEUED'
    event=json.loads(db.one("SELECT payload FROM call_reconciliation_events WHERE payload LIKE '%RECOVERY_GENERATION_AUTHORIZED%'")['payload'])
    assert event['task_family']==family and event['generation']==1
    calls=[];original=runner.gateway.extract_many
    monkeypatch.setattr(runner.gateway,'extract_many',lambda current,items:(calls.append(len(items)),original(current,items))[1])
    runner.process(run['id'])
    assert calls==[2]
    assert {row['status'] for row in db.all('SELECT status FROM evidence WHERE run_id=?',(run['id'],))}=={'EXTRACTED'}


def test_completed_mock_run_records_stage_performance(client,project):
    rid=run_demo(client,project['id']);run=client.get(f'/api/analysis-runs/{rid}').json()
    assert {'stage.parse','stage.vision','stage.extract','stage.assemble','stage.verify'}<=set(run['performance'])
    assert all(value['samples']==1 and value['total_ms']>=0 for value in run['performance'].values())
    assert run['progress']['percent']==100 and run['progress']['estimated_finish_epoch'] is None


def test_running_progress_reports_percentage_and_estimated_finish(client,project):
    upload(client,project['id'],'a.txt',b'hello')
    runner=client.app.state.runner;db=client.app.state.db;run=runner.create(project['id'])
    coverage={'fragments_total':10,'fragments_extracted':4,'fragments_need_review':1}
    db.execute("UPDATE runs SET status='RUNNING',stage=?,coverage=? WHERE id=?",
               ('一次读取，联合提取材料与检查要求',json.dumps(coverage),run['id']))
    current=runner.get(run['id']);progress=runner.progress(current,current['started_epoch']+100)
    assert progress['percent']==57
    assert progress['estimate'] and progress['remaining_seconds']>0
    assert progress['estimated_finish_epoch']>current['started_epoch']+100

def test_failed_local_publish_can_resume_without_repeating_model_work(client,project):
    upload(client,project['id'],'a.txt',b'hello')
    run=client.post(f'/api/projects/{project["id"]}/analysis-runs').json();rid=run['id']
    db=client.app.state.db
    db.execute("UPDATE runs SET status='FAILED',stage='生成可审核记录与设计差异',message='内部处理错误：ValidationError' WHERE id=?",(rid,))

    resumed=client.post(f'/api/analysis-runs/{rid}/resume')

    assert resumed.status_code==200,resumed.text
    payload=resumed.json()
    assert payload['status']=='QUEUED'
    assert '已完成的模型调用不会重发' in payload['message']
    assert client.get(f'/api/analysis-runs/{rid}/cost').json()['calls']==0

def test_failed_run_outside_local_publish_stage_stays_closed(client,project):
    upload(client,project['id'],'a.txt',b'hello')
    run=client.post(f'/api/projects/{project["id"]}/analysis-runs').json();rid=run['id']
    client.app.state.db.execute("UPDATE runs SET status='FAILED',stage='解析文件' WHERE id=?",(rid,))
    assert client.post(f'/api/analysis-runs/{rid}/resume').status_code==409


def test_old_provider_run_cannot_resume_or_process_under_current_service(client,project):
    upload(client,project['id'],'a.txt',b'hello')
    run=client.post(f'/api/projects/{project["id"]}/analysis-runs').json();rid=run['id']
    db=client.app.state.db;runner=client.app.state.runner
    db.execute("UPDATE runs SET status='PAUSED',provider='deepseek' WHERE id=?",(rid,))
    assert client.post(f'/api/analysis-runs/{rid}/resume').status_code==409
    db.execute("UPDATE runs SET status='QUEUED' WHERE id=?",(rid,))
    runner.process(rid)
    current=client.get(f'/api/analysis-runs/{rid}').json()
    assert current['status']=='PAUSED_PROVIDER' and '未调用API' in current['message']
    assert client.get(f'/api/analysis-runs/{rid}/records').json()==[]
    assert client.get(f'/api/analysis-runs/{rid}/cost').json()['calls']==0


def test_unknown_call_requires_explicit_reconciliation_before_resume(client,project):
    upload(client,project['id'],'a.txt',b'hello')
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id']
    db=client.app.state.db;runner=client.app.state.runner
    db.execute("UPDATE runs SET status='RUNNING' WHERE id=?",(rid,))
    run=runner.get(rid);runner.parse_one(run,run['document_ids'][0])
    pending=json.loads(db.one("SELECT payload FROM evidence WHERE run_id=? AND status='PENDING'",(rid,))['payload'])
    call_id=db.reserve(project['id'],rid,pending['evidence_id'],Decimal('1'),'model','hash',Decimal('1'),Decimal('2'))
    db.unknown(call_id,json.dumps({'kind':'HTTP_STATUS','status':429,'class':'RATE_LIMIT_OR_QUOTA',
                                   'provider_code':'RESOURCE_EXHAUSTED',
                                   'retry_after':'Wed, 21 Oct 2015 07:28:00 GMT',
                                   'provider_request_id':'unsafe request id',
                                   'message':'must-not-leak'}),
               'safe-request-1')
    db.execute("UPDATE runs SET status='PAUSED_PROVIDER' WHERE id=?",(rid,))
    unresolved=client.get(f'/api/projects/{project["id"]}/unresolved-model-calls').json()
    assert len(unresolved)==1 and unresolved[0]['id']==call_id
    assert unresolved[0]['diagnostic']=={'kind':'HTTP_STATUS','status':429,
                                         'class':'RATE_LIMIT_OR_QUOTA','provider_code':'RESOURCE_EXHAUSTED',
                                         'retry_after':'Wed, 21 Oct 2015 07:28:00 GMT'}
    assert 'must-not-leak' not in json.dumps(unresolved)
    assert client.post(f'/api/analysis-runs/{rid}/resume').status_code==409
    assert client.post(f'/api/model-calls/{call_id}/reconcile',json={
        'resolution':'NOT_BILLED','actual_cny':'0'}).status_code==422
    assert client.post(f'/api/model-calls/{call_id}/reconcile',json={
        'resolution':'NOT_BILLED','actual_cny':'1','confirmation':'PROVIDER_BILLING_CHECKED'}).status_code==400
    result=client.post(f'/api/model-calls/{call_id}/reconcile',json={
        'resolution':'NOT_BILLED','actual_cny':'0','confirmation':'PROVIDER_BILLING_CHECKED',
        'note':'Checked provider billing'}).json()
    assert result['call']['state']=='RECONCILED_ZERO'
    assert result['cost']['spent_cny']=='0.000000' and result['cost']['reserved_cny']=='0.000000'
    assert client.get(f'/api/projects/{project["id"]}/unresolved-model-calls').json()==[]
    events=client.get(f'/api/projects/{project["id"]}/call-reconciliation-events').json()
    assert len(events)==1 and events[0]['call_id']==call_id
    assert events[0]['payload']['resolution']=='NOT_BILLED'
    assert client.post(f'/api/model-calls/{call_id}/reconcile',json={
        'resolution':'NOT_BILLED','actual_cny':'0','confirmation':'PROVIDER_BILLING_CHECKED'}).status_code==409
    resumed=client.post(f'/api/analysis-runs/{rid}/resume').json()
    assert resumed['status']=='QUEUED' and '1个显式付费恢复代次' in resumed['message']
    events=client.get(f'/api/projects/{project["id"]}/call-reconciliation-events').json()
    assert [event['payload']['event_type'] for event in events]==[
        'CALL_RECONCILED','RECOVERY_GENERATION_AUTHORIZED']
    assert events[1]['payload']['next_task_key']==pending['evidence_id']+':manual-requeue:1'


def test_reconciled_unknown_task_key_fails_closed_without_generation_or_http(client,project):
    upload(client,project['id'],'a.txt',b'hello')
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id']
    db=client.app.state.db
    db.execute("UPDATE runs SET status='RUNNING' WHERE id=?",(rid,))
    call_id=db.reserve(project['id'],rid,'EV-not-a-real-pending-evidence',Decimal('0.1'),
                       'model','hash',Decimal('1'),Decimal('2'))
    db.unknown(call_id,json.dumps({'kind':'NETWORK_ERROR','class':'TIMEOUT'}))
    db.execute("UPDATE runs SET status='PAUSED_PROVIDER' WHERE id=?",(rid,))
    db.reconcile_call(call_id,'NOT_BILLED',Decimal('0'),'offline provider ledger check')
    blocked=client.post(f'/api/analysis-runs/{rid}/resume')
    assert blocked.status_code==409 and '不对应真实待处理证据' in blocked.json()['detail']
    assert db.one("SELECT status FROM runs WHERE id=?",(rid,))['status']=='PAUSED_PROVIDER'
    events=db.reconciliation_events(project['id'])
    assert len(events)==1 and events[0]['payload']['event_type']=='CALL_RECONCILED'
    assert db.one("SELECT COUNT(*) AS n FROM model_calls WHERE run_id=?",(rid,))['n']==1


@pytest.mark.parametrize(('resolution','actual_cny','old_state'),(
    ('NOT_BILLED','0','RECONCILED_ZERO'),
    ('BILLED','0.000001','RECONCILED_CHARGED'),
))
def test_real_pending_extraction_reconcile_resume_sends_exactly_one_new_generation_http(
        client,project,monkeypatch,resolution,actual_cny,old_state):
    """Exercise the former deadlock with the real pending evidence/task key."""
    upload(client,project['id'],'pending.txt',
           b'DEMO_MATERIAL|M-1|Concrete|-|strength=5000 psi')
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id']
    db=client.app.state.db;runner=client.app.state.runner
    live=replace(runner.s,provider='deepseek',live_enabled=True,prices_confirmed=True,
                 api_key='offline-not-real',input_rate=Decimal('1'),output_rate=Decimal('2'))
    requests=[]

    def handler(request):
        requests.append(request)
        if len(requests)==1:
            raise httpx.ReadTimeout('synthetic uncertain request')
        row=db.one("SELECT payload FROM evidence WHERE run_id=? AND status='PENDING'",(rid,))
        evidence=json.loads(row['payload']);data=mock_extract(evidence)
        return httpx.Response(200,json={
            'id':'explicit-generation-1',
            'choices':[{'finish_reason':'stop','message':{'content':json.dumps(data)}}],
            'usage':{'prompt_tokens':100,'completion_tokens':40},
        })

    gateway=Gateway(live,db,httpx.Client(transport=httpx.MockTransport(handler)))
    runner.s=live;runner.gateway=gateway
    monkeypatch.setattr(runner.verifier,'refresh',lambda *args,**kwargs: {})
    db.execute("UPDATE runs SET provider='deepseek' WHERE id=?",(rid,))
    runner.process(rid)

    pending=db.one("SELECT id,payload,status FROM evidence WHERE run_id=?",(rid,))
    evidence_id=json.loads(pending['payload'])['evidence_id']
    first=db.one("SELECT id,state,task_key FROM model_calls WHERE run_id=?",(rid,))
    assert len(requests)==1 and pending['status']=='PENDING'
    assert first['state']=='UNKNOWN' and first['task_key']==evidence_id
    assert runner.get(rid)['status']=='PAUSED_PROVIDER'

    reconciled=client.post(f'/api/model-calls/{first["id"]}/reconcile',json={
        'resolution':resolution,'actual_cny':actual_cny,
        'confirmation':'PROVIDER_BILLING_CHECKED','note':'offline provider ledger check'})
    assert reconciled.status_code==200,reconciled.text
    assert reconciled.json()['call']['state']==old_state
    resumed=client.post(f'/api/analysis-runs/{rid}/resume')
    assert resumed.status_code==200,resumed.text
    runner.process(rid)

    calls=db.all("SELECT id,state,task_key,response FROM model_calls WHERE run_id=? ORDER BY created_at,id",(rid,))
    assert len(requests)==2 and len(calls)==2
    assert [row['task_key'] for row in calls]==[
        evidence_id,evidence_id+':manual-requeue:1']
    assert calls[0]['id']==first['id'] and calls[0]['state']==old_state and calls[0]['response'] is None
    assert calls[1]['state']=='SETTLED' and calls[1]['response'] is not None
    assert db.one("SELECT status FROM evidence WHERE run_id=?",(rid,))['status']=='EXTRACTED'
    events=db.reconciliation_events(project['id'])
    assert [event['payload']['event_type'] for event in events]==[
        'CALL_RECONCILED','RECOVERY_GENERATION_AUTHORIZED']
    assert events[1]['payload']['reconciliation_resolution']==resolution
    assert events[1]['payload']['attempt_number']==2
    gateway.close()


def test_reconciled_paused_visual_task_resumes_with_one_audited_generation_http(client,project):
    document=upload(client,project['id'],'drawing.pdf',b'offline-placeholder')
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id']
    db=client.app.state.db;runner=client.app.state.runner
    live=replace(runner.s,provider='deepseek',live_enabled=True,prices_confirmed=True,
                 api_key='offline-not-real',vision_enabled=True,
                 input_rate=Decimal('1'),output_rate=Decimal('2'))
    db.execute("UPDATE runs SET status='RUNNING',provider='deepseek' WHERE id=?",(rid,))
    summary={'warnings':[],'visual_tasks':[{'page':1,'status':'PAUSED_PROVIDER'}]}
    db.execute('INSERT INTO document_results VALUES(?,?,?,?)',
               (rid,document['document_id'],'PARTIAL',json.dumps(summary)))
    family=f'vision:{document["document_id"]}:1'
    first=db.reserve(project['id'],rid,family,Decimal('0.1'),'deepseek-v4-flash-vision-exp',
                     'hash',Decimal('1'),Decimal('2'))
    db.unknown(first,json.dumps({'kind':'NETWORK_ERROR','class':'TIMEOUT'}))
    db.execute("UPDATE runs SET status='PAUSED_PROVIDER' WHERE id=?",(rid,))
    db.reconcile_call(first,'NOT_BILLED',Decimal('0'),'offline provider ledger check')
    runner.s=live
    resumed=client.post(f'/api/analysis-runs/{rid}/resume')
    assert resumed.status_code==200,resumed.text
    stored=json.loads(db.one('SELECT summary FROM document_results WHERE run_id=?',(rid,))['summary'])
    task=stored['visual_tasks'][0]
    assert task['status']=='PENDING' and task['billing_generation']==1

    db.execute("UPDATE runs SET status='RUNNING' WHERE id=?",(rid,));run=runner.get(rid);requests=[]
    vision={'page_type':'DRAWING','sheet_id':'A-1','important_visible_text':'Visible plan.',
            'observations':[],'explicit_quantity_texts':[],'scale_text':None,
            'limitations':['Review source.'],'needs_review':True}
    response={'id':'visual-generation-1','choices':[{'finish_reason':'stop','message':{'content':json.dumps(vision)}}],
              'usage':{'prompt_tokens':500,'completion_tokens':100}}
    gateway=Gateway(live,db,httpx.Client(transport=httpx.MockTransport(
        lambda request:(requests.append(request),httpx.Response(200,json=response))[1])))
    result=gateway.vision(run,f'{document["document_id"]}:1',b'fake-png',{
        'document_id':document['document_id'],'page':1,'coordinate_system':'image-pixels-top-left',
        'width':80,'height':60},billing_generation=task['billing_generation'])
    calls=db.all("SELECT state,task_key FROM model_calls WHERE run_id=? ORDER BY created_at,id",(rid,))
    assert result.data==vision and len(requests)==1
    assert [row['task_key'] for row in calls]==[family,family+':manual-requeue:1']
    assert [row['state'] for row in calls]==['RECONCILED_ZERO','SETTLED']
    gateway.close()


def test_extraction_family_stops_after_three_explicit_reconciled_generations(client,project,monkeypatch):
    upload(client,project['id'],'bounded.txt',b'One real pending evidence fragment.')
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id']
    db=client.app.state.db;runner=client.app.state.runner
    live=replace(runner.s,provider='deepseek',live_enabled=True,prices_confirmed=True,
                 api_key='offline-not-real',input_rate=Decimal('1'),output_rate=Decimal('2'))
    requests=[]
    gateway=Gateway(live,db,httpx.Client(transport=httpx.MockTransport(
        lambda request:(requests.append(request),(_ for _ in ()).throw(
            httpx.ReadTimeout('synthetic repeated uncertainty')))[1])))
    runner.s=live;runner.gateway=gateway
    monkeypatch.setattr(runner.verifier,'refresh',lambda *args,**kwargs: {})
    db.execute("UPDATE runs SET provider='deepseek' WHERE id=?",(rid,))
    runner.process(rid)
    for generation in (1,2):
        unknown=db.one("SELECT id FROM model_calls WHERE run_id=? AND actual_units IS NULL",(rid,))
        db.reconcile_call(unknown['id'],'NOT_BILLED',Decimal('0'),'offline provider ledger check')
        response=client.post(f'/api/analysis-runs/{rid}/resume')
        assert response.status_code==200,response.text
        runner.process(rid)
    third=db.one("SELECT id FROM model_calls WHERE run_id=? AND actual_units IS NULL",(rid,))
    db.reconcile_call(third['id'],'NOT_BILLED',Decimal('0'),'offline provider ledger check')
    blocked=client.post(f'/api/analysis-runs/{rid}/resume')
    assert blocked.status_code==409 and '三次累计调用上限' in blocked.json()['detail']
    calls=db.all("SELECT task_key,state FROM model_calls WHERE run_id=? ORDER BY created_at,id",(rid,))
    family=calls[0]['task_key']
    assert len(requests)==3 and [row['task_key'] for row in calls]==[
        family,family+':manual-requeue:1',family+':manual-requeue:2']
    assert all(row['state']=='RECONCILED_ZERO' for row in calls)
    gateway.close()


def test_unresolved_call_in_another_run_blocks_project_resume(client,project):
    upload(client,project['id'],'a.txt',b'hello')
    first=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id']
    db=client.app.state.db
    db.execute("UPDATE runs SET status='PAUSED' WHERE id=?",(first,))
    second=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id']
    db.execute("UPDATE runs SET status='PAUSED_PROVIDER' WHERE id=?",(second,))
    db.execute("UPDATE runs SET status='RUNNING' WHERE id=?",(first,))
    call_id=db.reserve(project['id'],first,'EV-old',Decimal('0.1'),'model','hash',Decimal('1'),Decimal('2'))
    db.unknown(call_id,json.dumps({'kind':'NETWORK_ERROR','class':'TIMEOUT'}))
    db.execute("UPDATE runs SET status='PAUSED_PROVIDER' WHERE id=?",(first,))
    assert client.post(f'/api/projects/{project["id"]}/analysis-runs').status_code==409
    db.execute("UPDATE runs SET status='QUEUED' WHERE id=?",(second,));client.app.state.runner.process(second)
    assert client.get(f'/api/analysis-runs/{second}').json()['status']=='PAUSED_PROVIDER'
    response=client.post(f'/api/analysis-runs/{second}/resume')
    assert response.status_code==409 and '待对账' in response.json()['detail']


def test_interrupted_reserved_call_becomes_reconcilable_unknown(client,project):
    upload(client,project['id'],'a.txt',b'hello')
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id'];db=client.app.state.db
    db.execute("UPDATE runs SET status='RUNNING' WHERE id=?",(rid,))
    call_id=db.reserve(project['id'],rid,'EV-interrupted',Decimal('0.1'),'model','hash',Decimal('1'),Decimal('2'))
    runner=client.app.state.runner;runner.start()
    try:
        call=db.one('SELECT state,error FROM model_calls WHERE id=?',(call_id,))
        assert call['state']=='UNKNOWN' and json.loads(call['error'])=={
            'kind':'CLIENT_ERROR','class':'PROCESS_INTERRUPTED'}
        assert client.get(f'/api/analysis-runs/{rid}').json()['status']=='INTERRUPTED'
        assert client.get(f'/api/projects/{project["id"]}/unresolved-model-calls').json()[0]['state']=='UNKNOWN'
    finally:runner.close()


def test_resume_processes_only_pending_fragments_after_reconciliation(client,project,monkeypatch):
    from app.db import dumps
    from app.gateway import mock_extract

    upload(client,project['id'],'one.txt',b'DEMO_MATERIAL|M-1|Concrete|-|strength=5000 psi')
    upload(client,project['id'],'two.txt',b'DEMO_TEST|T-2|Concrete cylinder test|-|frequency=each pour')
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id']
    db=client.app.state.db;runner=client.app.state.runner
    db.execute("UPDATE runs SET status='RUNNING' WHERE id=?",(rid,))
    run=runner.get(rid);run['model']='mock'
    for did in run['document_ids']:runner.parse_one(run,did)
    rows=db.all('SELECT id,payload FROM evidence WHERE run_id=? ORDER BY id',(rid,))
    completed=rows[0];completed_ev=json.loads(completed['payload'])
    db.execute("UPDATE evidence SET status='EXTRACTED',extraction=? WHERE id=?",
               (dumps({'data':mock_extract(completed_ev),'request_id':None,'cached':False}),completed['id']))
    pending_ev=json.loads(rows[1]['payload'])
    call_id=db.reserve(project['id'],rid,pending_ev['evidence_id'],Decimal('0.1'),'model','hash',Decimal('1'),Decimal('2'))
    db.unknown(call_id,dumps({'kind':'NETWORK_ERROR','class':'TIMEOUT'}))
    db.execute("UPDATE runs SET status='PAUSED_PROVIDER' WHERE id=?",(rid,))
    client.post(f'/api/model-calls/{call_id}/reconcile',json={
        'resolution':'NOT_BILLED','actual_cny':'0','confirmation':'PROVIDER_BILLING_CHECKED'}).raise_for_status()
    assert client.post(f'/api/analysis-runs/{rid}/resume').json()['status']=='QUEUED'
    called=[];original=runner.gateway.extract
    def tracked(run_arg,evidence):
        called.append(evidence['evidence_id'])
        return original(run_arg,evidence)
    monkeypatch.setattr(runner.gateway,'extract',tracked)
    monkeypatch.setattr(runner.verifier,'refresh',lambda *args,**kwargs: {})
    runner.process(rid)
    assert called==[json.loads(rows[1]['payload'])['evidence_id']]
    assert db.one('SELECT COUNT(*) AS n FROM document_results WHERE run_id=?',(rid,))['n']==2
    assert [x['status'] for x in db.all('SELECT status FROM evidence WHERE run_id=? ORDER BY id',(rid,))]==['EXTRACTED','EXTRACTED']
    current=client.get(f'/api/analysis-runs/{rid}').json()
    assert current['status']=='PARTIAL' and current['coverage']['fragments_extracted']==2


def test_resume_rebuilds_merged_record_and_invalidates_early_human_review(client,project,monkeypatch):
    from app.db import dumps
    from app.gateway import mock_extract

    for name,suffix in (('one.txt',b'\nONE'),('two.txt',b'\nTWO')):
        upload(client,project['id'],name,b'DEMO_MATERIAL|M-1|Concrete|-|strength=placeholder'+suffix)
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id'];db=client.app.state.db
    runner=client.app.state.runner;db.execute("UPDATE runs SET status='RUNNING' WHERE id=?",(rid,))
    run=runner.get(rid);run['model']='mock'
    for did in run['document_ids']:runner.parse_one(run,did)
    rows=db.all('SELECT id,payload FROM evidence WHERE run_id=? ORDER BY id',(rid,))
    assert len(rows)==2
    # Publish the lexicographically later evidence first.  The resumed item
    # sorts before it, covering the former record-id instability directly.
    early_row,new_row=rows[1],rows[0]
    for row,value,date in ((early_row,'4000 psi','2026-01-01'),(new_row,'5000 psi','2026-02-01')):
        ev=json.loads(row['payload']);ev['raw_text']=f'DEMO_MATERIAL|M-1|Concrete|-|strength={value}'
        ev['internal_revision_date']=date;db.execute('UPDATE evidence SET payload=? WHERE id=?',(dumps(ev),row['id']))
    first_ev=json.loads(db.one('SELECT payload FROM evidence WHERE id=?',(early_row['id'],))['payload'])
    db.execute("UPDATE evidence SET status='EXTRACTED',extraction=? WHERE id=?",
               (dumps({'data':mock_extract(first_ev),'request_id':None,'cached':False}),early_row['id']))
    monkeypatch.setattr(runner.verifier,'refresh',lambda *args,**kwargs: {})
    runner.publish(run)
    early=db.one("SELECT id,envelope FROM records WHERE run_id=? AND kind='MATERIAL'",(rid,))
    early_record=json.loads(early['envelope'])
    assert early_record['candidate']['design_properties'][0]['value']=='4000 psi'
    # Simulate a record persisted by the pre-stable-key build.  Recovery must
    # migrate it in place rather than inserting a second MG-key record.
    legacy_id='REC-legacy-material';legacy_key='C-legacy-material'
    early_record['meta']['record_id']=legacy_id;early_record['candidate']['candidate_key']=legacy_key
    db.execute('UPDATE records SET id=?,envelope=?,logical_key=? WHERE id=?',
               (legacy_id,dumps(early_record),legacy_key,early['id']))
    early={'id':legacy_id,'envelope':dumps(early_record)}
    db.execute("UPDATE runs SET status='PAUSED_PROVIDER' WHERE id=?",(rid,))
    client.post(f'/api/records/{early["id"]}/review',json={'action':'ACCEPTED','expected_version':0}).raise_for_status()
    db.execute("UPDATE runs SET status='RUNNING' WHERE id=?",(rid,))
    pending_ev=json.loads(new_row['payload'])
    call_id=db.reserve(project['id'],rid,pending_ev['evidence_id'],Decimal('0.1'),'model','hash',Decimal('1'),Decimal('2'))
    db.unknown(call_id,dumps({'kind':'NETWORK_ERROR','class':'TIMEOUT'}))
    db.execute("UPDATE runs SET status='PAUSED_PROVIDER' WHERE id=?",(rid,))
    client.post(f'/api/model-calls/{call_id}/reconcile',json={
        'resolution':'NOT_BILLED','actual_cny':'0','confirmation':'PROVIDER_BILLING_CHECKED'}).raise_for_status()
    client.post(f'/api/analysis-runs/{rid}/resume').raise_for_status();runner.process(rid)
    current=db.one('SELECT envelope,review_version FROM records WHERE id=?',(early['id'],))
    record=json.loads(current['envelope'])
    assert record['candidate']['design_properties'][0]['value']=='5000 psi'
    assert len(record['candidate']['evidence_ids'])==2 and record['review']['status']=='PENDING'
    assert db.one("SELECT COUNT(*) AS n FROM records WHERE run_id=? AND kind='MATERIAL'",(rid,))['n']==1
    assert current['review_version']==2
    assert not db.one("SELECT envelope FROM records WHERE run_id=? AND kind='CONFLICT'",(rid,),required=False)
    history=client.get(f'/api/records/{early["id"]}/history').json()
    assert [event['action'] for event in history]==['ACCEPTED','INVALIDATED_BY_NEW_EVIDENCE']


@pytest.mark.parametrize(('field','edited_value'),(
    ('support_note','工程师补充说明，必须原样保留。'),
    ('name','Engineer edited material name'),
))
def test_legacy_material_identity_migration_preserves_human_edits(client,project,monkeypatch,field,edited_value):
    from app.db import dumps
    from app.gateway import mock_extract

    upload(client,project['id'],'legacy.txt',b'DEMO_MATERIAL|M-1|Concrete|-|strength=4000 psi')
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id'];db=client.app.state.db
    runner=client.app.state.runner;db.execute("UPDATE runs SET status='RUNNING' WHERE id=?",(rid,))
    run=runner.get(rid);run['model']='mock';runner.parse_one(run,run['document_ids'][0])
    row=db.one('SELECT id,payload FROM evidence WHERE run_id=?',(rid,));evidence=json.loads(row['payload'])
    db.execute("UPDATE evidence SET status='EXTRACTED',extraction=? WHERE id=?",
               (dumps({'data':mock_extract(evidence),'request_id':None,'cached':False}),row['id']))
    monkeypatch.setattr(runner.verifier,'refresh',lambda *args,**kwargs: {})
    runner.publish(run)
    original=db.one("SELECT id,envelope FROM records WHERE run_id=? AND kind='MATERIAL'",(rid,))
    envelope=json.loads(original['envelope']);legacy_id='REC-legacy-'+field;legacy_key='C-legacy-'+field
    envelope['meta']['record_id']=legacy_id;envelope['candidate']['candidate_key']=legacy_key
    db.execute('UPDATE records SET id=?,envelope=?,logical_key=? WHERE id=?',
               (legacy_id,dumps(envelope),legacy_key,original['id']))
    db.execute("UPDATE runs SET status='PAUSED_PROVIDER' WHERE id=?",(rid,))
    edited=deepcopy(envelope['candidate']);edited[field]=edited_value
    response=client.post(f'/api/records/{legacy_id}/review',json={
        'action':'EDITED','expected_version':0,'note':'保留人工修改','candidate':edited})
    assert response.status_code==200,response.text

    runner.publish(run)
    current=db.one('SELECT envelope,review_version,logical_key FROM records WHERE id=?',(legacy_id,))
    candidate=json.loads(current['envelope'])['candidate']
    assert candidate[field]==edited_value and candidate['candidate_key'].startswith('MG-')
    assert json.loads(current['envelope'])['review']['status']=='EDITED'
    assert current['review_version']==1 and current['logical_key']==candidate['candidate_key']
    assert db.one("SELECT COUNT(*) AS n FROM records WHERE run_id=? AND kind='MATERIAL'",(rid,))['n']==1
    history=client.get(f'/api/records/{legacy_id}/history').json()
    assert [event['action'] for event in history]==['EDITED']


def test_reconciled_charge_is_recorded_and_over_reservation_freezes_budget(client,project):
    upload(client,project['id'],'a.txt',b'hello')
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id']
    db=client.app.state.db
    db.execute("UPDATE runs SET status='RUNNING' WHERE id=?",(rid,))
    call_id=db.reserve(project['id'],rid,'EV-test',Decimal('0.1'),'model','hash',Decimal('1'),Decimal('2'))
    db.unknown(call_id,json.dumps({'kind':'NETWORK_ERROR','class':'TIMEOUT'}))
    result=client.post(f'/api/model-calls/{call_id}/reconcile',json={
        'resolution':'BILLED','actual_cny':'1.25','confirmation':'PROVIDER_BILLING_CHECKED'}).json()
    assert result['call']['state']=='RECONCILED_CHARGED'
    assert result['cost']['spent_cny']=='1.250000' and result['cost']['reserved_cny']=='0.000000'
    assert result['cost']['frozen']

def test_mock_end_to_end_revision_and_export(client,project):
    rid=run_demo(client,project['id'])
    run=client.get(f'/api/analysis-runs/{rid}').json()
    assert run['status']=='PARTIAL' and run['coverage']['files_processed']==2
    rows=client.get(f'/api/analysis-runs/{rid}/records').json()
    mats=[x['record'] for x in rows if x['record']['kind']=='MATERIAL']
    assert len(mats)==2
    pad=next(m for m in mats if 'concrete' in m['candidate']['name'])
    assert pad['candidate']['design_properties'][0]['value']=='5000 psi'
    assert not any(x['record']['kind']=='CONFLICT' for x in rows)
    eid=pad['candidate']['evidence_ids'][0]
    source=client.get(f'/api/analysis-runs/{rid}/evidence/{eid}').json()
    assert pad['meta']['parser_version']==source['evidence']['parser_version']==PARSER_VERSION
    assert source['evidence']['raw_text'] and source['file_name']
    assert client.get(f'/api/analysis-runs/{rid}/cost').json()['calls']==0
    data=client.get(f'/api/analysis-runs/{rid}/exports/json').json()
    assert data['export_version']=='0.2.6-readable-en-2'
    assert data['summary']['run_status']=='Partial'
    assert len(data['materials_and_equipment'])==2
    assert len(data['inspections_and_tests'])==2
    assert 'conflicts' not in data and 'missing_information' not in data
    exported_pad=next(item for item in data['materials_and_equipment'] if 'concrete' in item['material_or_equipment_name'])
    assert exported_pad['quantity'] is None and exported_pad['evidence']
    readable_json=json.dumps(data,ensure_ascii=False)
    assert '"record_id"' not in readable_json and '"evidence_id"' not in readable_json
    export=client.get(f'/api/analysis-runs/{rid}/exports/xlsx')
    wb=load_workbook(io.BytesIO(export.content))
    assert {'Summary','Materials & Equipment','Inspections & Tests'}==set(wb.sheetnames)
    assert wb['Materials & Equipment']['A4'].value=='Material or Equipment Name'

def test_review_and_optimistic_lock(client,project):
    rid=run_demo(client,project['id'])
    row=client.get(f'/api/analysis-runs/{rid}/records').json()[0];record_id=row['record']['meta']['record_id']
    uri=f'/api/records/{record_id}/review'
    assert client.post(uri,json={'action':'ACCEPTED','expected_version':0}).status_code==200
    assert client.post(uri,json={'action':'REJECTED','expected_version':0}).status_code==409
    assert len(client.get(f'/api/records/{record_id}/history').json())==1
    exported=client.get(f'/api/analysis-runs/{rid}/exports/json?reviewed_only=true').json()
    assert sum(len(exported[key]) for key in ('materials_and_equipment','inspections_and_tests'))==1

def test_edit_validates_evidence_and_keeps_meta(client,project):
    rid=run_demo(client,project['id'])
    row=client.get(f'/api/analysis-runs/{rid}/records?kind=MATERIAL').json()[0];c=deepcopy(row['record']['candidate'])
    uri=f'/api/records/{row["record"]["meta"]["record_id"]}/review'
    c['evidence_ids']=['EV-foreign']
    r=client.post(uri,json={'action':'EDITED','expected_version':0,'note':'edit','candidate':c})
    assert r.status_code==400
    c=deepcopy(row['record']['candidate']);c['support_note']='人工说明'
    assert client.post(uri,json={'action':'EDITED','expected_version':0,'note':'说明更新','candidate':c}).status_code==200

def test_real_text_not_faked_by_mock(client,project):
    upload(client,project['id'],'real.txt','实际项目：请提供混凝土。'.encode())
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id'];client.app.state.runner.process(rid)
    assert client.get(f'/api/analysis-runs/{rid}/records?kind=MATERIAL').json()==[]
    assert client.get(f'/api/analysis-runs/{rid}/records?kind=MISSING').json()==[]

def test_unsupported_not_silently_dropped(client,project):
    upload(client,project['id'],'drawing.dwg',b'AC1032-not-real')
    upload(client,project['id'],'arbitrary.xlsx',b'not-supported')
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id'];client.app.state.runner.process(rid)
    run=client.get(f'/api/analysis-runs/{rid}').json()
    assert run['status']=='PARTIAL' and run['coverage']['files_partial_or_failed']==2
    assert client.get(f'/api/analysis-runs/{rid}/records?kind=MISSING').json()==[]


def test_dxf_takeoff_endpoint_keeps_object_count_pending_human_review(client,project,tmp_path):
    import ezdxf
    path=tmp_path/'objects.dxf';doc=ezdxf.new();doc.units=4
    block=doc.blocks.new('ANCHOR');block.add_circle((0,0),1)
    model=doc.modelspace();model.add_blockref('ANCHOR',(10,10));model.add_blockref('ANCHOR',(20,20))
    doc.saveas(path);upload(client,project['id'],path.name,path.read_bytes())
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id'];client.app.state.runner.process(rid)
    response=client.get(f'/api/analysis-runs/{rid}/takeoffs')
    assert response.status_code==200
    item=next(x for x in response.json()[0]['takeoffs'] if x['kind']=='BLOCK_COUNT')
    assert (item['label'],item['value'],item['unit'],item['review_status'])==('ANCHOR',2,'EA','PENDING')
    assert item['evidence_ids'] and item['method']=='CAD_OBJECT_COUNT'
    exported=client.get(f'/api/analysis-runs/{rid}/exports/json').json()
    exported_item=exported['quantity_takeoffs'][0]
    assert exported_item['quantity']==2 and exported_item['human_review']=='Pending human review'
    reviewed=client.get(f'/api/analysis-runs/{rid}/exports/json?reviewed_only=true').json()
    assert reviewed['quantity_takeoffs']==[]
    assert reviewed['summary']['review_scope']=='Accepted or edited material and QA items only'
    workbook=load_workbook(io.BytesIO(client.get(f'/api/analysis-runs/{rid}/exports/xlsx').content))
    assert workbook['Quantity Takeoffs']['A5'].value=='ANCHOR' and workbook['Quantity Takeoffs']['B5'].value==2


def test_runner_persists_visual_evidence_and_page_preview_without_key_exposure(client,project,tmp_path):
    image=Image.new('RGB',(320,180),'white');source=io.BytesIO();image.save(source,format='PNG')
    document=upload(client,project['id'],'drawing.png',source.getvalue())
    runner=client.app.state.runner
    live=replace(runner.s,provider='deepseek',live_enabled=True,prices_confirmed=True,
                 api_key='synthetic-not-real',vision_enabled=True,
                 input_rate=Decimal('4.40'),output_rate=Decimal('13.20'))
    runner.s=live;runner.gateway.s=live
    run=runner.create(project['id']);client.app.state.db.execute("UPDATE runs SET status='RUNNING' WHERE id=?",(run['id'],))
    run=runner.get(run['id'])
    runner.parse_one(run,document['document_id'])
    visual={'page_type':'DRAWING','sheet_id':'A-1','important_visible_text':'DOOR D-1',
            'observations':['One visible door tag.'],'explicit_quantity_texts':[],
            'scale_text':None,'limitations':['Synthetic test image.'],'needs_review':True}
    runner.gateway.vision=lambda *args,**kwargs:ModelResult(visual,'REQ-safe',False,'deepseek')
    runner.process_visual_tasks(run)
    payloads=[json.loads(row['payload']) for row in client.app.state.db.all(
        'SELECT payload FROM evidence WHERE run_id=?',(run['id'],))]
    evidence=next(item for item in payloads if item['extraction_method']=='VISION')
    assert evidence['locator']['sheet']=='A-1' and evidence['confidence'] is None
    assert evidence['content_basis']=='MODEL_VISION_OUTPUT'
    assert evidence['image_crop_uri'].endswith('/pages/1/image') and 'synthetic-not-real' not in json.dumps(evidence)
    preview=client.get(evidence['image_crop_uri'])
    assert preview.status_code==200 and preview.headers['content-type']=='image/png' and preview.content.startswith(b'\x89PNG')
    # Simulate a process dying after evidence commit but before the task-state update.
    row=client.app.state.db.one('SELECT summary FROM document_results WHERE run_id=? AND document_id=?',
                                (run['id'],document['document_id']))
    summary=json.loads(row['summary']);summary['visual_tasks'][0]['status']='PENDING'
    client.app.state.db.execute('UPDATE document_results SET summary=? WHERE run_id=? AND document_id=?',
                                (json.dumps(summary),run['id'],document['document_id']))
    runner.gateway.vision=lambda *args,**kwargs:(_ for _ in ()).throw(AssertionError('must not resend'))
    runner.process_visual_tasks(run)
    repaired=json.loads(client.app.state.db.one(
        'SELECT summary FROM document_results WHERE run_id=? AND document_id=?',
        (run['id'],document['document_id']))['summary'])
    assert repaired['visual_tasks'][0]['status']=='VISION_EXTRACTED'


def test_run_page_image_returns_a_validated_source_coordinate_crop(client,project,tmp_path):
    from reportlab.pdfgen import canvas
    path=tmp_path/'crop.pdf';drawing=canvas.Canvas(str(path),pagesize=(600,800))
    drawing.drawString(100,650,'EQUIPMENT SCHEDULE');drawing.rect(100,300,300,300);drawing.save()
    document=upload(client,project['id'],path.name,path.read_bytes())
    run=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()
    base=f'/api/analysis-runs/{run["id"]}/documents/{document["document_id"]}/pages/1/image'

    response=client.get(base+'?x0=100&y0=100&x1=400&y1=300')

    assert response.status_code==200 and response.content.startswith(b'\x89PNG')
    preview=Image.open(io.BytesIO(response.content))
    assert preview.width / preview.height == pytest.approx(300 / 200,rel=0.01)
    assert client.get(base+'?x0=100&y0=100').status_code==422

def test_parser_timeout_keeps_completed_page_progress(client,project,monkeypatch,tmp_path):
    document=upload(client,project['id'],'large.pdf',b'not-read-by-fake-worker')
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id']
    runner=client.app.state.runner;run=runner.get(rid)
    checkpoint={'status':'PARTIAL','fragments':[{
        'text':'Provide cement.','locator':{'page_number':1,'sheet':None,'section':None,'paragraph':None,
        'bbox':[1,2,3,4],'coordinate_system':'pdf-points-top-left','text_line_start':None,
        'text_line_end':None,'native_element_id':None},'method':'TEXT_LAYER',
        'internal_revision_date':None,'revision_label':None,'text_map':[]}],
        'pages':[{'page':1,'status':'TEXT_EXTRACTED'}],'page_count':3,
        'warnings':[],'parser_version':'text-anchors-3'}
    def timeout(command,**kwargs):
        Path(command[-1]).write_text(json.dumps(checkpoint),encoding='utf-8')
        raise subprocess.TimeoutExpired(command,kwargs['timeout'])
    monkeypatch.setattr('app.runner.subprocess.run',timeout)
    runner.parse_one(run,document['document_id']);runner.update_coverage(rid)
    coverage=runner.get(rid)['coverage']
    assert coverage['fragments_total']==1 and coverage['files_partial_or_failed']==1
    assert coverage['warnings']==['解析超过300秒；已保留1/3页，未完成页已标记NOT_PROCESSED。']
    summary=json.loads(runner.db.one('SELECT summary FROM document_results WHERE run_id=?',(rid,))['summary'])
    assert [page['status'] for page in summary['pages']]==['TEXT_EXTRACTED','NOT_PROCESSED','NOT_PROCESSED']

def test_parser_worker_does_not_inherit_secrets(client,project,monkeypatch):
    document=upload(client,project['id'],'plain.txt',b'hello')
    run=client.app.state.runner.create(project['id']);runner=client.app.state.runner;seen=[]
    monkeypatch.setenv('CIRP_API_KEY','synthetic-key')
    monkeypatch.setenv('ANOTHER_SECRET','synthetic-secret')
    monkeypatch.setenv('CAD_TOOL_PASSWORD','synthetic-password')
    monkeypatch.setenv('AWS_ACCESS_KEY_ID','synthetic-access-key')
    monkeypatch.setenv('UNRELATED_SETTING','kept')
    def worker(command,**kwargs):
        seen.append(kwargs['env'])
        Path(command[-1]).write_text(json.dumps({'status':'SUCCESS','fragments':[],'pages':[],
            'warnings':[],'parser_version':'text-anchors-3'}),encoding='utf-8')
    monkeypatch.setattr('app.runner.subprocess.run',worker)
    runner.parse_one(run,document['document_id'])
    assert seen[0]['UNRELATED_SETTING']=='kept'
    assert 'CIRP_API_KEY' not in seen[0] and 'ANOTHER_SECRET' not in seen[0]
    assert 'CAD_TOOL_PASSWORD' not in seen[0] and 'AWS_ACCESS_KEY_ID' not in seen[0]

def test_shutdown_wait_covers_parser_timeout(client):
    waits=[]
    class Thread:
        def join(self,timeout):waits.append(timeout)
    runner=client.app.state.runner;runner.thread=Thread();runner.close();runner.thread=None
    assert waits==[runner.s.parser_timeout+15]

def test_export_formula_injection():
    from app.exporter import as_xlsx
    geometry={'page':2,'calibration':{'ratio':48},'primitive_counts':{'lines':3},
              'paper_length_points':72.0,'paper_rectangle_area_points2':144.0,
              'calibrated_total_length':{'value':4.0,'unit':'FT'},'material_quantity':None,
              'scope_note':'audit only'}
    data={'run':{'project_id':'P','id':'R','status':'PARTIAL','coverage':{}},'notice':'=HYPERLINK("evil")',
          'cost':{},'records':[],'evidence':[],
          'analysis_support':[{'file_name':'a.pdf','document_id':'D','takeoffs':[],
                               'geometry_summaries':[geometry]}]}
    wb=load_workbook(io.BytesIO(as_xlsx(data)))
    summary=wb['Summary']
    cell=next(row[1] for row in summary.iter_rows(min_row=5) if row[0].value=='Review notice')
    assert cell.data_type=='s' and cell.value.startswith('=')
    assert 'PDF Geometry Audit' not in wb.sheetnames and 'Quantity Takeoffs' not in wb.sheetnames

def test_snapshot_does_not_expand(client,project):
    upload(client,project['id'],'a.txt',b'first')
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id']
    upload(client,project['id'],'b.txt',b'second')
    assert len(client.get(f'/api/analysis-runs/{rid}').json()['document_ids'])==1
