"""FR-CITATION-001 / FR-VERIFY-001/002 / FR-EXPORT-002. Offline fixtures only."""
from __future__ import annotations
from copy import deepcopy
from dataclasses import replace
from decimal import Decimal
import io
import json
import sqlite3
import threading
import time

import httpx
import pytest
from openpyxl import load_workbook
from PIL import Image

from app.db import Database, DomainError, BudgetError, dumps
from app.gateway import Gateway, InvalidModelOutput, ModelResult, ProviderPaused
from app.verification import (VerificationService, citation, exact_quote, anchors, digest, fields_for, summarize, statement_span)
from app.settings import Settings, ROOT
from contracts.runtime_rules import validate_schema
from .conftest import run_demo, upload


def evidence(text='The valve body shall be 316 stainless steel.'):
    e=json.loads((ROOT/'examples/evidence.json').read_text(encoding='utf-8'))[0]
    e.update(raw_text=text, file_name='Spec.txt', text_map=[], internal_revision_date='2026-08-01', revision_label='2')
    return e


def field(ev,claim='316',path='/grade'):
    return {'path':path,'claim':claim,'label':'grade','evidence_ids':[ev['evidence_id']], 'basis':'DIRECT',
            'context':{'name':'valve'},'check_key':digest([path,claim]),'status':'NEEDS_SEMANTIC','method':'LITERAL_LOCATION_ONLY',
            'citations':[],'issues':[],'request_id':None}


def result(f,e,**kw):
    row={'path':f['path'],'status':'SUPPORTED','citations':[{'evidence_id':e['evidence_id'],'quote':e['raw_text'],'role':'SUPPORT'}],'reason':'Source explicitly states this.'}
    row.update(kw)
    return {'checks':[row]}


def body(data,**kw):
    b={'id':'synthetic-verifier','usage':{'prompt_tokens':100,'completion_tokens':60},
       'choices':[{'finish_reason':'stop','message':{'content':dumps(data)}}]}
    b.update(kw);return b


@pytest.mark.parametrize('text,needle',[
 ('Revision: 2\nConcrete shall be 5000 psi. Steel is Grade 60.','5000'),
 ('设备应使用316不锈钢。支撑不得使用木材。','316'),
 ('🌏 Valve body shall be 316 SS.\nAnother line.','316'),
 ('Use 1.5 inch pipe. Do not use copper.','1.5'),
])
def test_exact_slice_and_unicode_offsets(text,needle):
    e=evidence(text);span=anchors(needle,e)[0]
    assert span['quote']==text[span['start']:span['end']]
    assert needle in span['quote']
    assert span['file_sha256']==e['file_sha256']
    assert span['role']=='CONTEXT'


@pytest.mark.parametrize('quote',['304 stainless steel','', 'The valve body shall be 316 stainless steel!'])
def test_false_quotes_rejected(quote):
    with pytest.raises(DomainError):exact_quote(evidence(),quote)


def test_duplicate_sentence_requires_unambiguous_quote():
    with pytest.raises(DomainError):exact_quote(evidence('A valve. A valve.'),'A valve.')


@pytest.mark.parametrize('start,end',[(0,999),(-1,4),(4,4),(1.0,5),(True,4)])
def test_invalid_span_rejected(start,end):
    with pytest.raises(DomainError):citation(evidence(),start,end)


def test_pdf_text_map_coordinates():
    from app.parsers import pdf_fragment
    words=[{'text':'Valve','x0':10,'top':20,'x1':35,'bottom':30},{'text':'316','x0':40,'top':20,'x1':56,'bottom':30}]
    fragment=pdf_fragment(words,2,('2026-08-01','2'))
    e=evidence(fragment.text);e['locator']=fragment.locator;e['text_map']=fragment.text_map
    q=exact_quote(e,'316')
    assert q['word_boxes']==[[40,20,56,30]] and q['locator']['bbox']==[40,20,56,30]
    assert q['locator']['page_number']==2


def test_vision_context_is_not_an_original_sentence_or_semantic_support():
    e=evidence('Visual observation: DOOR D-1 appears near the north wall.')
    e.update(extraction_method='VISION',content_basis='MODEL_VISION_OUTPUT',
             locator={**e['locator'],'page_number':3,'bbox':[0,0,612,792],
                      'coordinate_system':'pdf-cropbox-points-top-left'})
    q=anchors('DOOR D-1',e)[0]
    assert q['text_basis']=='MODEL_VISION_OUTPUT'
    assert q['role']=='CONTEXT' and q['granularity']=='PAGE_VISUAL_CONTEXT'
    assert (q['start'],q['end'],q['word_boxes'])==(0,len(e['raw_text']),[])
    f=field(e);f['citations']=[q]
    validate_schema('citation-report',{'version':'1.0','policy_version':'test','record_id':'REC-test',
        'candidate_hash':'x','evidence_fingerprint':'x','checked_at':'now','status':'PENDING','counts':{},
        'fields':[f],'coverage_note':'x','limitations':[],'source_run_id':'RUN-test',
        'processing_basis':None,'verifier_identity':{'provider':'mock','model_id':'mock','prompt_sha256':'x','endpoint_sha256':'x','policy_version':'test'}})
    with pytest.raises(DomainError):
        exact_quote(e,e['raw_text'])
    semantic=field(e)
    with pytest.raises(InvalidModelOutput):
        VerificationService.apply_model([semantic],result(semantic,e),{e['evidence_id']:e},'CALL-test')


def test_vision_context_schema_rejects_support_role():
    e=evidence('Visual observation: symbol.')
    e.update(extraction_method='VISION',content_basis='MODEL_VISION_OUTPUT')
    q=anchors('',e)[0];q['role']='SUPPORT'
    f=field(e);f['citations']=[q]
    report={'version':'1.0','policy_version':'test','record_id':'REC-test','candidate_hash':'x',
        'evidence_fingerprint':'x','checked_at':'now','status':'PENDING','counts':{},'fields':[f],
        'coverage_note':'x','limitations':[],'source_run_id':'RUN-test','processing_basis':None,
        'verifier_identity':{'provider':'mock','model_id':'mock','prompt_sha256':'x','endpoint_sha256':'x','policy_version':'test'}}
    with pytest.raises(Exception):validate_schema('citation-report',report)


def test_literal_mismatch_is_never_self_verified():
    e=evidence();f=field(e,'304')
    data=result(f,e)
    VerificationService.apply_model([f],data,{e['evidence_id']:e},'CALL-test')
    assert f['status']=='NEEDS_CONTEXT' # a provider's unsupported numeric approval is downgraded


def test_valid_semantic_result_preserves_server_quote():
    e=evidence();f=field(e)
    VerificationService.apply_model([f],result(f,e),{e['evidence_id']:e},'CALL-test')
    assert f['status']=='SUPPORTED' and f['citations'][0]['quote']==e['raw_text']
    assert f['request_id']=='CALL-test'


@pytest.mark.parametrize('change', ['missing','duplicate','foreign','invented_quote','no_support','no_contradiction','both','review_field'])
def test_bad_verifier_output_rejected_atomically(change):
    e=evidence();f=field(e);d=result(f,e);row=d['checks'][0]
    if change=='missing':d['checks']=[]
    if change=='duplicate':d['checks'].append(deepcopy(row))
    if change=='foreign':row['citations'][0]['evidence_id']='EV-foreign'
    if change=='invented_quote':row['citations'][0]['quote']='This sentence was never uploaded.'
    if change=='no_support':row['citations']=[]
    if change=='no_contradiction':row.update(status='CONTRADICTED',citations=[])
    if change=='both':row['citations'].append({**row['citations'][0],'role':'CONTRADICT'})
    if change=='review_field':row['review_status']='ACCEPTED'
    before=deepcopy(f)
    with pytest.raises(Exception):VerificationService.apply_model([f],d,{e['evidence_id']:e},None)
    assert f==before


def test_contradiction_has_two_distinct_states():
    e=evidence();f=field(e,'304')
    d=result(f,e,status='CONTRADICTED',citations=[{'evidence_id':e['evidence_id'],'quote':e['raw_text'],'role':'CONTRADICT'}])
    VerificationService.apply_model([f],d,{e['evidence_id']:e},None)
    assert f['status']=='CONTRADICTED'


@pytest.mark.parametrize('states,expected',[
 (['NEEDS_SEMANTIC'],'PENDING'),(['NON_DOCUMENT'],'NON_DOCUMENT'),(['SUPPORTED','NEEDS_CONTEXT'],'PARTIAL'),
 (['SUPPORTED','NON_DOCUMENT'],'SUPPORTED'),(['SUPPORTED','CONTRADICTED'],'CONTRADICTED'),(['SUPPORTED','UNSUPPORTED'],'UNSUPPORTED')])
def test_report_status(states,expected):
    assert summarize({'fields':[{'status':s} for s in states]})['status']==expected


@pytest.fixture
def demo(client,project):
    rid=run_demo(client,project['id']);rows=client.get(f'/api/analysis-runs/{rid}/records').json()
    return rid,rows


def test_every_output_has_report_no_fake_mock_pass(client,demo):
    rid,rows=demo
    assert rows
    for row in rows:
        assert row['verification']['fields']
        assert row['verification']['status']!='SUPPORTED'
        assert row['record']['review']['status']=='PENDING'
    assert client.app.state.db.cost(rows[0]['record']['meta']['project_id'])['calls']==0


def test_run_verification_batches_fields_from_records_with_the_same_evidence(client,project):
    rid=run_demo(client,project['id']);db=client.app.state.db
    rows=client.get(f'/api/analysis-runs/{rid}/records').json()
    by_scope={}
    for row in rows:
        pending=[field for field in row['verification']['fields'] if field['status']=='NEEDS_SEMANTIC']
        for item in pending:by_scope.setdefault(tuple(sorted(item['evidence_ids'])),set()).add(row['record']['meta']['record_id'])
    record_ids=next(sorted(ids) for ids in by_scope.values() if len(ids)>1)
    settings=replace(client.app.state.settings,provider='deepseek',cheap_model='offline-verifier')
    calls=[]
    class FakeGateway:
        def verify_claims(self,run,fields,bundle,job_id=None):
            calls.append(deepcopy(fields))
            return ModelResult({'checks':[{'path':item['path'],'status':'NEEDS_CONTEXT','citations':[],
                                           'reason':'Synthetic offline batch.'} for item in fields]},'CALL-batch')
    service=VerificationService(db,settings,FakeGateway())
    reports=service.refresh_many(record_ids,allow_model=True)
    assert reports and any(len({json.dumps(item['context'],sort_keys=True) for item in batch})>1 for batch in calls)


def test_exact_citation_endpoint_and_read_has_no_model_call(client,demo):
    rid,rows=demo
    row=next(r for r in rows if r['record']['kind']=='MATERIAL');report=row['verification']
    q=next(q for f in report['fields'] for q in f['citations'])
    url=f"/api/records/{report['record_id']}/citations/{q['citation_id']}"
    data=client.get(url).json();assert data['citation']['quote']==q['quote']
    ev=client.get(f'/api/analysis-runs/{rid}/evidence/{q["evidence_id"]}').json()['evidence']
    assert ev['raw_text'][q['start']:q['end']]==q['quote']
    assert client.get(url+'/preview').status_code==422
    assert client.get(url+'-fake').status_code==404


def test_edit_invalidation_reuses_unchanged_fields(client,demo):
    rid,rows=demo;row=next(r for r in rows if r['record']['kind']=='MATERIAL' and r['record']['candidate']['design_properties'])
    service=client.app.state.runner.verifier;id_=row['record']['meta']['record_id'];report=service.get(id_)
    # Simulate an earlier accepted verifier verdict, with real evidence anchors.
    for f in report['fields']:
        if f['basis']=='DIRECT':f.update(status='SUPPORTED',method='CHEAP_MODEL_SEMANTIC',request_id='CALL-old')
    summarize(report);assert service.save(report)
    candidate=deepcopy(row['record']['candidate'])
    if not candidate['design_properties']:
        pytest.fail('Synthetic demo material must have a property')
    candidate['design_properties'][0]['value']='99999'
    result_=client.post(f'/api/records/{id_}/review',json={'action':'EDITED','expected_version':0,'candidate':candidate,'note':'Change design value'})
    assert result_.status_code==200,result_.text
    latest=result_.json()['verification']
    assert next(f for f in latest['fields'] if f['path']=='/design_properties/0/value')['status']=='NEEDS_SEMANTIC'
    assert next(f for f in latest['fields'] if f['path']=='/name')['request_id']=='CALL-old'
    assert result_.json()['record']['review']['status']=='EDITED'
    assert len(client.get(f'/api/records/{id_}/verification-history').json())>=2


def test_stale_hash_never_displayed_as_supported(client,demo):
    rid,rows=demo;row=rows[0];id_=row['record']['meta']['record_id'];db=client.app.state.db
    r=deepcopy(row['record']);r['candidate']['subject']='Changed subject'
    # Direct DB mutation emulates stale cache or future clients; not an API approval.
    db.execute('UPDATE records SET envelope=? WHERE id=?',(dumps(r),id_))
    assert client.get(f'/api/records/{id_}/verification').json()['status']=='STALE'


def test_changed_source_invalidates_verification(client,demo):
    rid,rows=demo;db=client.app.state.db
    e=db.one('SELECT * FROM evidence WHERE run_id=?',(rid,));payload=json.loads(e['payload']);payload['raw_text']+='Changed.'
    db.execute('UPDATE evidence SET payload=? WHERE id=?',(dumps(payload),e['id']))
    affected=[r for r in rows if payload['evidence_id'] in {q['evidence_id'] for f in r['verification']['fields'] for q in f['citations']}]
    assert affected
    assert client.get(f'/api/records/{affected[0]["record"]["meta"]["record_id"]}/verification').json()['status']=='STALE'


def test_mock_semantic_button_cannot_trigger_paid_call(client,demo):
    row=demo[1][0];uri=f'/api/records/{row["record"]["meta"]["record_id"]}/verification'
    assert client.post(uri,json={'expected_version':0,'semantic':True}).status_code==409
    assert client.post(uri,json={'expected_version':99}).status_code==409
    assert client.post(uri,json={'expected_version':0}).status_code==200


def test_export_places_readable_evidence_with_each_item(client,demo):
    rid,_=demo
    data=client.get(f'/api/analysis-runs/{rid}/exports/json').json()
    assert data['export_version']=='0.2.6-readable-en-2'
    assert all(item['evidence'] for key in ('materials_and_equipment','inspections_and_tests')
               for item in data[key])
    assert 'conflicts' not in data and 'missing_information' not in data
    assert 'citation_catalog' not in data and 'verifications' not in data
    assert not any(key.endswith('_id') or key.endswith('_ids')
                   for key in data for _ in [None])
    wb=load_workbook(io.BytesIO(client.get(f'/api/analysis-runs/{rid}/exports/xlsx').content))
    assert {'Evidence','Field Verification','Citations','PDF Geometry Audit'}.isdisjoint(wb.sheetnames)
    material=wb['Materials & Equipment']
    assert material.cell(4,14).value=='Evidence Source' and material.cell(4,15).value=='Evidence Text'
    assert material.cell(5,14).value and material.cell(5,15).value
    for ws in wb:
        for row in ws:
            assert all(c.data_type!='f' for c in row)


def test_rejected_item_not_in_reviewed_only_export(client,demo):
    rid,rows=demo;row=rows[0];id_=row['record']['meta']['record_id']
    client.post(f'/api/records/{id_}/review',json={'action':'REJECTED','expected_version':0})
    data=client.get(f'/api/analysis-runs/{rid}/exports/json?reviewed_only=true').json()
    assert all(not data[key] for key in ('materials_and_equipment','inspections_and_tests','quantity_takeoffs'))
    assert id_ not in json.dumps(data,ensure_ascii=False)


def test_additive_migration_keeps_legacy_budget(tmp_path):
    path=tmp_path/'old.db'
    c=sqlite3.connect(path);c.executescript((ROOT/'migrations/001_initial.sql').read_text())
    c.execute("INSERT INTO projects VALUES('P','Legacy','today')")
    c.execute("INSERT INTO budget_accounts VALUES('P',300000000,1000000,0)");c.commit();c.close()
    db=Database(path)
    assert db.cost('P')['spent_cny']=='1.000000'
    assert db.one('SELECT version FROM schema_migrations WHERE version=2')['version']==2
    assert db.one('SELECT version FROM schema_migrations WHERE version=3')['version']==3
    assert db.one('SELECT version FROM schema_migrations WHERE version=4')['version']==4
    assert db.one('SELECT version FROM schema_migrations WHERE version=5')['version']==5
    assert any(row['name']=='ix_uploads_document' for row in db.all("PRAGMA index_list('uploads')"))
    assert db.one("SELECT name FROM sqlite_master WHERE type='table' AND name='call_reconciliation_events'")
    assert db.one("SELECT name FROM sqlite_master WHERE type='table' AND name='run_metrics'")
    assert db.one("SELECT name FROM sqlite_master WHERE type='table' AND name='budget_limit_events'")
    assert Database(path).cost('P')['spent_cny']=='1.000000'


@pytest.fixture
def live_context(client,project,tmp_path):
    upload(client,project['id'],'spec.txt',b'The valve body shall be 316 stainless steel.')
    run=client.app.state.runner.create(project['id']);db=client.app.state.db
    db.execute("UPDATE runs SET status='RUNNING',provider='deepseek' WHERE id=?",(run['id'],));run=client.app.state.runner.get(run['id'])
    ev=evidence();ev.update(project_id=project['id'],input_snapshot_id=run['snapshot_id'])
    s=Settings(tmp_path,provider='deepseek',live_enabled=True,prices_confirmed=True,api_key='offline-not-real',input_rate=Decimal('1'),output_rate=Decimal('2'),start_worker=False)
    return db,run,ev,s


def test_verifier_non_thinking_budget_and_cache(live_context):
    db,run,e,s=live_context;f=field(e);calls=[]
    def handler(r):calls.append(json.loads(r.content));return httpx.Response(200,json=body(result(f,e)))
    g=Gateway(s,db,httpx.Client(transport=httpx.MockTransport(handler)))
    assert not g.verify_claims(run,[f],{e['evidence_id']:e}).cached
    assert g.verify_claims(run,[f],{e['evidence_id']:e}).cached
    assert len(calls)==1 and calls[0]['thinking']=={'type':'disabled'}
    assert calls[0]['model']==s.cheap_model and calls[0]['max_tokens']==1400
    assert db.cost(run['project_id'])['spent_cny']=='0.000220'
    performance=db.performance(run['id'])
    assert performance['model.request_wait']['samples']==1
    assert performance['model.response']['samples']==1


def test_verifier_crash_after_atomic_commit_recovers_without_second_http(live_context,monkeypatch):
    db,run,e,s=live_context;f=field(e);requests=[]
    g=Gateway(s,db,httpx.Client(transport=httpx.MockTransport(
        lambda request:(requests.append(request),httpx.Response(200,json=body(result(f,e))))[1])))
    committed=db.finalize_model_call
    def commit_then_crash(*args,**kwargs):
        committed(*args,**kwargs)
        raise SystemExit('synthetic crash after verification commit')
    monkeypatch.setattr(db,'finalize_model_call',commit_then_crash)
    with pytest.raises(SystemExit,match='verification commit'):
        g.verify_claims(run,[f],{e['evidence_id']:e})
    monkeypatch.setattr(db,'finalize_model_call',committed)

    recovered=g.verify_claims(run,[f],{e['evidence_id']:e})
    assert recovered.cached and recovered.request_id and recovered.data==result(f,e)
    assert len(requests)==1 and db.cost(run['project_id'])['calls']==1


def test_gemini_verifier_uses_minimal_and_accounts_total_output(live_context):
    db,run,e,s=live_context;run={**run,'provider':'gemini'};f=field(e);calls=[]
    s=replace(s,provider='gemini',api_base_url='https://generativelanguage.googleapis.com/v1beta/openai',
              cheap_model='gemini-3.6-flash')
    response=body(result(f,e));response['usage']['total_tokens']=175
    def handler(r):calls.append(json.loads(r.content));return httpx.Response(200,json=response)
    g=Gateway(s,db,httpx.Client(transport=httpx.MockTransport(handler)))
    value=g.verify_claims(run,[f],{e['evidence_id']:e})
    assert value.mode=='gemini' and calls[0]['reasoning_effort']=='minimal' and 'thinking' not in calls[0]
    assert db.cost(run['project_id'])['output_tokens']==75


@pytest.mark.parametrize('case',['budget','disabled','oversized'])
def test_no_http_when_not_authorized_or_no_capacity(live_context,case):
    db,run,e,s=live_context;f=field(e)
    if case=='budget':db.execute('UPDATE budget_accounts SET spent_units=300000000 WHERE project_id=?',(run['project_id'],))
    if case=='disabled':s=replace(s,live_enabled=False)
    if case=='oversized':e['raw_text']='a'*12000
    g=Gateway(s,db,httpx.Client(transport=httpx.MockTransport(lambda r:pytest.fail('No HTTP permitted'))))
    with pytest.raises((BudgetError,ProviderPaused,InvalidModelOutput)):g.verify_claims(run,[f],{e['evidence_id']:e})
    assert db.cost(run['project_id'])['calls']==0


@pytest.mark.parametrize('case',['timeout','usage','truncated','bad_quote','thinking'])
def test_verifier_failure_never_repeats_silently(live_context,case):
    db,run,e,s=live_context;f=field(e);requests=[];data=result(f,e)
    def handler(r):
        requests.append(r)
        if case=='timeout':raise httpx.ReadTimeout('synthetic timeout')
        b=body(data)
        if case=='usage':b.pop('usage')
        if case=='truncated':b['choices'][0]['finish_reason']='length'
        if case=='bad_quote':b['choices'][0]['message']['content']=dumps(result(f,e,citations=[{'evidence_id':e['evidence_id'],'quote':'Fake.','role':'SUPPORT'}]))
        if case=='thinking':b['choices'][0]['message']['reasoning_content']='not permitted'
        return httpx.Response(200,json=b)
    g=Gateway(s,db,httpx.Client(transport=httpx.MockTransport(handler)))
    with pytest.raises((ProviderPaused,InvalidModelOutput)):g.verify_claims(run,[f],{e['evidence_id']:e})
    assert len(requests)==1
    if case in ('timeout','usage'):
        with pytest.raises(ProviderPaused):g.verify_claims(run,[f],{e['evidence_id']:e})
        assert len(requests)==1 and db.cost(run['project_id'])['unknown_calls']==1
    else:
        with pytest.raises((ProviderPaused,InvalidModelOutput)):
            g.verify_claims(run,[f],{e['evidence_id']:e})
        row=db.one('SELECT state,response,error FROM model_calls WHERE run_id=?',(run['id'],))
        assert len(requests)==1 and Decimal(db.cost(run['project_id'])['spent_cny'])>0
        assert row['state']=='SETTLED_ERROR' and row['response'] is None and row['error']


def test_verifier_http_status_uses_shared_safe_diagnostic(live_context):
    db,run,e,s=live_context;f=field(e);requests=[]
    def handler(request):
        requests.append(request)
        return httpx.Response(429,headers={'Retry-After':'7','x-request-id':'verify-request-1'},
                              json={'error':{'status':'RESOURCE_EXHAUSTED','message':'private upstream text'}})
    g=Gateway(s,db,httpx.Client(transport=httpx.MockTransport(handler)))
    with pytest.raises(ProviderPaused,match='HTTP 429') as caught:
        g.verify_claims(run,[f],{e['evidence_id']:e})
    assert 'Retry-After=7' in str(caught.value) and len(requests)==1
    row=db.one('SELECT state,error,provider_request_id FROM model_calls WHERE run_id=?',(run['id'],))
    diagnostic=json.loads(row['error'])
    assert diagnostic=={'kind':'HTTP_STATUS','status':429,'class':'RATE_LIMIT_OR_QUOTA',
                        'provider_code':'RESOURCE_EXHAUSTED','retry_after':'7',
                        'provider_request_id':'verify-request-1'}
    assert row['state']=='UNKNOWN' and row['provider_request_id']=='verify-request-1'
    assert 'private upstream text' not in row['error']


def test_unsupported_source_does_not_create_reviewer_item(client,project):
    upload(client,project['id'],'drawing.dwg',b'Unsupported synthetic DWG')
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id'];client.app.state.runner.process(rid)
    assert client.get(f'/api/analysis-runs/{rid}/records').json()==[]
    run=client.get(f'/api/analysis-runs/{rid}').json()
    assert run['coverage']['files_partial_or_failed']==1


def test_citation_xss_preserved_as_text_and_not_generated_html():
    e=evidence('<script>alert(1)</script>\nValve body is 316.')
    q=exact_quote(e,'<script>alert(1)</script>')
    assert q['quote']=='<script>alert(1)</script>'
    source=(ROOT/'web/app.js').read_text(encoding='utf-8')
    assert 'innerHTML' not in source and 'document.createTextNode' in source


def prepare_semantic_job(client,demo):
    rid,rows=demo;db=client.app.state.db
    row=next(x for x in rows if x['record']['kind']=='MATERIAL')
    db.execute("UPDATE runs SET provider='deepseek' WHERE id=?",(rid,))
    s=replace(client.app.state.settings,provider='deepseek',live_enabled=True,prices_confirmed=True,api_key='synthetic-key',input_rate=Decimal('1'),output_rate=Decimal('2'))
    service=VerificationService(db,s,None)
    def handler(request):
        payload=json.loads(request.content);content=json.loads(payload['messages'][1]['content'])
        checks=[]
        for f in content['fields']:
            e=next(e for e in content['evidence'] if e['evidence_id'] in f['evidence_ids'])
            checks.append({'path':f['path'],'status':'NEEDS_CONTEXT','citations':[{'evidence_id':e['evidence_id'],'quote':e['text'],'role':'CONTEXT'}], 'reason':'Synthetic demo protocol is not design evidence.'})
        return httpx.Response(200,json=body({'checks':checks}))
    gateway=Gateway(s,db,httpx.Client(transport=httpx.MockTransport(handler)))
    service.gateway=gateway
    return db,service,row,rid


def test_persistent_job_can_verify_finished_run_without_changing_review(client,demo):
    db,service,row,rid=prepare_semantic_job(client,demo);record_id=row['record']['meta']['record_id']
    job=service.enqueue(record_id,0);assert service.enqueue(record_id,0)['id']==job['id']
    service.process_job(job['id'])
    assert db.one('SELECT state FROM verification_jobs WHERE id=?',(job['id'],))['state']=='DONE'
    assert db.cost(row['record']['meta']['project_id'])['calls']>=1
    current=json.loads(db.one('SELECT envelope FROM records WHERE id=?',(record_id,))['envelope'])
    assert current['review']==row['record']['review']
    assert db.one('SELECT status FROM runs WHERE id=?',(rid,))['status']=='PARTIAL'


def test_reconciled_verification_job_requires_new_explicit_job_generation(client,demo):
    db,service,row,rid=prepare_semantic_job(client,demo);requests=[]
    original_client=service.gateway.client

    def handler(request):
        requests.append(request)
        if len(requests)==1:
            raise httpx.ReadTimeout('synthetic verification uncertainty')
        payload=json.loads(request.content);content=json.loads(payload['messages'][1]['content'])
        checks=[]
        for item in content['fields']:
            evidence=next(value for value in content['evidence'] if value['evidence_id'] in item['evidence_ids'])
            checks.append({'path':item['path'],'status':'NEEDS_CONTEXT',
                           'citations':[{'evidence_id':evidence['evidence_id'],'quote':evidence['text'],'role':'CONTEXT'}],
                           'reason':'Synthetic demo protocol is not design evidence.'})
        return httpx.Response(200,json=body({'checks':checks}))

    service.gateway.client=httpx.Client(transport=httpx.MockTransport(handler))
    record_id=row['record']['meta']['record_id']
    first_job=service.enqueue(record_id,0);service.process_job(first_job['id'])
    assert db.one('SELECT state FROM verification_jobs WHERE id=?',(first_job['id'],))['state']=='PAUSED_PROVIDER'
    call=db.one("SELECT id,state,task_key FROM model_calls WHERE run_id=?",(rid,))
    assert call['state']=='UNKNOWN' and call['task_key'].startswith('verify:')
    db.reconcile_call(call['id'],'NOT_BILLED',Decimal('0'),'offline provider ledger check')

    second_job=service.enqueue(record_id,0)
    assert second_job['id']!=first_job['id']
    service.process_job(second_job['id'])
    assert db.one('SELECT state FROM verification_jobs WHERE id=?',(second_job['id'],))['state']=='DONE'
    calls=db.all("SELECT id,state,task_key FROM model_calls WHERE run_id=? ORDER BY created_at,id",(rid,))
    assert len(requests)==2 and len(calls)==2
    assert calls[0]['id']==call['id'] and calls[0]['state']=='RECONCILED_ZERO'
    assert calls[1]['task_key']==call['task_key']+':manual-requeue:1' and calls[1]['state']=='SETTLED'
    events=db.reconciliation_events(row['record']['meta']['project_id'])
    recovery=next(event for event in events if event['payload'].get('event_type')=='RECOVERY_GENERATION_AUTHORIZED')
    assert recovery['payload']['trigger']=='VERIFICATION_ENQUEUE'
    assert recovery['payload']['context_id']==second_job['id']
    service.gateway.client.close();original_client.close()


def test_verification_family_stops_after_three_explicit_reconciled_jobs(client,demo):
    db,service,row,rid=prepare_semantic_job(client,demo);requests=[]
    original_client=service.gateway.client

    def handler(request):
        requests.append(request)
        raise httpx.ReadTimeout('synthetic repeated verification uncertainty')

    service.gateway.client=httpx.Client(transport=httpx.MockTransport(handler))
    record_id=row['record']['meta']['record_id']
    for _ in range(3):
        job=service.enqueue(record_id,0);service.process_job(job['id'])
        assert db.one('SELECT state FROM verification_jobs WHERE id=?',(job['id'],))['state']=='PAUSED_PROVIDER'
        unknown=db.one("SELECT id FROM model_calls WHERE run_id=? AND actual_units IS NULL",(rid,))
        db.reconcile_call(unknown['id'],'NOT_BILLED',Decimal('0'),'offline provider ledger check')
    blocked=service.enqueue(record_id,0);service.process_job(blocked['id'])
    blocked_row=db.one('SELECT state,message FROM verification_jobs WHERE id=?',(blocked['id'],))
    assert blocked_row['state']=='PAUSED_PROVIDER' and '三次累计调用上限' in blocked_row['message']
    calls=db.all("SELECT task_key,state FROM model_calls WHERE run_id=? ORDER BY created_at,id",(rid,))
    family=calls[0]['task_key']
    assert len(requests)==3 and [call['task_key'] for call in calls]==[
        family,family+':manual-requeue:1',family+':manual-requeue:2']
    assert all(call['state']=='RECONCILED_ZERO' for call in calls)
    service.gateway.client.close();original_client.close()


def test_new_analysis_blocked_while_verification_queued(client,demo):
    db,service,row,rid=prepare_semantic_job(client,demo)
    service.enqueue(row['record']['meta']['record_id'],0)
    response=client.post(f'/api/projects/{row["record"]["meta"]["project_id"]}/analysis-runs')
    assert response.status_code==409


def test_stale_job_does_not_make_model_call(client,demo):
    db,service,row,rid=prepare_semantic_job(client,demo);record_id=row['record']['meta']['record_id']
    job=service.enqueue(record_id,0)
    db.execute('UPDATE records SET review_version=review_version+1 WHERE id=?',(record_id,))
    service.process_job(job['id'])
    assert db.one('SELECT state FROM verification_jobs WHERE id=?',(job['id'],))['state']=='STALE'
    assert db.cost(row['record']['meta']['project_id'])['calls']==0


def test_expired_run_no_semantic_job(client,demo):
    db,service,row,rid=prepare_semantic_job(client,demo)
    db.execute('UPDATE runs SET deadline_epoch=? WHERE id=?',(time.time()-1,rid))
    with pytest.raises(DomainError):service.enqueue(row['record']['meta']['record_id'],0)


def test_job_budget_exhaustion_retains_citations(client,demo):
    db,service,row,rid=prepare_semantic_job(client,demo);record_id=row['record']['meta']['record_id']
    db.execute('UPDATE budget_accounts SET spent_units=300000000 WHERE project_id=?',(row['record']['meta']['project_id'],))
    job=service.enqueue(record_id,0);service.process_job(job['id'])
    assert db.one('SELECT state FROM verification_jobs WHERE id=?',(job['id'],))['state']=='PAUSED_BUDGET'
    assert service.get(record_id)['status']=='PENDING'
    assert any(f['citations'] for f in service.get(record_id)['fields'])
    assert db.cost(row['record']['meta']['project_id'])['calls']==0


def test_close_during_verification_throttle_prevents_reserve_and_http(client,demo):
    db,service,row,rid=prepare_semantic_job(client,demo);record_id=row['record']['meta']['record_id']
    service.s=replace(service.s,min_request_interval_seconds=.3)
    requests=[]
    service.gateway.s=service.s;service.gateway._last_request_started=time.monotonic()
    service.gateway.client=httpx.Client(transport=httpx.MockTransport(lambda request: requests.append(request)))
    job=service.enqueue(record_id,0)
    worker=threading.Thread(target=service.process_job,args=(job['id'],));worker.start()
    deadline=time.time()+2
    while time.time()<deadline and db.one('SELECT state FROM verification_jobs WHERE id=?',(job['id'],))['state']!='RUNNING':
        time.sleep(.005)
    client.app.state.runner.close();worker.join(timeout=2)
    assert not worker.is_alive() and requests==[]
    assert db.one('SELECT state FROM verification_jobs WHERE id=?',(job['id'],))['state']=='INTERRUPTED'
    assert db.one('SELECT COUNT(*) AS n FROM model_calls WHERE run_id=?',(rid,))['n']==0


def test_pdf_preview_matches_source_and_blocks_modified_file(client,project,tmp_path):
    from reportlab.pdfgen import canvas
    from app.parsers import parse_file
    from app.assemble import envelopes
    from app.gateway import mock_extract
    p=tmp_path/'source.pdf';c=canvas.Canvas(str(p));c.drawString(72,700,'DEMO_MATERIAL|PAD-01|Concrete|-|strength=5000 psi');c.save()
    upload(client,project['id'],'source.pdf',p.read_bytes())
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id'];client.app.state.runner.process(rid)
    rows=client.get(f'/api/analysis-runs/{rid}/records?kind=MATERIAL').json()
    assert rows,client.get(f'/api/analysis-runs/{rid}').json()
    row=rows[0];q=next(q for f in row['verification']['fields'] for q in f['citations'])
    url=f'/api/records/{row["record"]["meta"]["record_id"]}/citations/{q["citation_id"]}/preview'
    response=client.get(url);assert response.status_code==200,response.text[:300] if response.status_code!=200 else ''
    assert response.content.startswith(b'\x89PNG')
    doc=client.app.state.db.one('SELECT * FROM documents WHERE id=?',(q['document_id'],))
    client.app.state.uploads.object_path(doc).write_bytes(b'changed-file')
    assert client.get(url).status_code==409


def test_pdf_citation_preview_draws_cropbox_local_locator_in_visible_position(client,project,tmp_path):
    """A non-zero PDF CropBox must not shift a local evidence rectangle twice."""
    from reportlab.pdfgen import canvas
    path=tmp_path/'cropped-source.pdf';drawing=canvas.Canvas(str(path),pagesize=(600,800))
    drawing.setCropBox([100,100,500,600])
    drawing.drawString(150,400,'DEMO_MATERIAL|PAD-01|Concrete|-|strength=5000 psi')
    drawing.save()
    upload(client,project['id'],'cropped-source.pdf',path.read_bytes())
    rid=client.post(f'/api/projects/{project["id"]}/analysis-runs').json()['id']
    client.app.state.runner.process(rid)
    row=client.get(f'/api/analysis-runs/{rid}/records?kind=MATERIAL').json()[0]
    citation=next(q for field in row['verification']['fields'] for q in field['citations'])
    assert citation['locator']['coordinate_system']=='pdf-cropbox-points-top-left'
    box=citation['word_boxes'][0]
    response=client.get(f'/api/records/{row["record"]["meta"]["record_id"]}/citations/{citation["citation_id"]}/preview')
    assert response.status_code==200
    preview=Image.open(io.BytesIO(response.content)).convert('RGB')
    scale=preview.width / 400
    # The word begins 50 CropBox points from the left, so an erroneous second
    # CropBox subtraction would put it off-canvas. The correct location is tinted.
    x=round((box[0]+1)*scale);y=round((box[1]+1)*scale)
    nearby=[preview.getpixel((px,py)) for px in range(max(0,x-2),min(preview.width,x+3))
            for py in range(max(0,y-2),min(preview.height,y+3))]
    assert any(red>150 and green>80 and blue<180 for red,green,blue in nearby)


def test_inferred_items_cite_trigger_not_fabricated_design(client,demo):
    rid,rows=demo;row=next(r for r in rows if r['record']['kind']=='MATERIAL')
    c=deepcopy(row['record']['candidate']);c['requirement_status']='INFERRED_TO_VERIFY';c['inference_rule_id']='company-method-unapproved'
    record={**row['record'],'candidate':c}
    assert all(f['basis'] in ('INFERRED','CLASSIFICATION') for f in fields_for(record))


def test_remote_verification_endpoints_still_require_auth(tmp_path):
    from fastapi.testclient import TestClient
    from app.main import create_app
    s=Settings(tmp_path,remote_enabled=True,preview_password='synthetic-password-not-secret-012345',start_worker=False)
    with TestClient(create_app(s)) as c:
        assert c.get('/api/records/any/verification').status_code==401
        assert c.get('/api/records/any/citations/any/preview').status_code==401
        assert c.post('/api/records/any/verification',json={'expected_version':0}).status_code==401


def test_supported_cache_invalidates_on_model_change(client,demo):
    row=demo[1][0];service=client.app.state.runner.verifier;rid=row['record']['meta']['record_id']
    assert service.get(rid)['status']!='STALE'
    changed=VerificationService(service.db,replace(service.s,cheap_model='different-model'),service.gateway)
    assert changed.get(rid)['status']=='STALE'


def test_report_saves_verifier_prompt_identity(client,demo):
    report=demo[1][0]['verification']
    assert report['verifier_identity']['provider']=='mock'
    assert len(report['verifier_identity']['prompt_sha256'])==64
    assert 'api_key' not in dumps(report)


def test_model_cache_keeps_original_billing_reference(live_context):
    db,run,e,s=live_context;f=field(e)
    with httpx.Client(transport=httpx.MockTransport(lambda r:httpx.Response(200,json=body(result(f,e))))) as client:
        g=Gateway(s,db,client);first=g.verify_claims(run,[f],{e['evidence_id']:e});second=g.verify_claims(run,[f],{e['evidence_id']:e})
    assert first.request_id==second.request_id and first.request_id
    assert second.cached and db.cost(run['project_id'])['calls']==1


def test_support_notes_never_implicitly_verified(client,demo):
    row=next(r for r in demo[1] if r['record']['kind']=='MATERIAL')
    fields=fields_for(row['record'])
    note=next(f for f in fields if f['path']=='/support_note')
    assert note['basis']=='INFERRED'
