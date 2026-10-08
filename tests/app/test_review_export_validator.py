import json

import pytest

from openpyxl import load_workbook

from app import exporter
from scripts import validate_readable_export as validator
from .test_review_export_integrity import QUOTE, data, material


def _export(tmp_path):
    record=material('PANEL_A1',quantity_review='VERIFIED')
    record['candidate']['design_properties']=[{'name':'model_number','value':'AB_123','unit':None}]
    payload=data(record)
    payload['verifications']={'synthetic-record':{'fields':[{'path':'/name','status':'SUPPORTED','citations':[
        {'quote':QUOTE,'file_name':'SYNTHETIC_PANEL_A1.txt','locator':{}}]}]}}
    json_path=tmp_path/'review.json';xlsx_path=tmp_path/'review.xlsx'
    json_path.write_bytes(exporter.as_json(payload));xlsx_path.write_bytes(exporter.as_xlsx(payload))
    return json_path,xlsx_path


def test_validator_accepts_exact_source_quote_and_business_identifiers(tmp_path):
    report,errors=validator.validate(*_export(tmp_path))
    assert report['status']=='PASS' and errors==[],errors


def test_validator_still_rejects_injected_formula(tmp_path):
    json_path,xlsx_path=_export(tmp_path)
    workbook=load_workbook(xlsx_path);workbook['Materials & Equipment'].cell(5,1).value='=1+1';workbook.save(xlsx_path)
    report,errors=validator.validate(json_path,xlsx_path)
    assert report['status']=='FAIL' and 'Workbook contains 1 formulas.' in errors


def test_validator_preserves_visual_observation_language(tmp_path):
    json_path,xlsx_path=_export(tmp_path)
    payload=json.loads(json_path.read_text(encoding='utf-8'))
    payload['materials_and_equipment'][0]['evidence'][0]['source_type']='Visual model observation'
    json_path.write_text(json.dumps(payload,ensure_ascii=False),encoding='utf-8')
    report,errors=validator.validate(json_path,xlsx_path)
    assert report['status']=='PASS' and errors==[],errors


def test_validator_preserves_private_glyph_in_source_but_not_application_label(tmp_path):
    json_path,xlsx_path=_export(tmp_path)
    payload=json.loads(json_path.read_text(encoding='utf-8'))
    payload['materials_and_equipment'][0]['evidence'][0]['text']+='\ue123'
    json_path.write_text(json.dumps(payload,ensure_ascii=False),encoding='utf-8')
    workbook=load_workbook(xlsx_path)
    sheet=workbook['Materials & Equipment']
    column=next(cell.column for cell in sheet[4] if cell.value=='Evidence Text (Excerpt)')
    sheet.cell(5,column).value+='\ue123';workbook.save(xlsx_path)
    _,errors=validator.validate(json_path,xlsx_path)
    assert errors==[],errors
    sheet.cell(4,1).value+='\ue123';workbook.save(xlsx_path)
    _,errors=validator.validate(json_path,xlsx_path)
    assert 'Workbook contains 1 private-use glyphs.' in errors


def test_validator_still_rejects_untranslated_application_enum(tmp_path):
    json_path,xlsx_path=_export(tmp_path)
    payload=json.loads(json_path.read_text(encoding='utf-8'))
    payload['materials_and_equipment'][0]['quantity_basis']='DESIGN_NET'
    json_path.write_text(json.dumps(payload),encoding='utf-8')
    _,errors=validator.validate(json_path,xlsx_path)
    assert 'JSON contains 1 untranslated application enum values.' in errors


@pytest.mark.parametrize('requirement',['Hydrostatic test using approved product data',
                                      'Leakage test at no additional cost'])
def test_validator_does_not_reject_real_tests_for_modifiers(tmp_path,requirement):
    record={'kind':'INSPECTION','meta':{},'review':{'status':'PENDING'},
            'candidate':{'qa_type':'FIELD_TEST','activity':'test','requirement':requirement,'evidence_ids':[]}}
    payload=data(record)
    json_path=tmp_path/'review.json';xlsx_path=tmp_path/'review.xlsx'
    json_path.write_bytes(exporter.as_json(payload));xlsx_path.write_bytes(exporter.as_xlsx(payload))
    report,errors=validator.validate(json_path,xlsx_path)
    assert report['status']=='PASS' and errors==[],errors


def test_validator_still_rejects_document_as_primary_qa_object(tmp_path):
    json_path,xlsx_path=_export(tmp_path)
    payload=json.loads(json_path.read_text(encoding='utf-8'))
    payload['inspections_and_tests']=[{'test_or_inspection_name':'Shop drawings for pump','activity':'submit'}]
    json_path.write_text(json.dumps(payload),encoding='utf-8')
    _,errors=validator.validate(json_path,xlsx_path)
    assert any('administrative or drawing submittals' in error for error in errors)
