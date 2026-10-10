"""Synthetic audit regressions: saved facts and independent human decisions survive export."""
import io
import json
from copy import deepcopy

import pytest
from openpyxl import load_workbook

from app import exporter
from .conftest import run_demo


QUOTE = 'Do not connect to PANEL_A1. Model AB_123 is prohibited.\n混凝土强度不得低于 4000 psi。'


def material(name='Isolation valve', *, status='ACCEPTED', quantity_review='PENDING'):
    return {'kind': 'MATERIAL', 'meta': {'record_id': 'synthetic-record'},
            'review': {'status': status}, 'quantity_review': quantity_review,
            'candidate': {'name': name, 'material_kind': 'PERMANENT',
                          'quantity': {'value': 7, 'unit': 'EA', 'basis': 'DESIGN_NET'},
                          'design_properties': [], 'options': [], 'entity_ids': [],
                          'location': None, 'condition': None, 'csi_sections': [],
                          'requirement_status': 'CONFIRMED', 'evidence_ids': ['synthetic-evidence']}}


def data(record, *, reviewed_only=False):
    return {'records': [record], 'reviewed_only': reviewed_only,
            'evidence': [{'evidence_id': 'synthetic-evidence', 'raw_text': QUOTE,
                          'file_name': 'SYNTHETIC_PANEL_A1.txt', 'locator': {}}],
            'verifications': {}, 'analysis_support': [], 'run': {'status': 'PARTIAL'}}


@pytest.mark.parametrize('source_type', ['Source text', 'OCR text', 'Visual model observation'])
def test_raw_evidence_preserves_all_source_characters(source_type):
    assert exporter._evidence_text(source_type, QUOTE) == QUOTE


def test_json_and_excel_keep_citation_words_and_business_identifiers():
    record = material('PANEL_A1', quantity_review='VERIFIED')
    record['candidate']['design_properties'] = [{'name': 'model_number', 'value': 'AB_123', 'unit': None}]
    payload = data(record)
    payload['verifications'] = {'synthetic-record': {'fields': [
        {'path': '/name', 'status': 'SUPPORTED', 'citations': [
            {'quote': QUOTE, 'file_name': 'SYNTHETIC_PANEL_A1.txt', 'locator': {}}]}]}}
    item = json.loads(exporter.as_json(payload))['materials_and_equipment'][0]
    assert item['evidence'][0]['text'] == QUOTE
    assert item['material_or_equipment_name'] == 'PANEL_A1'
    assert 'AB_123' in item['design_properties']
    workbook = load_workbook(io.BytesIO(exporter.as_xlsx(payload)))
    values = [cell.value for row in workbook['Materials & Equipment'] for cell in row
              if isinstance(cell.value, str)]
    assert any(QUOTE in value for value in values)
    assert 'PANEL_A1' in values


@pytest.mark.parametrize(('name','quantity'),[
    ('2 inch PVC valve',2),('Single phase transformer',1),('Two pole circuit breaker',2),
])
def test_reviewed_export_preserves_dimension_phase_and_pole_in_complete_name(name,quantity):
    record=material(name,quantity_review='VERIFIED')
    record['candidate']['quantity']['value']=quantity
    payload=data(record,reviewed_only=True)
    exported=json.loads(exporter.as_json(payload))['materials_and_equipment'][0]
    assert exported['material_or_equipment_name']==name
    workbook=load_workbook(io.BytesIO(exporter.as_xlsx(payload)))
    values=[cell.value for row in workbook['Materials & Equipment'] for cell in row]
    assert name in values


@pytest.mark.parametrize('separator',['\n','\r\n','. ','; '])
def test_fallback_excerpt_keeps_first_character_after_real_boundary(separator):
    phrase='No asbestos pipe shall be used.'
    evidence={'content_basis':'SOURCE_TEXT','raw_text':'Header'+separator+phrase}
    record={'candidate':{'name':'No asbestos pipe'}}
    assert exporter._focused_source_text(evidence,record).startswith(phrase)


@pytest.mark.parametrize('name', ['model_number', 'Number of poles', 'quantity'])
def test_export_never_infers_quantity_from_properties(name):
    record = material()
    record['candidate']['quantity'] = None
    record['candidate']['design_properties'] = [{'name': name, 'value': '12345', 'unit': None}]
    item = json.loads(exporter.as_json(data(record)))['materials_and_equipment'][0]
    assert item['quantity'] is None
    assert '12345' in item['design_properties']


@pytest.mark.parametrize('status', ['PENDING', 'ACCEPTED', 'EDITED'])
@pytest.mark.parametrize('name', ['Manual Transfer Switch', 'Manual isolation valve',
                                 'Owner-furnished flow meter', 'Spare parts kit for pump'])
def test_legitimate_materials_are_visible_and_exported(name, status):
    record = material(name, status=status)
    assert exporter.reviewer_record_is_visible(record)
    assert len(json.loads(exporter.as_json(data(record)))['materials_and_equipment']) == 1


@pytest.mark.parametrize('status', ['PENDING', 'ACCEPTED', 'EDITED'])
@pytest.mark.parametrize('requirement', ['Leakage test at no additional cost',
                                       'Hydrostatic test using approved product data'])
def test_real_tests_are_not_hidden_by_incidental_administrative_words(requirement, status):
    record = {'kind': 'INSPECTION', 'meta': {}, 'review': {'status': status},
              'candidate': {'qa_type': 'FIELD_TEST', 'activity': 'test',
                            'requirement': requirement, 'evidence_ids': []}}
    assert exporter.reviewer_record_is_visible(record)
    assert len(json.loads(exporter.as_json(data(record)))['inspections_and_tests']) == 1


@pytest.mark.parametrize('status', ['PENDING', 'REJECTED', 'VERIFIED'])
@pytest.mark.parametrize('reviewed_only', [False, True])
def test_quantity_review_is_explicit_and_unapproved_quantity_is_not_reviewed_output(status, reviewed_only):
    payload = data(material(quantity_review=status), reviewed_only=reviewed_only)
    item = json.loads(exporter.as_json(payload))['materials_and_equipment'][0]
    assert item['quantity_review'] == exporter._label(status)
    assert item['quantity'] == (None if reviewed_only and status != 'VERIFIED' else 7)
    workbook = load_workbook(io.BytesIO(exporter.as_xlsx(payload)))
    sheet = workbook['Materials & Equipment']
    headers = [cell.value for cell in sheet[4]]
    assert 'Quantity Review' in headers
    assert sheet.cell(5, headers.index('Quantity Review') + 1).value == exporter._label(status)
    assert sheet.cell(5, headers.index('Quantity') + 1).value == item['quantity']


def test_item_review_does_not_reset_quantity_review_and_quantity_edit_does(client, project):
    rid = run_demo(client, project['id'])
    row = next(item for item in client.get(f'/api/analysis-runs/{rid}/records?kind=MATERIAL').json()
               if item['record']['candidate']['material_kind'] == 'PERMANENT')
    candidate = deepcopy(row['record']['candidate'])
    candidate['quantity'] = {'value': 7, 'unit': 'EA', 'basis': 'DESIGN_NET',
                             'method': 'EXPLICIT_DOCUMENT', 'scope_key': candidate['candidate_key'],
                             'entity_ids': candidate['entity_ids'], 'evidence_ids': candidate['evidence_ids'],
                             'formula': None, 'operands': [], 'calibration': None}
    uri = f'/api/records/{row["record"]["meta"]["record_id"]}/review'

    def review(action, **kwargs):
        nonlocal row
        response = client.post(uri, json={'action': action, 'expected_version': row['review_version'], **kwargs})
        assert response.status_code == 200, response.text
        row = response.json()
        return row['record']

    assert review('EDITED', candidate=candidate, note='Synthetic quantity review', quantity_action='REJECTED')['quantity_review'] == 'REJECTED'
    history=client.get(uri.rsplit('/',1)[0]+'/history').json()
    assert history[-1]['action']=='EDITED'
    assert json.loads(history[-1]['before_json'])['quantity_review']=='NOT_APPLICABLE'
    assert json.loads(history[-1]['after_json'])['quantity_review']=='REJECTED'
    assert review('ACCEPTED')['quantity_review'] == 'REJECTED'
    exported = client.get(f'/api/analysis-runs/{rid}/exports/json?reviewed_only=true').json()
    item = exported['materials_and_equipment'][0]
    assert item['quantity'] is None and item['quantity_review'] == 'Rejected'
    assert review('ACCEPTED', quantity_action='VERIFIED')['quantity_review'] == 'VERIFIED'
    original_review = deepcopy(row['record']['review'])
    assert review(None, quantity_action='REJECTED')['quantity_review'] == 'REJECTED'
    assert row['record']['review'] == original_review
    assert client.get(uri.rsplit('/',1)[0]+'/history').json()[-1]['action']=='QUANTITY_REJECTED'
    assert client.post(uri, json={'expected_version': row['review_version']}).status_code == 400
    assert review(None, quantity_action='VERIFIED')['quantity_review']=='VERIFIED'
    candidate['design_properties'].append({'name':'building','value':'Synthetic Building B',
                                           'unit':None,'evidence_ids':candidate['evidence_ids']})
    assert review('EDITED',candidate=candidate,note='Changed installation scope')['quantity_review']=='PENDING'
    candidate['quantity']['value'] = 8
    assert review('EDITED', candidate=candidate, note='Changed quantity')['quantity_review'] == 'PENDING'
    candidate['quantity'] = None
    assert review('EDITED', candidate=candidate, note='Removed quantity')['quantity_review'] == 'NOT_APPLICABLE'
