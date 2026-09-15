"""Export saved results as reviewer-readable English files without model calls."""
from __future__ import annotations

import io
import json
import re
import unicodedata
from collections import defaultdict
from datetime import datetime, timezone

from app.db import Database
from app.settings import ROOT


_LABELS = {
    'PARTIAL': 'Partial', 'COMPLETED': 'Completed', 'CANCELLED': 'Cancelled',
    'FAILED': 'Failed', 'PAUSED_PROVIDER': 'Paused by provider',
    'PENDING': 'Pending human review', 'ACCEPTED': 'Accepted', 'EDITED': 'Edited',
    'REJECTED': 'Rejected', 'NOT_CHECKED': 'Not checked', 'SUPPORTED': 'Supported',
    'PARTIAL_SUPPORT': 'Partially supported', 'UNSUPPORTED': 'Unsupported',
    'CONTRADICTED': 'Contradicted', 'NON_DOCUMENT': 'System-generated notice',
    'STALE': 'Needs recheck', 'CONFIRMED': 'Confirmed requirement',
    'CONDITIONAL': 'Conditional requirement', 'CONFLICT': 'Conflict',
    'INFERRED_TO_VERIFY': 'Inferred; verify before use',
    'PERMANENT': 'Permanent', 'TEMPORARY': 'Temporary',
    'INITIAL_INSPECTION': 'Initial inspection', 'FIELD_TEST': 'Field test',
    'TEST_REPORT': 'Test report', 'SUBMITTAL': 'Submittal',
    'DESIGN_PROPERTY': 'Design property', 'LATEST_APPLIED': 'Latest source applied',
    'UNRESOLVED': 'Unresolved', 'NO_CONFLICT': 'No conflict',
    'PENDING_REVIEW': 'Pending human review', 'VERIFIED': 'Verified',
    'TEXT_LAYER': 'Source text', 'OCR': 'OCR text', 'PARSER_TEXT': 'Source text',
    'SOURCE_TEXT': 'Source text', 'SOURCE_PARSED_TEXT': 'Source text',
    'MODEL_VISION_OUTPUT': 'Visual model observation', 'VISION': 'Visual model observation',
    'BLOCK_COUNT': 'Object count', 'LENGTH': 'Length', 'AREA': 'Area',
    'DESIGN_NET': 'Design net quantity', 'EXPLICIT_DOCUMENT': 'Stated in document',
    'SCHEDULE_EXTRACTION': 'Extracted from schedule', 'CAD_OBJECT_COUNT': 'CAD object count',
    'CAD_GEOMETRY': 'CAD geometry', 'CALCULATED': 'Calculated', 'DIRECT': 'Direct',
    'CLASSIFICATION': 'Classification', 'INFERRED': 'Inferred',
    'REVISION_POLICY': 'Revision comparison', 'SEARCH_RECORD': 'Search result',
    'NEEDS_SEMANTIC': 'Needs semantic review', 'NEEDS_CONTEXT': 'Needs more context',
}

_FIELD_LABELS = {
    '/name': 'Name', '/material_kind': 'Material type', '/quantity': 'Quantity',
    '/activity': 'Activity', '/requirement': 'Requirement', '/timing': 'Timing',
    '/frequency': 'Frequency', '/acceptance_criteria': 'Acceptance criteria',
    '/standard_reference': 'Standard', '/report_name': 'Report name',
    '/submission_trigger': 'Submission trigger', '/submission_offset_days': 'Submission offset',
    '/performer_as_stated': 'Performer', '/witness_as_stated': 'Witness',
    '/qa_type': 'Inspection or test type', '/subject': 'Subject',
    '/resolution_status': 'Resolution', '/reason': 'Resolution basis',
    '/missing_field_or_document': 'Missing information', '/blocking_reason': 'Impact',
    '/recommended_action': 'Recommended action', '/condition': 'Condition',
    '/location': 'Location', '/support_note': 'System note',
}

_KNOWN_SYSTEM_TEXT = {
    '合成演示，不代表工程结论。':
        'Synthetic demonstration only; this is not an engineering conclusion.',
    '模拟Provider不分析任意施工文件。请配置真实API；本片段尚未进行语义审查。':
        'Mock mode does not analyze uploaded construction files. Configure a live provider; this fragment has not received semantic review.',
    '工程审核候选；未审核或部分结果不可直接用作采购/施工依据。':
        'Engineering review candidates. Unreviewed or partial results are not procurement or construction instructions.',
    '仅导出已接受/已编辑的正式记录；未单独审批的 CAD 量算、PDF 几何审计和视觉任务未包含。':
        'Only accepted or edited records are included. Unapproved quantity aids are excluded.',
    '当前内容审查无法完成，需要处理后重新分析。':
        'The current review cannot be completed until the missing information is addressed.',
    '当前要求涉及未完成的父条款/上下文/选项关联；原始要求已保留。':
        'A parent requirement, context, or option relationship still needs review.',
    '原子要求已抽取，但尚无对应输出映射。':
        'The extracted requirement is not yet mapped to a final output category.',
    '核查已上传详图或补充依据。':
        'Check the uploaded details or provide the missing supporting document.',
    '这是整页矢量图元审计值，可能包含图框、标注和重复视图；未映射到材料，不作为设计净量。':
        'Whole-page vector audit only. It may include borders, annotations, or repeated views; it is not a material design quantity.',
    '模型空间块实例计数；可能包含图例或参考对象，人工确认范围后才能作为设计净量。':
        'Model-space block instance count. It may include legend or reference objects; confirm the scope before using it as a design quantity.',
    '模型空间图层几何合计；未自动排除图框、图例、辅助线或重复表达。':
        'Model-space layer geometry total. Borders, legends, construction lines, and repeated representations are not automatically excluded.',
    '模型空间图层封闭图元面积合计；未自动映射材料或去除重复表达。':
        'Model-space closed-entity area total. Materials are not automatically mapped and repeated representations are not removed.',
    '个': 'EA',
    '单杆吊架用吊杆': 'Hanger Rods for Single-Rod Hangers',
    '所有防火封堵材料': 'All Firestopping Materials',
    '结构化布线系统': 'Structured Cabling System',
    '给排水工程': 'Plumbing Work',
    '管道及管件': 'Piping and Fittings',
    '所有管道': 'All Piping',
    '管道系统保温': 'Piping System Insulation',
    '结构构件': 'Structural Members',
    '聚氨酯密封胶': 'Polyurethane Sealant',
    '加热水管道接头': 'Heating Water Piping Joints',
    '吊杆、螺母、垫圈、U形夹等': 'Hanger Rods, Nuts, Washers and U-Bolts',
    '低电压电缆(小于100伏)': 'Low-Voltage Cables (Below 100 V)',
    '密封材料安装表面': 'Sealant Application Surfaces',
    '机柜': 'Cabinet',
    '管道配件(2英寸及以下)': 'Pipe Fittings (2 Inches and Smaller)',
    '消防保护设施': 'Fire Protection Systems',
    '消防栓及配套管道、阀门、支架': 'Fire Hydrants with Associated Piping, Valves and Supports',
    '防风雨保温屋顶固定基座': 'Weatherproof Insulated Roof-Mounted Base',
    '3/4 HP及更小的电机': 'Motors 3/4 HP and Smaller',
    '给排水及湿式喷淋管道的防火封堵系统': 'Firestopping Systems for Plumbing and Wet-Pipe Sprinkler Piping',
    '相关电气底盒、导管、杂项电缆和电源': 'Associated Electrical Backboxes, Conduits, Miscellaneous Cabling and Power',
    '暴露于光线、交通、潮湿或物理损害的贯穿防火封堵系统': 'Through-Penetration Firestop Systems Exposed to Light, Traffic, Moisture or Physical Damage',
    '屋顶安装机组': 'Roof-Mounted Units',
    'F. Piping - 1-1/2" through 15" 的 Adapters': 'Adapters for 1-1/2-Inch Through 15-Inch Piping',
    '加热水管道配件': 'Heating Water Piping Fittings',
    '铝制铭牌': 'Aluminum Nameplates',
    'F. Piping - 1-1/2" through 15" 的 Joints': 'Joints for 1-1/2-Inch Through 15-Inch Piping',
    '独立建筑材料检测实验室': 'Independent Construction Materials Testing Laboratory',
    '暴露于视线的贯穿防火封堵系统': 'Exposed Through-Penetration Firestop Systems',
    '塑壳断路器': 'Molded-Case Circuit Breaker',
    '接缝背衬': 'Joint Backing',
    '防火封堵材料': 'Firestopping Materials',
    '给排水及湿式喷淋管道用防火封堵材料': 'Firestopping Materials for Plumbing and Wet-Pipe Sprinkler Piping',
    '所有管道标记': 'All Pipe Markers',
    'CK-13双碟片对夹式止回阀': 'CK-13 Dual-Disc Wafer Check Valve',
    '空心混凝土砌块膨胀锚栓': 'Expansion Anchors for Hollow Concrete Masonry Units',
    '1 HP及更大的电机': 'Motors 1 HP and Larger',
    '入水口防护': 'Inlet Protection',
    '现有干式喷淋系统': 'Existing Dry-Pipe Sprinkler System',
    '穿越所有耐火构造的防火封堵系统': 'Firestopping Systems at All Fire-Resistance-Rated Construction',
    '机械压接连接过滤器': 'Mechanical Press-Connection Strainer',
    '底涂': 'Primer',
    '塑料铭牌': 'Plastic Nameplates',
    '外护套': 'Outer Jacket',
    '给排水及湿式喷淋管道防火封堵系统': 'Firestopping Systems for Plumbing and Wet-Pipe Sprinkler Piping',
    '外墙': 'Exterior Walls',
    'ST-2过滤器': 'ST-2 Strainer',
    '直径3/4"至14"的管道': 'Piping 3/4 Inch Through 14 Inches in Diameter',
    '屋顶': 'Roof',
    '现有湿式喷淋系统': 'Existing Wet-Pipe Sprinkler System',
    '所有管道标记(购买或模印)': 'All Pipe Markers (Purchased or Stenciled)',
    'ST-1过滤器': 'ST-1 Strainer',
    '接缝清洁剂': 'Joint Cleaner',
    '其他未列名制造商': 'Other Unlisted Manufacturers',
    'D. Piping 21” and Over 的 Joints': 'Joints for Piping 21 Inches and Larger',
    '砌体锚固': 'Masonry Anchorage',
    '空气静压箱内的贯穿防火封堵系统': 'Through-Penetration Firestop Systems in Air Plenums',
    'E. Piping - All Sizes 的 Adapters': 'Adapters for All Piping Sizes',
    '材料与系统': 'Materials and Systems',
}

_TRANSLATION_PATH = ROOT / 'outputs' / 'export_english_translation_map.json'
_TRANSLATION_MTIME = None
_TRANSLATIONS: dict[str, str] = {}
_CJK = re.compile(r'[\u3400-\u9fff]')
_VISUAL_REVIEW_NOTE = 'Visual model observation recorded for this page. Review the cited source page before use.'
_INTERNAL_REFERENCE = re.compile(r'\b(?:REC|EV|DOC|CALL|RUN)-[0-9a-fA-F]{16,64}\b')
_NUMBER_TOKEN = re.compile(r'\d+(?:[./]\d+)*')
_SNAKE_CODE = re.compile(r'\b[a-z][a-z0-9]*(?:_[a-z0-9]+)+\b', re.I)
_PDF_PAGE_DIAGNOSTIC = re.compile(r'PDF\s*第\s*(\d+)\s*页', re.I)
_GENERIC_TRANSLATION = re.compile(
    r'\b(?:this|the) input contains (?:a )?visual description of a drawing page\b', re.I
)
_VISUAL_SEGMENT = re.compile(r'(Visual observation|Visual limitation|Explicit quantity text):', re.I)
_ACTOR_TEXT = re.compile(
    r'\b(?:owner|contractor|subcontractor|design[- ]builder|architect|engineer|installer|'
    r'manufacturer|supplier|testing agency|inspection agency|inspector|qc manager|'
    r'quality control manager|laboratory|fabricator|commissioning authority|balancing contractor)\b',
    re.I,
)
_MATERIAL_DESCRIPTOR = re.compile(
    r'^(?:(?:minimum\s+)?insulation thickness(?: per pipe size)?(?:\s+\S+\s*(?:inch(?:es)?|in\.?|mm))?|'
    r'piping(?: system)? insulation thickness|'
    r'water service pipe material and size|\d+(?:\.\d+)?\s*(?:gauge|awg)|'
    r'(?:minimum|maximum)?\s*(?:thickness|size|diameter|pressure|temperature|voltage|flow rate)|'
    r'design pressure|output voltage|moisture resistance?|pipe type|control joint spacing|'
    r'(?:suspended )?rod diameter specification table|'
    r'one size|#\d+\s*bars?\s*@\s*\S+\s*o\.?c\.?|'
    r'ul system no\.?|ansi\s+[a-z0-9.-]+)$',
    re.I,
)
_NON_MATERIAL_NAME = re.compile(
    r'\b(?:shop drawings?|submittals?|table of contents|schedule of values|record documents?|'
    r'operation and maintenance|o\s*&\s*m\b|manuals?|warrant(?:y|ies)|schematic|diagrams?|'
    r'transmittals?|calculations?|reports?|certificates?|instructions?|spare parts?|extra materials?)\b',
    re.I,
)
_QA_ACTION = re.compile(
    r'\b(?:inspect(?:ed|ing|ion)?|test(?:ed|ing)?|check(?:ed|ing)?|verify|verification|witness|hold point|'
    r'commission(?:ing)?|start[- ]?up|balance|balancing|calibrat(?:e|ion)|functional performance|'
    r'hydrostatic|pressure test|leak(?:age)? test|continuity|megger|torque test|'
    r'bacteriological|airflow|sound level|vibration)\b',
    re.I,
)
_NON_QA_DOCUMENT = re.compile(
    r'\b(?:shop drawings?|product data|submittal register|schedule of values|table of contents|'
    r'transmittals?|record documents?|schematic|wiring diagrams?|coordination drawings?|'
    r'operation and maintenance|o\s*&\s*m\b|manuals?|warrant(?:y|ies)|cost|bid|'
    r'time impact analysis|schedule extension|total float|critical path|unstamped submittals?)\b',
    re.I,
)
_ADMINISTRATIVE_VERIFY = re.compile(
    r'\bverify\b.{0,80}\b(?:this information|dimensions?|relationships?|existing equipment sizes?|interferences?|'
    r'contract documents?|routing|quantities submitted)\b|\bbase\b.{0,100}\btest data\b',
    re.I,
)
_NON_QA_REQUIREMENT = re.compile(
    r'^\s*(?:include all test results|complete all applicable tests|coordinate\b|refer to applicable code|'
    r'include in the punch list|visit the premises|review and document attendance|audit and approve|'
    r'correct any items noted)|'
    r'\b(?:applicable tests, certifications, forms, and matrices|test(?:ing)? plan(?: summary| and log)?|'
    r'must be installed prior to testing|tests? that fail.{0,80}remedied|'
    r'(?:attachment points.{0,80}written certification|written certification.{0,80}attachment points)|'
    r'locations? of wall-mounted devices.{0,80}drawings)\b',
    re.I,
)
_DESCRIPTIVE_MATERIAL_PHRASE = re.compile(
    r'\b(?:including|constructed from|made of|classified by|listed by|designed for use|'
    r'flame-spread|smoke-developed|receive insulation thickness|sizes? up to)\b',
    re.I,
)
_REPORT_DELIVERABLE = re.compile(
    r'\b(?:report|log|record|results?|certif(?:y|ied|icate|ication)|test data)\b',
    re.I,
)
_GENERIC_QA_NAME = re.compile(
    r'^(?:all\s+)?(?:testing|test|inspection|report|results? qa report|qa report|inspection report|test report|'
    r'pipes? test|penetrations? (?:test|inspection)|site inspection and testing)$',
    re.I,
)
_ATTRIBUTE_ONLY_MATERIAL = re.compile(
    r'^(?:all new materials?\b|all additional hardware\b|all documentation\b|products? that\b|'
    r'materials? and systems?\b|state fire marshal approved\b|wet location listed\b|'
    r'(?:efficiencies|standards?) listed\b|listed (?:category|standards?)\b|'
    r'low emitting material limits\b|operating temperature range|maximum (?:ambient temperature|input voltage)|'
    r'design pressure and temperature|mppt operating voltage range|temperature rise|dc voltage ripple|'
    r'to give a control voltage\b|with \d+(?:\.\d+)?\s*lb\b|\d+(?:\.\d+)?\s*psi\b|'
    r'minimum of \d+(?:\.\d+)?\s*gauge\b|'
    r'(?:aluminum|iron|steel|brass|copper|bronze)(?:,\s*(?:aluminum|iron|steel|brass|copper|bronze|galvanized steel|stainless steel)){3,})',
    re.I,
)
_ENGLISH_PUNCTUATION = str.maketrans({
    '，': ',', '。': '.', '；': ';', '：': ':', '、': ',', '！': '!', '？': '?',
    '（': '(', '）': ')', '【': '[', '】': ']', '《': '"', '》': '"',
})
_MONTH_NUMBERS = {
    'january': '1', 'february': '2', 'march': '3', 'april': '4', 'may': '5', 'june': '6',
    'july': '7', 'august': '8', 'september': '9', 'october': '10', 'november': '11', 'december': '12',
}
_WORD_NUMBERS = {
    'zero': '0', 'one': '1', 'two': '2', 'three': '3', 'four': '4', 'five': '5',
    'six': '6', 'seven': '7', 'eight': '8', 'nine': '9', 'ten': '10', 'eleven': '11',
    'twelve': '12', 'thirteen': '13', 'fourteen': '14', 'fifteen': '15', 'sixteen': '16',
    'seventeen': '17', 'eighteen': '18', 'nineteen': '19', 'twenty': '20',
}
_ORDINAL_NUMBERS = {
    'first': '1', 'second': '2', 'third': '3', 'fourth': '4', 'fifth': '5',
    'sixth': '6', 'seventh': '7', 'eighth': '8', 'ninth': '9', 'tenth': '10',
    'eleventh': '11', 'twelfth': '12', 'thirteenth': '13', 'fourteenth': '14',
    'fifteenth': '15', 'sixteenth': '16', 'seventeenth': '17', 'eighteenth': '18',
    'nineteenth': '19', 'twentieth': '20',
}


def _number_tokens(value: str) -> set[str]:
    tokens = _NUMBER_TOKEN.findall(value)
    lowered = value.casefold()
    tokens.extend(number for month, number in _MONTH_NUMBERS.items() if month in lowered)
    tokens.extend(number for word, number in _WORD_NUMBERS.items()
                  if re.search(rf'\b{word}\b', lowered))
    tokens.extend(number for word, number in _ORDINAL_NUMBERS.items()
                  if re.search(rf'\b{word}\b', lowered))
    return set(tokens)


def _translation_is_safe(source: str, translation: str) -> bool:
    if not translation or _CJK.search(translation) or _GENERIC_TRANSLATION.search(translation):
        return False
    return not (_number_tokens(source) - _number_tokens(translation))


def _filter_translations(values: dict[str, str]) -> dict[str, str]:
    """Reject lossy or obviously collapsed local translations before display."""
    reverse = defaultdict(list)
    for source, translation in values.items():
        if translation and not _CJK.search(translation):
            reverse[re.sub(r'\s+', ' ', translation).strip().casefold()].append(source)
    collapsed = {translation for translation, sources in reverse.items()
                 if len(sources) >= 3 and len(translation) >= 80}
    return {source: translation for source, translation in values.items()
            if _translation_is_safe(source, translation)
            if re.sub(r'\s+', ' ', translation).strip().casefold() not in collapsed}


def _sanitize_visible_text(value: str, humanize_codes: bool = False) -> str:
    """Remove application references and PDF extraction codes from reviewer text."""
    # Some embedded PDF fonts expose common engineering symbols through the
    # Unicode private-use area. Normalize the two verified mappings before
    # applying ordinary Unicode and punctuation cleanup so reviewers see the
    # intended temperature and fiber-size units rather than font glyph codes.
    text = value.translate({0xF0B0: '°', 0xF06D: 'μ'})
    text = unicodedata.normalize('NFKC', text).translate(_ENGLISH_PUNCTUATION)
    text = re.sub(
        r'\b(?:Document ID|document_id)\s*(?::|is)?\s*DOC-[0-9a-fA-F]{16,64}'
        r'(?:\s+used as reference only)?',
        'source document reference', text, flags=re.I,
    )
    text = re.sub(r'\s*\((?:REC|EV|DOC|CALL|RUN)-[0-9a-fA-F]{16,64}\)', '', text)
    text = _INTERNAL_REFERENCE.sub('internal reference', text)
    text = re.sub(r'\(cid:176\)', '°', text, flags=re.I)
    text = re.sub(r'\(cid:\d+\)', ' ', text, flags=re.I)
    text = text.replace('\ufffd', '')
    text = re.sub(r'\bnull\b', 'not stated', text, flags=re.I)
    if humanize_codes:
        replacements = {
            'SPEC_PAGE': 'Specification page',
            'NEEDS_CONTEXT': 'Needs context',
            'ONE_OF': 'One of',
            'FRAGMENT_START': 'Fragment start',
        }
        for code, label in replacements.items():
            text = re.sub(rf'\b{code}\b', label, text, flags=re.I)
        text = _SNAKE_CODE.sub(lambda match: match.group(0).replace('_', ' '), text)
    text = re.sub(r',(?=[A-Za-z])', ', ', text)
    text = re.sub(r'\s+([,.;:])', r'\1', text)
    return re.sub(r'\s+', ' ', text).strip()


def _evidence_text(source_type: str, value) -> str:
    text = _sanitize_visible_text(_plain(value), humanize_codes=True)
    text = text.translate(_ENGLISH_PUNCTUATION)
    text = _CJK.sub(' ', text)
    text = re.sub(r'\bsheet id\b\s*:?(?=\s*[,.;]|\s*$)', ' ', text, flags=re.I)
    text = re.sub(r'\bsheet id\b\s*:?', 'Sheet ', text, flags=re.I)
    text = re.sub(r'\bscale text\s+(?:null|none|not stated)\b', ' ', text, flags=re.I)
    text = re.sub(r'\bnull\b', 'not stated', text, flags=re.I)

    def tidy(fragment: str) -> str:
        fragment = re.sub(r'\s+([,.;:!?])', r'\1', fragment)
        fragment = re.sub(r'([,.;:!?])(?:\s*[,.;:!?])+', r'\1', fragment)
        fragment = re.sub(r'(^|\s)[,.;:/\\-]+(?=\s|$)', ' ', fragment)
        fragment = re.sub(r',(?=\S)', ', ', fragment)
        cleaned = re.sub(r'\s+', ' ', fragment).strip()
        return cleaned if re.search(r'[A-Za-z0-9]', cleaned) else ''

    parts = _VISUAL_SEGMENT.split(text)
    if len(parts) == 1:
        return tidy(text)
    result = [tidy(parts[0])]
    for index in range(1, len(parts), 2):
        marker = parts[index]
        payload = tidy(parts[index + 1] if index + 1 < len(parts) else '')
        if re.search(r'[A-Za-z0-9]', payload):
            result.append(f'{marker}: {payload}')
    return re.sub(r'\s+', ' ', ' '.join(part for part in result if part)).strip()


def _plain(value) -> str:
    """Return compact display text without code fences or layout-only whitespace."""
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'Yes' if value else 'No'
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, dict):
        parts = []
        for key, item in value.items():
            text = _plain(item)
            if text:
                parts.append(f"{_human_key(key)}: {text}")
        return '; '.join(parts)
    if isinstance(value, (list, tuple)):
        return '; '.join(text for item in value if (text := _plain(item)))
    text = _KNOWN_SYSTEM_TEXT.get(str(value), str(value))
    text = unicodedata.normalize('NFKC', text)
    text = re.sub(r'```(?:json|javascript|text|markdown)?', ' ', text, flags=re.I)
    text = text.replace('```', ' ').replace('\u00ad', '')
    text = re.sub(r'[\u200b-\u200d\ufeff]', '', text)
    return re.sub(r'\s+', ' ', text).strip()


def _human_key(value) -> str:
    return _plain(str(value).replace('_', ' ')).title()


def _mapped_plain(value) -> str:
    raw = '' if value is None else str(value)
    text = _plain(value)
    if raw in _KNOWN_SYSTEM_TEXT:
        return text
    translations = _load_translations()
    return _plain(translations.get(raw, translations.get(text, text)))


def _label(value, default='Not stated') -> str:
    if value is None or value == '':
        return default
    return _LABELS.get(str(value), _human_key(value))


def _normalized_property_name(value) -> str:
    return re.sub(r'[^a-z0-9]+', '_', _mapped_plain(value).casefold()).strip('_')


def _candidate_property(candidate: dict, *names: str) -> dict | None:
    wanted={re.sub(r'[^a-z0-9]+','_',name.casefold()).strip('_') for name in names}
    return next((prop for prop in candidate.get('design_properties') or []
                 if _normalized_property_name(prop.get('name')) in wanted),None)


def _is_actor_text(value: str) -> bool:
    return bool(value and _ACTOR_TEXT.search(value))


def _title_if_upper(value: str) -> str:
    return value.title() if value and value.upper()==value else value


def _field_status(report: dict, path: str) -> str:
    field=next((item for item in report.get('fields',[]) if item.get('path')==path),None)
    return str((field or {}).get('status') or '')


def _reviewer_material_name(candidate: dict, report: dict) -> tuple[str,str]:
    """Return a tangible item name and any location recovered from a schedule identity."""
    name=_mapped_plain(candidate.get('name'))
    entities=[_mapped_plain(value) for value in candidate.get('entity_ids') or []]
    location=_mapped_plain(candidate.get('location'))
    raw_name=name
    if re.fullmatch(r'water service pipe material and size',name,re.I):
        service=_candidate_property(candidate,'type','service type')
        if service:
            name=_title_if_upper(_mapped_plain(service.get('value')))+' Pipe'
        if entities:
            match=re.match(r'(?P<tag>W\d+)\s+at\s+(?P<address>.+?)\s+combined\b',entities[0],re.I)
            if match and not location:location=f"{match.group('tag')} at {match.group('address')}"
    usable_entities=[value for value in entities
                     if value and not _is_actor_text(value)
                     and not _MATERIAL_DESCRIPTOR.fullmatch(value)
                     and value.casefold() not in {'equipment','material','materials','product','products','item','items','all piping'}]
    bad_name=(_MATERIAL_DESCRIPTOR.fullmatch(name) or
              name.casefold() in {'equipment','material','materials','product','products','item','items'} or
              _field_status(report,'/name')=='UNSUPPORTED' or
              (bool(usable_entities) and (
                  len(name)>max(64,len(usable_entities[0])+28) or
                  bool(_DESCRIPTIVE_MATERIAL_PHRASE.search(name)) or
                  bool(re.search(r'^\s*(?:heavy gauge|molded thermoplastic|direct expansion type|products? with|materials? and systems?)\b',name,re.I))
              )))
    if bad_name:
        replacement=next(iter(usable_entities),'')
        if replacement:name=replacement
    if _MATERIAL_DESCRIPTOR.fullmatch(name) and any('insulation' in value.casefold() for value in entities):
        name='Piping System Insulation'
    if not usable_entities:
        name=re.sub(r'^complete\s+(.+?)\s+including\b.*$',r'\1',name,flags=re.I)
        name=re.sub(r'^seismic restraint and equipment\b.*$', 'Seismic Restraints', name, flags=re.I)
        if re.search(r'\bpiping\b',name,re.I) and any('insulation thickness' in value.casefold() for value in entities):
            name=f'{name} Insulation'
    gauge_name=re.match(r'^(?:minimum of\s+)?(?P<gauge>\d+(?:\.\d+)?)\s*gauge\s+(?P<item>.+)$',name,re.I)
    if gauge_name and len(gauge_name.group('item').split())>=2:
        name=gauge_name.group('item')
    canonical_patterns=(
        (r'^all fire hydrants and associated piping, valves, and supports\b.*$',
         'Fire Hydrants with Associated Piping, Valves and Supports'),
        (r'^control low \(24V\) and control line \(120V\) voltage wiring, conduit, (?:and )?related switches and relays$',
         'Low-Voltage Control Wiring, Conduit, Switches and Relays'),
        (r'^(?:required\s+)?grounding lugs or other hardware for each piece of technology equipment\b.*$',
         'Technology Equipment Grounding Hardware'),
        (r'^factory fabricated metal edge system\b.*$', 'Factory-Fabricated Metal Edge System'),
        (r'^Type L hard drawn seamless copper tube\b.*$', 'Type L Hard-Drawn Seamless Copper Tube'),
        (r'^Schedule 5 Type 304 316 stainless steel\b.*$', 'Schedule 5 Type 304/316 Stainless Steel Tube'),
        (r'^factory applied protective coatings consisting of a polyethylene plastic film\b.*$',
         'Factory-Applied Polyethylene Protective Coating'),
    )
    for pattern,replacement in canonical_patterns:
        if re.search(pattern,name,re.I):name=replacement;break
    if re.fullmatch(r'one size',raw_name,re.I) and any(value.casefold()=='rods' for value in entities):
        name='Double-Rod Hanger Rods'
    name=re.sub(r'^\s*(?:provide|furnish and install|required)\s+', '', name, flags=re.I)
    name=re.sub(r'\s+exposed to view\s*$', '', name, flags=re.I)
    name=re.sub(r'\s+', ' ', name).strip(' .;:-')
    if name and name[0].islower():name=name[0].upper()+name[1:]
    if raw_name.casefold()=='18 gauge' and name.casefold()=='structural members':
        name='Structural Members'
    return name,location


def _reviewer_material_allowed(name: str) -> bool:
    if not name or len(name)>180 or _is_actor_text(name):return False
    if _NON_MATERIAL_NAME.search(name) or _MATERIAL_DESCRIPTOR.fullmatch(name):return False
    if name.casefold() in {
        'equipment','material','materials','product','products','item','items','piping system','all piping',
        'pipe','piping','fitting','fittings','valve','valves','conductor','conductors','coil','coils',
        'joint','joints','insulation','plumbing work','moisture resistant','software',
    }:return False
    if re.search(r'\bsoftware\b',name,re.I):return False
    if re.fullmatch(r'(?:astm|ansi|nfpa|ul|ashrae|iecc|scaqmd)\b.*',name,re.I):return False
    if re.search(r'\b(?:shall|must|responsibility|responsible for|approval is secured|used only by change orders)\b',name,re.I):return False
    if re.match(r'^(?:applicable|requirements for|for use with|any material|all equipment that|equipment in finished areas)\b',name,re.I):return False
    if _ATTRIBUTE_ONLY_MATERIAL.search(name):return False
    return True


def _inspection_performer(candidate: dict) -> str:
    stated=_mapped_plain(candidate.get('performer_as_stated'))
    if stated:return stated
    activity=_mapped_plain(candidate.get('activity'))
    if _is_actor_text(activity):return activity
    return next((_mapped_plain(value) for value in candidate.get('entity_ids') or []
                 if _is_actor_text(_mapped_plain(value))), '')


def _inspection_activity(candidate: dict) -> str:
    activity=_mapped_plain(candidate.get('activity'))
    requirement=_mapped_plain(candidate.get('requirement'))
    if requirement and _QA_ACTION.search(requirement) and (_is_actor_text(activity) or not _QA_ACTION.search(activity)):
        return requirement
    if activity and requirement and not _QA_ACTION.search(activity) and not _QA_ACTION.search(requirement):
        return f'{activity}: {requirement}'
    return activity or requirement


def _inspection_name(candidate: dict) -> str:
    qa_type=str(candidate.get('qa_type') or '')
    report_name=_mapped_plain(candidate.get('report_name'))
    requirement=_mapped_plain(candidate.get('requirement'))
    activity=_mapped_plain(candidate.get('activity'))
    if (activity and not _is_actor_text(activity) and _QA_ACTION.search(activity)
            and len(activity)+20<len(requirement)):
        source=activity
    elif (qa_type in {'FIELD_TEST','LAB_TEST'} and activity and not _is_actor_text(activity)
            and not _QA_ACTION.search(requirement)):
        source=activity
    elif (re.search(r'factory assembled and tested',requirement,re.I) and activity and not _is_actor_text(activity)):
        source=f'{activity} Factory Test'
    elif re.search(r'conform to\s+([A-Z]+[- ]?\d[\w.-]*)',requirement,re.I):
        standard=re.search(r'conform to\s+([A-Z]+[- ]?\d[\w.-]*)',requirement,re.I).group(1)
        source=f'{standard} Compliance Test'
    else:
        source=report_name if report_name and _QA_ACTION.search(report_name) and not _NON_QA_DOCUMENT.search(report_name) else requirement
    canonical_activity=activity if activity and not _is_actor_text(activity) else ''
    canonical_rules=(
        (r'inspect material for damage', 'Material Damage Inspection'),
        (r'test pipes in chases, walls, or above non-accessible ceilings',
         f'{canonical_activity} Test' if canonical_activity else 'Concealed Piping Test'),
        (r'swept frequency testing through 500\s*MHz', 'Swept-Frequency Channel Performance Test'),
        (r'(?:verify|be certain).{0,40}in place all strainer screens', 'Strainer Screen Installation Inspection'),
        (r'have inlet protection', 'Stormwater Inlet Protection Inspection'),
        (r'verify availability in proper location and ready for use', 'Required Utilities Readiness Inspection'),
        (r'(?:ensure|check).{0,30}correct installation',
         f'{canonical_activity} Installation Inspection' if canonical_activity else 'Installation Inspection'),
        (r'(?:dirt|dirt, dust).{0,80}(?:rust|loose material)', 'Sealant Contact Surface Cleanliness Inspection'),
        (r'verify (?:that )?abandoned utilities', 'Abandoned Utilities Verification'),
        (r'follow-on inspections?.{0,40}deficien', 'Follow-Up Deficiency Inspection'),
        (r'(?:inspect and )?certify equipment installation', 'Equipment Installation Inspection and Certification'),
        (r'voltage, frequency, and amperes', 'Electrical Voltage, Frequency and Amperage Test'),
        (r'visual destructive review of installed firestop', 'Installed Firestop Destructive Inspection'),
    )
    for pattern,replacement in canonical_rules:
        if re.search(pattern,requirement,re.I):source=replacement;break
    source=re.sub(r'^\s*(?:shall\s+)?(?:be\s+)?(?:provide[sd]?|perform(?:ed)?|conduct(?:ed)?|complete[sd]?|submit(?:ted)?)\s+', '', source, flags=re.I)
    source=re.sub(r'^\s*(?:shall|must|will)\s+be\s+', '', source, flags=re.I)
    source=re.sub(r'^\s*(?:inspect|verify|check|test)\s+', '', source, flags=re.I)
    source=re.sub(r'^\s*(?:and\s+)', '', source, flags=re.I)
    source=re.split(r'\b(?:before|after|when|where|at the time)\b',source,maxsplit=1,flags=re.I)[0]
    source=re.sub(r'^\s*(?:all|each|the|a|an)\s+', '', source, flags=re.I).strip(' .;:-')
    suffix='Inspection' if 'INSPECTION' in qa_type else 'Test' if qa_type in {'FIELD_TEST','LAB_TEST'} else 'QA Report'
    if source and not re.search(r'\b(?:inspection|test|testing|commissioning|startup|start-up|balancing|calibration|report)\b',source,re.I):
        source=f'{source} {suffix}'
    if source and source[0].islower():source=source[0].upper()+source[1:]
    return source


def _reviewer_inspection_allowed(candidate: dict, name: str, activity: str) -> bool:
    qa_type=str(candidate.get('qa_type') or '')
    allowed={'PREPARATORY_INSPECTION','INITIAL_INSPECTION','FOLLOWUP_INSPECTION','HOLD_POINT',
             'WITNESS_POINT','SPECIAL_INSPECTION','FIELD_TEST','LAB_TEST','MANUFACTURER_INSPECTION',
             'STARTUP','COMMISSIONING','TEST_REPORT','DAILY_FIELD_REPORT'}
    text=' '.join(value for value in (name,activity,_mapped_plain(candidate.get('requirement')),
                                      _mapped_plain(candidate.get('report_name'))) if value)
    if qa_type not in allowed or not name or len(name)>180:return False
    if _GENERIC_QA_NAME.fullmatch(name.strip()):return False
    if not _QA_ACTION.search(text):return False
    if _NON_QA_DOCUMENT.search(text) or _ADMINISTRATIVE_VERIFY.search(text):return False
    if _NON_QA_REQUIREMENT.search(text):return False
    if re.fullmatch(r'\d{2}\s+\d{2}\s+\d{2}\s+(?:testing|inspection|test)',name,re.I):return False
    if qa_type in {'TEST_REPORT','DAILY_FIELD_REPORT'} and not _REPORT_DELIVERABLE.search(
            ' '.join(value for value in (_mapped_plain(candidate.get('report_name')),
                                         _mapped_plain(candidate.get('requirement'))) if value)):
        return False
    if re.search(r'\b(?:may include|will be rejected|is responsible to verify)\b',text,re.I):return False
    return not _is_actor_text(name)


def reviewer_record_is_visible(record: dict, report: dict | None = None) -> bool:
    """Apply the same conservative reviewer scope to the UI and exports."""
    candidate=record.get('candidate') or {};report=report or {}
    if record.get('kind')=='MATERIAL':
        name,_=_reviewer_material_name(candidate,report)
        return _reviewer_material_allowed(name)
    if record.get('kind')=='INSPECTION':
        name=_inspection_name(candidate);activity=_inspection_activity(candidate)
        return _reviewer_inspection_allowed(candidate,name,activity)
    return False


def reviewer_record_display(record: dict, report: dict | None = None) -> dict:
    """Return reader-facing list fields without changing the stored candidate."""
    candidate=record.get('candidate') or {};report=report or {}
    if record.get('kind')=='MATERIAL':
        name,location=_reviewer_material_name(candidate,report)
        return {'name':name,'location':location,'performer':'','activity':'',
                'specification_section':', '.join(candidate.get('csi_sections') or [])}
    if record.get('kind')=='INSPECTION':
        return {'name':_inspection_name(candidate),'location':_mapped_plain(candidate.get('location')),
                'performer':_inspection_performer(candidate),'activity':_inspection_activity(candidate),
                'specification_section':', '.join(candidate.get('csi_sections') or [])}
    return {'name':'','location':'','performer':'','activity':'','specification_section':''}


def _load_translations() -> dict[str, str]:
    """Read the optional local compatibility cache; never invokes a model."""
    global _TRANSLATION_MTIME, _TRANSLATIONS
    try:
        mtime = _TRANSLATION_PATH.stat().st_mtime_ns
        if mtime != _TRANSLATION_MTIME:
            payload = json.loads(_TRANSLATION_PATH.read_text(encoding='utf-8'))
            raw = {str(key): str(value) for key, value in payload.get('translations', {}).items()}
            filtered = _filter_translations(raw)
            normalized_aliases = {
                _plain(source): translation for source, translation in filtered.items()
                if source not in _KNOWN_SYSTEM_TEXT
            }
            _TRANSLATIONS = {**normalized_aliases, **filtered}
            _TRANSLATION_MTIME = mtime
    except (FileNotFoundError, OSError, ValueError, TypeError):
        return _TRANSLATIONS
    return _TRANSLATIONS


def _translate_display(value, path=()):
    """Apply prebuilt English phrases while preserving source-evidence wording."""
    translations = _load_translations()
    if isinstance(value, dict):
        return {key: _translate_display(item, path + (key,)) for key, item in value.items()}
    if isinstance(value, list):
        return [_translate_display(item, path + (index,)) for index, item in enumerate(value)]
    if isinstance(value, str):
        is_evidence_text = bool(path and path[-1] in {'text', 'evidence_text'})
        translated = value if is_evidence_text else translations.get(value, value)
        if not is_evidence_text and _CJK.search(translated):
            field_path = '/'.join(str(part) for part in path)
            raise ValueError(
                'English export cache is incomplete for saved legacy text. '
                f'Rebuild the local translation cache before exporting. Field: {field_path}'
            )
        if is_evidence_text:
            return _evidence_text('', translated)
        return _sanitize_visible_text(translated, humanize_codes=True)
    return value


def _finalize_conflicts(readable: dict) -> dict:
    """Show each recorded value beside a concise, unchanged source passage."""
    for conflict in readable.get('conflicts', []):
        unique_statements = []
        seen = set()
        for statement in conflict.get('statements', []):
            identity = tuple(_plain(statement.get(key)) for key in (
                'value', 'source', 'location', 'revision_date', 'source_type', 'evidence_text'
            ))
            if identity in seen:
                continue
            seen.add(identity)
            unique_statements.append(statement)
        conflict['statements'] = unique_statements
        for statement in unique_statements:
            value = _plain(statement.get('value'))
            source_text = _plain(statement.get('evidence_text'))
            position = source_text.casefold().find(value.casefold()) if value else -1
            excerpt = ''
            if position >= 0:
                start = max(0, position - 140)
                end = min(len(source_text), position + len(value) + 180)
                for marker in ('. ', ': ', '; '):
                    boundary = source_text.rfind(marker, start, position)
                    if boundary >= start:
                        start = boundary + len(marker)
                        break
                boundaries = [source_text.find(marker, position + len(value), end)
                              for marker in ('. ', '; ')]
                boundaries = [boundary for boundary in boundaries if boundary >= 0]
                if boundaries:
                    end = min(boundaries) + 1
                excerpt = source_text[start:end].strip(' ,;:')
            kind = 'visual value' if statement.get('source_type') == 'Visual model observation' else 'conflicting value'
            statement['evidence_text'] = f'Recorded {kind}: {value}.'
            if excerpt:
                statement['evidence_text'] += f' Source passage: {excerpt}'
            elif source_text:
                statement['evidence_text'] += f' Source passage: {source_text[:360].rstrip()}'
            else:
                statement['evidence_text'] += ' Review the cited source page before use.'
    return readable


def _locator(value) -> str:
    if not isinstance(value, dict):
        return _plain(value)
    parts = []
    mapping = (('page_number', 'Page'), ('sheet', 'Sheet'), ('section', 'Section'),
               ('paragraph', 'Paragraph'), ('text_line_start', 'Line'))
    for key, name in mapping:
        item = value.get(key)
        if item not in (None, ''):
            parts.append(f'{name} {_plain(item)}')
    start, end = value.get('text_line_start'), value.get('text_line_end')
    if start not in (None, '') and end not in (None, '', start):
        parts[-1] = f'Lines {_plain(start)}-{_plain(end)}'
    return ' | '.join(parts)


def _field_label(field: dict, record: dict) -> str:
    path = field.get('path', '')
    match = re.fullmatch(r'/design_properties/(\d+)/value', path)
    if match:
        props = record.get('candidate', {}).get('design_properties') or []
        index = int(match.group(1))
        return _human_key(_mapped_plain(props[index].get('name'))) if index < len(props) else 'Design property'
    match = re.fullmatch(r'/claims/(\d+)/normalized_value', path)
    if match:
        return f'Statement {int(match.group(1)) + 1}'
    if path.startswith('/csi_sections/'):
        return 'Specification section'
    if path.startswith('/entity_ids/'):
        return 'Document tag'
    return _FIELD_LABELS.get(path, _human_key(path.rsplit('/', 1)[-1]) or 'Item')


def _citation_item(citation: dict, field: str, status: str) -> dict:
    basis = citation.get('text_basis') or citation.get('content_basis')
    source_type = _label(basis or 'SOURCE_PARSED_TEXT', 'Source text')
    text = _evidence_text(source_type, citation.get('quote') or citation.get('raw_text'))
    return {
        'supports': field,
        'verification': _label(status, 'Not checked'),
        'source': _plain(citation.get('file_name')),
        'location': _locator(citation.get('locator')),
        'revision_date': _plain(citation.get('internal_revision_date')),
        'source_type': source_type,
        'text': text,
    }


def _evidence_ids(record: dict) -> list[str]:
    candidate = record.get('candidate', {})
    values = list(candidate.get('evidence_ids') or candidate.get('related_evidence_ids') or [])
    values.extend(claim.get('evidence_id') for claim in (candidate.get('claims') or []) if claim.get('evidence_id'))
    quantity = candidate.get('quantity')
    if isinstance(quantity, dict):
        values.extend(quantity.get('evidence_ids', []))
    return list(dict.fromkeys(values))


def _focused_source_text(evidence: dict, record: dict | None = None) -> str:
    """Start a fallback excerpt at the clause most relevant to the displayed item."""
    source_type = _label(evidence.get('content_basis') or evidence.get('extraction_method'), 'Source text')
    text = _evidence_text(source_type, evidence.get('raw_text'))
    if not record or not text:
        return text
    candidate=record.get('candidate') or {}
    needles=[_mapped_plain(candidate.get(key)) for key in
             ('requirement','report_name','name','activity','acceptance_criteria')]
    needles=[re.sub(r'\s+',' ',value).strip(' .;:-') for value in needles if len(value)>=8]
    lowered=text.casefold()
    for needle in sorted(needles,key=len,reverse=True):
        position=lowered.find(needle.casefold())
        if position>=0:
            start=max(0,position-120);end=min(len(text),position+len(needle)+360)
            boundary=max(text.rfind('. ',start,position),text.rfind('; ',start,position),text.rfind('\n',start,position))
            if boundary>=start:start=boundary+2
            stops=[value for value in (text.find('. ',position+len(needle),end),
                                       text.find('; ',position+len(needle),end)) if value>=0]
            if stops:end=min(stops)+1
            return text[start:end].strip()
    keywords={token.casefold() for needle in needles for token in re.findall(r'[A-Za-z0-9-]{4,}',needle)
              if token.casefold() not in {'shall','with','from','that','this','each','provide','required','requirement'}}
    segments=[segment.strip() for segment in re.split(r'(?<=[.;])\s+',text) if len(segment.strip())>=20]
    if keywords and segments:
        best=max(segments,key=lambda segment:sum(token in segment.casefold() for token in keywords))
        if sum(token in best.casefold() for token in keywords)>=2:return best
    return text


def _fallback_evidence(evidence: dict, field='Item', record: dict | None = None) -> dict:
    source_type = _label(evidence.get('content_basis') or evidence.get('extraction_method'), 'Source text')
    text = _focused_source_text(evidence,record)
    return {
        'supports': field,
        'verification': 'Not checked',
        'source': _plain(evidence.get('file_name')),
        'location': _locator(evidence.get('locator')),
        'revision_date': _plain(evidence.get('internal_revision_date')),
        'source_type': source_type,
        'text': text,
    }


def _record_evidence(data: dict, record: dict, evidence_index: dict | None = None) -> list[dict]:
    evidence_index = evidence_index or {item.get('evidence_id'): item for item in data.get('evidence', [])}
    record_id = record.get('meta', {}).get('record_id')
    report = data.get('verifications', {}).get(record_id, {})
    grouped = {}
    for field in report.get('fields', []):
        label = _sanitize_visible_text(_mapped_plain(_field_label(field, record)), humanize_codes=True)
        for citation in field.get('citations', []):
            item = _citation_item(citation, label, field.get('status', 'NOT_CHECKED'))
            key = (item['source'], item['location'], item['revision_date'], item['source_type'], item['text'])
            current = grouped.setdefault(key, {**item, '_fields': [], '_states': []})
            if label not in current['_fields']:
                current['_fields'].append(label)
            state = item['verification']
            if state not in current['_states']:
                current['_states'].append(state)
    if not grouped:
        for evidence_id in _evidence_ids(record):
            if evidence_id in evidence_index:
                item = _fallback_evidence(evidence_index[evidence_id],record=record)
                key = (item['source'], item['location'], item['revision_date'], item['source_type'], item['text'])
                grouped.setdefault(key, {**item, '_fields': ['Item'], '_states': ['Not checked']})
    result = []
    for item in grouped.values():
        item['supports'] = ', '.join(item.pop('_fields'))
        item['verification'] = ', '.join(item.pop('_states'))
        result.append(item)
    return result


def _evidence_from_ids(ids: list[str], evidence_index: dict) -> list[dict]:
    return [_fallback_evidence(evidence_index[value], 'Quantity') for value in dict.fromkeys(ids)
            if value in evidence_index]


def _properties(candidate: dict) -> str:
    rows = []
    quantity_prop = _explicit_quantity_property(candidate) if not isinstance(candidate.get('quantity'), dict) else None
    for index, prop in enumerate(candidate.get('design_properties') or []):
        if quantity_prop and index == quantity_prop[0]:
            continue
        name = _human_key(_mapped_plain(prop.get('name'))) if prop.get('name') else 'Property'
        value = _mapped_plain(prop.get('value'))
        if value.isupper() and len(re.sub(r'[^A-Z]', '', value)) >= 6:
            value=value.title()
        unit = _mapped_plain(prop.get('unit'))
        rows.append(f"{name}: {value}{(' ' + unit) if unit and unit not in value else ''}")
    raw_name=_mapped_plain(candidate.get('name'))
    thickness=re.fullmatch(r'(?:minimum\s+)?insulation thickness(?: per pipe size)?\s+(.+)',raw_name,re.I)
    if thickness and not any(row.casefold().startswith('insulation thickness:') for row in rows):
        rows.append(f'Insulation Thickness: {thickness.group(1)}')
    gauge=re.match(r'^(?:minimum of\s+)?(?P<gauge>\d+(?:\.\d+)?)\s*gauge\b',raw_name,re.I)
    if gauge and not any(row.casefold().startswith(('gauge:','minimum thickness:')) for row in rows):
        rows.append(f"Gauge: {gauge.group('gauge')}")
    if re.fullmatch(r'one size',raw_name,re.I) and 'double rod hanger' in _mapped_plain(candidate.get('condition')).casefold():
        rows.append('Rod Size Adjustment: One size smaller than the single-rod hanger requirement')
    pressure_temperature=re.search(
        r'\((?:up to\s+)?(?P<pressure>\d+(?:\.\d+)?)\s*psig,\s*(?P<temperature>-?\d+(?:\.\d+)?)\s*°?F\)',
        raw_name,re.I,
    )
    if pressure_temperature:
        rows.extend((f"Service Pressure: {pressure_temperature.group('pressure')} psig",
                     f"Service Temperature: {pressure_temperature.group('temperature')} °F"))
    return '\n'.join(rows)


def _options(candidate: dict) -> str:
    rows = []
    for option in candidate.get('options') or []:
        values = [_mapped_plain(option.get(name)) for name in ('manufacturer', 'product', 'model_number', 'role')]
        text = ' | '.join(value for value in values if value)
        if text and text not in rows:
            rows.append(text)
    return '\n'.join(rows)


def _quantity(candidate: dict) -> tuple[object | None, str, str]:
    quantity = candidate.get('quantity')
    if not isinstance(quantity, dict):
        derived = _explicit_quantity_property(candidate)
        if derived:
            _, value, unit = derived
            return value, unit, 'Stated in document'
        leading=re.match(
            r'^\s*(?P<count>\d+|one|two|three|four|five|six|seven|eight|nine|ten)\s+'
            r'(?:additional\s+)?(?P<unit>sets?|each|pieces?|units?)\s+of\b',
            _mapped_plain(candidate.get('name')),re.I,
        )
        if leading:
            count=leading.group('count').casefold()
            value=int(_WORD_NUMBERS.get(count,count))
            return value, leading.group('unit').rstrip('s').title(), 'Stated in document'
        return None, '', ''
    value = quantity.get('value')
    return value, _plain(quantity.get('unit')), _label(quantity.get('basis'), '')


def _explicit_quantity_property(candidate: dict) -> tuple[int, object, str] | None:
    for index, prop in enumerate(candidate.get('design_properties') or []):
        raw_name = _plain(prop.get('name')).casefold()
        name = _mapped_plain(prop.get('name')).casefold()
        raw_value = prop.get('value')
        value_text = _plain(raw_value)
        is_quantity = (bool(re.search(r'\b(?:quantity|count|number|total)\b', name)) or
                       bool(re.search(r'(?:数量|个数|总数)$', raw_name)))
        if (not is_quantity or
                not re.fullmatch(r'-?\d+(?:\.\d+)?', value_text)):
            continue
        value = float(value_text) if '.' in value_text else int(value_text)
        return index, value, _mapped_plain(prop.get('unit'))
    return None


def _summary(data: dict, material_count: int, inspection_count: int) -> dict:
    files = sorted({_plain(item.get('file_name')) for item in data.get('evidence', []) if item.get('file_name')})
    cost = data.get('cost', {})
    return {
        'project': _plain(data.get('project_name')) or 'Local CIRP project',
        'run_status': _label(data.get('run', {}).get('status'), 'Unknown'),
        'generated_utc': _plain(data.get('generated_at')),
        'review_scope': ('Accepted or edited material and QA items only' if data.get('reviewed_only') else
                         'Tangible material/equipment items and executable test/inspection requirements'),
        'source_files': files,
        'material_and_equipment_items': material_count,
        'inspection_and_test_items': inspection_count,
        'model_calls': cost.get('calls', 0),
        'model_cost_cny': cost.get('spent_cny', '0.000000'),
        'unresolved_model_calls': cost.get('unknown_calls', 0),
        'notice': _plain(data.get('notice')),
        'evidence_display_note': 'Whitespace and isolated non-English OCR artifacts are normalized for readability; source wording is otherwise unchanged. Visual model observations are labeled and are not source-document quotations.',
        'quantity_note': 'A blank quantity means the source did not state a usable quantity or no reviewed design quantity was calculated. It does not mean zero.',
    }


def _material_item(data: dict, record: dict, evidence_index: dict) -> dict:
    candidate = record['candidate']; quantity, unit, basis = _quantity(candidate)
    report = data.get('verifications', {}).get(record.get('meta', {}).get('record_id'), {})
    name,location = _reviewer_material_name(candidate,report)
    if not _reviewer_material_allowed(name):return None
    if quantity is not None:
        name, minimum = _separate_name_quantity(
            name, quantity, unit, allow_bare=not isinstance(candidate.get('quantity'), dict)
        )
        if minimum:
            basis = 'Minimum stated in document'
        if quantity is not None and unit.casefold()=='set':
            name=re.sub(r'^(?:additional\s+)?sets?\s+of\s+(?:each\s+)?','',name,flags=re.I)
    return {
        'material_or_equipment_name': name,
        'material_type': _label(candidate.get('material_kind'), ''),
        'quantity': quantity,
        'unit': unit,
        'quantity_basis': basis,
        'design_properties': _properties(candidate),
        'manufacturer_product_model_options': _options(candidate),
        'specification_section': ', '.join(_plain(value) for value in (candidate.get('csi_sections') or []) if _plain(value)),
        'location': location,
        'condition': _plain(candidate.get('condition')),
        'requirement_status': _label(candidate.get('requirement_status'), ''),
        'verification': _label(report.get('status'), 'Not checked'),
        'human_review': _label(record.get('review', {}).get('status'), 'Pending human review'),
        'evidence': _record_evidence(data, record, evidence_index),
    }


def _separate_name_quantity(name: str, quantity, unit: str = '', *, allow_bare: bool = False) -> tuple[str, bool]:
    """Remove an item-count prefix without mistaking a dimensional size for quantity."""
    text = _plain(name)
    number = _plain(quantity)
    try:
        numeric = float(number)
        if numeric.is_integer():
            number = str(int(numeric))
    except (TypeError, ValueError):
        return text, False
    words = {value: key for key, value in _WORD_NUMBERS.items()}
    word = words.get(number, '')
    tokens = [rf'\(\s*{re.escape(number)}\s*\)']
    if word:
        tokens.insert(0, rf'{word}\s*\(\s*{re.escape(number)}\s*\)')
        tokens.append(re.escape(word))
    if number == '1':
        tokens.append(r'single')
    tokens.append(rf'{re.escape(number)}\s*(?:EA|each|pcs?|pieces?|units?)\b')
    if allow_bare or _plain(unit).casefold() in {'ea', 'each', 'pc', 'pcs', 'piece', 'pieces', 'unit', 'units'}:
        tokens.append(re.escape(number))
    pattern = re.compile(
        rf'^\s*(?P<minimum>(?:at\s+least|minimum(?:\s+of)?))?\s*(?:{"|".join(tokens)})(?=\s)',
        re.I,
    )
    match = pattern.match(text)
    if match:
        separated = text[match.end():].strip(' :-–—')
        return (separated or text), bool(match.group('minimum'))
    embedded = re.sub(rf'(?<![A-Za-z0-9])\(\s*{re.escape(number)}\s*\)(?![A-Za-z0-9])', ' ', text)
    embedded = re.sub(r'\s+', ' ', embedded).strip()
    return (embedded or text), False


def _inspection_item(data: dict, record: dict, evidence_index: dict) -> dict:
    candidate = record['candidate']
    report = data.get('verifications', {}).get(record.get('meta', {}).get('record_id'), {})
    activity=_inspection_activity(candidate)
    name=_inspection_name(candidate)
    if not _reviewer_inspection_allowed(candidate,name,activity):return None
    acceptance=_plain(candidate.get('acceptance_criteria'))
    requirement=_mapped_plain(candidate.get('requirement'))
    raw_activity=_mapped_plain(candidate.get('activity'))
    if not acceptance and requirement and not _QA_ACTION.search(requirement) and raw_activity and not _is_actor_text(raw_activity):
        acceptance=requirement
    standard=_plain(candidate.get('standard_reference'))
    if not standard:
        match=re.search(r'\b(?:conform to|in accordance with)\s+([A-Z]+[- ]?\d[\w.-]*)',requirement,re.I)
        if match:standard=match.group(1)
    return {
        'test_or_inspection_name': name,
        'activity': activity,
        'type': _label(candidate.get('qa_type'), ''),
        'specification_section': ', '.join(_plain(value) for value in (candidate.get('csi_sections') or []) if _plain(value)),
        'timing': _plain(candidate.get('timing')),
        'frequency': _plain(candidate.get('frequency')),
        'acceptance_criteria': acceptance,
        'standard': standard,
        'performer': _inspection_performer(candidate),
        'witness': _plain(candidate.get('witness_as_stated')),
        'requirement_status': _label(candidate.get('requirement_status'), ''),
        'verification': _label(report.get('status'), 'Not checked'),
        'human_review': _label(record.get('review', {}).get('status'), 'Pending human review'),
        'evidence': _record_evidence(data, record, evidence_index),
    }


def _claim(data: dict, record: dict, index: int, claim: dict, evidence_index: dict) -> dict:
    record_id = record.get('meta', {}).get('record_id')
    report = data.get('verifications', {}).get(record_id, {})
    field = next((item for item in report.get('fields', [])
                  if item.get('path') == f'/claims/{index}/normalized_value'), None)
    citations = field.get('citations', []) if field else []
    if citations:
        evidence = _citation_item(citations[0], f'Statement {index + 1}', field.get('status', 'NOT_CHECKED'))
    else:
        source = evidence_index.get(claim.get('evidence_id'), {})
        evidence = _fallback_evidence(source, f'Statement {index + 1}') if source else {
            'source': '', 'location': '', 'revision_date': _plain(claim.get('internal_revision_date')),
            'source_type': 'Source text', 'text': '', 'verification': 'Not checked',
        }
    return {
        'value': _plain(claim.get('normalized_value')),
        'source': evidence.get('source', ''),
        'location': evidence.get('location', ''),
        'revision_date': evidence.get('revision_date') or _plain(claim.get('internal_revision_date')),
        'source_type': evidence.get('source_type', 'Source text'),
        'evidence_text': evidence.get('text', ''),
        'verification': evidence.get('verification', 'Not checked'),
    }


def _resolution_note(candidate: dict) -> str:
    status = candidate.get('resolution_status')
    if status == 'LATEST_APPLIED':
        return 'The value from the source with the latest internal revision date is currently selected. Confirm during human review.'
    if status == 'UNRESOLVED':
        return 'No automatic selection was made. Compare the source statements during human review.'
    return _label(status, 'Pending human review')


def _conflict_item(data: dict, record: dict, evidence_index: dict) -> dict:
    candidate = record['candidate']
    statements = [_claim(data, record, index, claim, evidence_index)
                  for index, claim in enumerate(candidate.get('claims') or [])]
    selected_id = candidate.get('selected_evidence_id')
    selected = next((_plain(claim.get('normalized_value')) for claim in (candidate.get('claims') or [])
                     if selected_id and claim.get('evidence_id') == selected_id), 'None')
    return {
        'subject': _plain(candidate.get('subject')),
        'conflict_type': _label(candidate.get('conflict_type'), ''),
        'statements': statements,
        'current_selection': selected,
        'current_resolution': _label(candidate.get('resolution_status'), 'Pending human review'),
        'resolution_explanation': _resolution_note(candidate),
        'human_review': _label(record.get('review', {}).get('status'), 'Pending human review'),
    }


def _missing_item(data: dict, record: dict, evidence_index: dict) -> dict:
    candidate = record['candidate']
    subject = _plain(candidate.get('subject'))
    if re.fullmatch(r'(?:DOC|EV|REC|MI)-[A-Za-z0-9-]+', subject or ''):
        source = next((evidence_index[value].get('file_name') for value in _evidence_ids(record)
                       if value in evidence_index and evidence_index[value].get('file_name')), '')
        subject = _plain(source) or 'Unmapped or missing information'
    missing_text = _plain(candidate.get('missing_field_or_document'))
    action_text = _plain(candidate.get('recommended_action'))
    diagnostic_pages = sorted({int(value) for value in _PDF_PAGE_DIAGNOSTIC.findall(missing_text + ' ' + action_text)})
    if len(diagnostic_pages) >= 10:
        missing_text = (
            f'{len(diagnostic_pages)} PDF pages were flagged for visual review '
            f'(page span {diagnostic_pages[0]}-{diagnostic_pages[-1]}). '
            'Repeated per-page processing notices are condensed here. Low-text pages may include local OCR and require comparison with the page image.'
        )
        action_text = (
            'Review the flagged pages in the application and compare local OCR text with each page image '
            'before accepting extracted requirements.'
        )
    return {
        'subject_or_scope': subject,
        'what_is_missing': missing_text,
        'impact': _plain(candidate.get('blocking_reason')),
        'recommended_action': action_text,
        'human_review': _label(record.get('review', {}).get('status'), 'Pending human review'),
        'evidence': _record_evidence(data, record, evidence_index),
    }


def _quantity_takeoffs(data: dict, evidence_index: dict) -> list[dict]:
    rows = []
    for document in data.get('analysis_support', []):
        for takeoff in document.get('takeoffs', []):
            rows.append({
                'item': _plain(takeoff.get('label')),
                'quantity': takeoff.get('value'),
                'unit': _plain(takeoff.get('unit')),
                'quantity_type': _label(takeoff.get('kind'), ''),
                'basis': _label(takeoff.get('basis'), _plain(takeoff.get('basis'))),
                'source': _plain(document.get('file_name')),
                'scope': _plain(takeoff.get('scope_note')),
                'human_review': _label(takeoff.get('review_status'), 'Pending human review'),
                'evidence': _evidence_from_ids(takeoff.get('evidence_ids', []), evidence_index),
            })
    return rows


def _readable_untranslated(data: dict) -> dict:
    evidence_index = {item.get('evidence_id'): item for item in data.get('evidence', [])}
    records = data.get('records', [])
    materials=[item for record in records if record.get('kind')=='MATERIAL'
               if (item:=_material_item(data,record,evidence_index)) is not None]
    inspections=[item for record in records if record.get('kind')=='INSPECTION'
                 if (item:=_inspection_item(data,record,evidence_index)) is not None]
    readable = {
        'export_format': 'CIRP reviewer-readable English',
        'export_version': '0.2.6-readable-en-2',
        'summary': _summary(data,len(materials),len(inspections)),
        'materials_and_equipment': materials,
        'inspections_and_tests': inspections,
        'quantity_takeoffs': _quantity_takeoffs(data, evidence_index),
    }
    return readable


def _readable(data: dict) -> dict:
    return _translate_display(_readable_untranslated(data))


def collect(db: Database, run: dict, reviewed_only: bool = False, verifier=None):
    records = []
    record_rows = db.all('SELECT id,envelope,review_version FROM records WHERE run_id=? ORDER BY kind,id', (run['id'],))
    for row in record_rows:
        envelope = json.loads(row['envelope'])
        if reviewed_only and envelope['review']['status'] not in ('ACCEPTED', 'EDITED'):
            continue
        records.append(envelope)
    evidence = []
    for row in db.all('''SELECT e.payload,d.name AS file_name FROM evidence e
                         JOIN documents d ON d.id=e.document_id WHERE e.run_id=?''', (run['id'],)):
        item = json.loads(row['payload']); item['file_name'] = row['file_name']; evidence.append(item)
    report_rows = db.all('''SELECT vr.record_id,vr.payload FROM verification_reports vr
                            JOIN records r ON r.id=vr.record_id WHERE r.run_id=?''', (run['id'],))
    reports_by_id = {row['record_id']: json.loads(row['payload']) for row in report_rows}
    reports = {record['meta']['record_id']: reports_by_id.get(record['meta']['record_id'],
               {'record_id': record['meta']['record_id'], 'status': 'NOT_CHECKED', 'fields': []}) for record in records}
    analysis_support = []
    for row in db.all('''SELECT r.document_id,r.summary,d.name FROM document_results r
                         JOIN documents d ON d.id=r.document_id
                         WHERE r.run_id=? ORDER BY d.name''', (run['id'],)):
        summary = json.loads(row['summary'])
        analysis_support.append({'file_name': row['name'], 'takeoffs': summary.get('takeoffs', [])})
    if reviewed_only:
        analysis_support = []
    project = db.one('SELECT name FROM projects WHERE id=?', (run['project_id'],))
    notice = ('Only accepted or edited records are included. Unapproved quantity aids are excluded.'
              if reviewed_only else
              'Engineering review candidates. Unreviewed or partial results are not procurement or construction instructions.')
    return {'verifications': reports, 'notice': notice,
            'generated_at': datetime.now(timezone.utc).isoformat(), 'run': run,
            'project_name': project['name'] if project else None,
            'cost': db.cost(run['project_id']), 'reviewed_only': reviewed_only,
            'records': records, 'evidence': evidence, 'analysis_support': analysis_support}


def as_json(data: dict) -> bytes:
    return json.dumps(_readable(data), ensure_ascii=False, indent=2, allow_nan=False).encode('utf-8')


def _evidence_cells(items: list[dict], *, max_entries: int = 3, max_excerpt_chars: int = 320) -> tuple[str, str]:
    sources, texts = [], []
    visible_items = []
    seen = set()
    for item in items:
        identity = (_plain(item.get('source')), _plain(item.get('location')), _plain(item.get('text')))
        if identity in seen:
            continue
        seen.add(identity)
        visible_items.append(item)
    for index, item in enumerate(visible_items[:max_entries], 1):
        source_type = _plain(item.get('source_type'))
        source = ' | '.join(value for value in (
            _plain(item.get('source')), _plain(item.get('location')),
            ('Revision date ' + _plain(item.get('revision_date'))) if item.get('revision_date') else '',
            source_type if source_type == 'Visual model observation' else '',
        ) if value)
        evidence_text = _plain(item.get('text'))
        if len(evidence_text) > max_excerpt_chars:
            evidence_text = evidence_text[:max_excerpt_chars].rstrip(' ,;:.') + '...'
        sources.append(f'{index}. {source}')
        texts.append(f'{index}. {evidence_text}')
    remaining = len(visible_items) - len(sources)
    if remaining > 0:
        notice = f'{remaining} additional evidence entr' + ('y' if remaining == 1 else 'ies')
        notice += (' is' if remaining == 1 else ' are') + ' available in the JSON export and application.'
        texts.append(notice)
    return '\n'.join(sources), '\n'.join(texts)


def _limit_cell(value):
    if not isinstance(value, str) or len(value) <= 32600:
        return value
    return value[:32520].rstrip() + ' … Additional text is available in the application.'


def as_xlsx(data: dict) -> bytes:
    # The local application uses the broadly installable openpyxl package.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    readable = _readable(data)
    wb = Workbook(); wb.remove(wb.active)
    navy = '1F4E78'; blue = '4472C4'; stripe = 'F5F8FB'; gray = '666666'
    thin = Side(style='thin', color='D9E2F3')

    def add_sheet(title, subtitle, headers, rows, widths):
        ws = wb.create_sheet(title)
        last = get_column_letter(len(headers))
        ws.merge_cells(f'A1:{last}1'); ws['A1'] = title
        ws['A1'].font = Font(bold=True, color='FFFFFF', size=16)
        ws['A1'].fill = PatternFill('solid', fgColor=navy)
        ws['A1'].alignment = Alignment(vertical='center')
        ws.row_dimensions[1].height = 28
        ws.merge_cells(f'A2:{last}2'); ws['A2'] = subtitle
        ws['A2'].font = Font(italic=True, color=gray, size=10)
        ws['A2'].fill = PatternFill('solid', fgColor='EEF3F8')
        ws['A2'].alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[2].height = 32
        ws.append([]); ws.append(headers)
        for cell in ws[4]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor=blue)
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[4].height = 32
        for row_index, row in enumerate(rows, 5):
            ws.append([_limit_cell(value) for value in row])
            fill = PatternFill('solid', fgColor=stripe) if row_index % 2 == 0 else None
            for cell in ws[row_index]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                cell.border = Border(bottom=thin)
                if fill:
                    cell.fill = fill
            estimated_lines = 1
            for value, width in zip(row, widths):
                if isinstance(value, str) and value:
                    segments = value.splitlines() or [value]
                    width_chars = max(8, int(width))
                    estimated_lines = max(
                        estimated_lines,
                        sum(max(1, (len(segment) + width_chars - 1) // width_chars)
                            for segment in segments),
                    )
            ws.row_dimensions[row_index].height = min(390, max(30, estimated_lines * 13.5))
        ws.freeze_panes = 'A5'
        if rows:
            ws.auto_filter.ref = f'A4:{last}{4 + len(rows)}'
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        ws.sheet_view.showGridLines = False
        return ws

    summary = readable['summary']
    summary_rows = [
        ['Project', summary['project']], ['Run status', summary['run_status']],
        ['Generated (UTC)', summary['generated_utc']], ['Review scope', summary['review_scope']],
        ['Source files', '\n'.join(summary['source_files'])],
        ['Material and equipment items', summary['material_and_equipment_items']],
        ['Test and inspection items', summary['inspection_and_test_items']],
        ['Model calls', summary['model_calls']], ['Model cost (CNY)', summary['model_cost_cny']],
        ['Unresolved model calls', summary['unresolved_model_calls']],
        ['Review notice', summary['notice']], ['Quantity note', summary['quantity_note']],
        ['Evidence display note', summary['evidence_display_note']],
    ]
    add_sheet('Summary', 'Saved results only. Exporting does not call a model or change the project.',
              ['Item', 'Value'], summary_rows, [34, 110])

    material_rows = []
    for item in readable['materials_and_equipment']:
        sources, texts = _evidence_cells(item['evidence'])
        material_rows.append([
            item['material_or_equipment_name'], item['design_properties'], item['specification_section'],
            item['quantity'], item['unit'], item['quantity_basis'], item['material_type'],
            item['manufacturer_product_model_options'], item['location'], item['condition'], item['requirement_status'],
            item['verification'], item['human_review'], sources, texts,
        ])
    add_sheet('Materials & Equipment',
              'Each row names a tangible item. Dimensions, ratings, material, and other stated values remain in Design Properties. A blank quantity is unknown, not zero.',
              ['Material or Equipment Name', 'Design Properties', 'Specification Section', 'Quantity', 'Unit',
               'Quantity Basis', 'Type', 'Manufacturer / Product / Model Options', 'Location', 'Condition',
               'Requirement Status', 'Verification', 'Human Review',
               'Evidence Source', 'Evidence Text'], material_rows,
              [34, 42, 18, 12, 10, 18, 14, 34, 20, 26, 20, 20, 20, 52, 80])

    inspection_rows = []
    for item in readable['inspections_and_tests']:
        sources, texts = _evidence_cells(item['evidence'])
        inspection_rows.append([
            item['test_or_inspection_name'], item['activity'], item['type'], item['specification_section'],
            item['performer'], item['witness'], item['timing'], item['frequency'],
            item['acceptance_criteria'], item['standard'],
            item['requirement_status'], item['verification'], item['human_review'], sources, texts,
        ])
    add_sheet('Inspections & Tests',
              'Only executable test and inspection requirements are listed. Specification section, stated performer, and supporting excerpts stay on the same row.',
              ['Test or Inspection Name', 'Activity / Requirement', 'Type', 'Specification Section',
               'Performer', 'Witness', 'Timing', 'Frequency', 'Acceptance Criteria', 'Standard',
               'Requirement Status', 'Verification', 'Human Review',
               'Evidence Source', 'Evidence Text'], inspection_rows,
              [36, 38, 20, 18, 22, 20, 20, 22, 36, 24, 20, 20, 20, 52, 80])

    if readable['quantity_takeoffs']:
        takeoff_rows = []
        for item in readable['quantity_takeoffs']:
            sources, texts = _evidence_cells(item['evidence'])
            takeoff_rows.append([item['item'], item['quantity'], item['unit'], item['quantity_type'], item['basis'],
                                 item['source'], item['scope'], item['human_review'], sources, texts])
        add_sheet('Quantity Takeoffs',
                  'Only saved CAD quantity candidates are shown. PDF paper-space geometry is intentionally omitted.',
                  ['Item', 'Quantity', 'Unit', 'Quantity Type', 'Basis', 'Source File', 'Scope', 'Human Review',
                   'Evidence Source', 'Evidence Text'], takeoff_rows, [30, 12, 10, 18, 22, 34, 44, 20, 52, 80])

    out = io.BytesIO(); wb.save(out); return out.getvalue()
