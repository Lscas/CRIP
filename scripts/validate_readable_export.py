"""Validate reviewer-readable CIRP JSON/XLSX exports without reading the database."""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


REQUIRED_SHEETS = {'Summary', 'Materials & Equipment', 'Inspections & Tests'}
BANNED_SHEETS = {
    'Evidence', 'Field Verification', 'Citations', 'PDF Geometry Audit',
    'Conflicts', 'Missing Information',
}
INTERNAL_ID = re.compile(r'\b(?:REC|EV|DOC|CALL|RUN)-[0-9a-fA-F]{16,64}\b')
CJK = re.compile(r'[\u3400-\u9fff]')
EXTRACTION_CODE = re.compile(r'\(cid:\d+\)|\ufffd', re.I)
PRIVATE_USE = re.compile(r'[\ue000-\uf8ff\U000f0000-\U000ffffd\U00100000-\U0010fffd]')
MACHINE_CODE = re.compile(r'\b[a-z][a-z0-9]*(?:_[a-z0-9]+)+\b', re.I)
RAW_NULL = re.compile(r'\bnull\b', re.I)
NON_ENGLISH_PUNCTUATION = re.compile(r'[，。；：、！？（）【】《》]')
GENERIC_VISUAL_NOTE = 'Visual model observation recorded for this page. Review the cited source page before use.'
GENERIC_TRANSLATION_DISCLAIMER = re.compile(
    r'\b(?:this|the) input contains (?:a )?visual description of a drawing page\b', re.I
)
LIST_KEYS = ('materials_and_equipment', 'inspections_and_tests', 'quantity_takeoffs')
REMOVED_KEYS = {'conflicts', 'missing_information'}
MATERIAL_DESCRIPTOR_NAME = re.compile(
    r'^(?:equipment|materials?|products?|items?|insulation thickness(?: per pipe size)?|'
    r'water service pipe material and size|\d+(?:\.\d+)?\s*(?:gauge|awg)|'
    r'(?:minimum|maximum)?\s*(?:thickness|size|diameter|pressure|temperature|voltage|flow rate)|'
    r'(?:suspended )?rod diameter specification table|ul system no\.?|ansi\s+[a-z0-9.-]+)$',
    re.I,
)
NON_QA_ITEM = re.compile(
    r'\b(?:shop drawings?|product data|schedule of values|table of contents|schematic|wiring diagrams?|'
    r'coordination drawings?|operation and maintenance|o\s*&\s*m\b|manuals?|submittal register)\b',
    re.I,
)


def walk(value, path=()):
    if isinstance(value, dict):
        for key, item in value.items():
            yield from walk(item, path + (key,))
    elif isinstance(value, list):
        for index, item in enumerate(value):
            yield from walk(item, path + (index,))
    else:
        yield path, value


def validate(json_path: Path, xlsx_path: Path) -> tuple[dict, list[str]]:
    data = json.loads(json_path.read_text(encoding='utf-8'))
    errors = []
    if data.get('export_version') != '0.2.6-readable-en-2':
        errors.append('Unexpected JSON export version.')
    if any(key not in data for key in LIST_KEYS):
        errors.append('JSON output groups are incomplete.')
    if REMOVED_KEYS & set(data):
        errors.append('JSON still contains removed conflict or missing-information output groups.')

    internal_keys = []
    internal_values = []
    whitespace_values = []
    non_evidence_cjk = []
    visual_evidence_cjk = []
    source_evidence_cjk = []
    code_artifacts = []
    private_use_glyphs = 0
    machine_code_values = []
    raw_null_values = []
    punctuation_artifacts = []
    generic_visual_notes = []
    generic_translation_disclaimers = []
    for path, value in walk(data):
        key = path[-1] if path else None
        if isinstance(key, str) and (key.endswith('_id') or key.endswith('_ids') or
                                     key in {'hash', 'sha256', 'candidate', 'meta', 'path'}):
            internal_keys.append(path)
        if not isinstance(value, str):
            continue
        if value and not value.strip():
            whitespace_values.append(path)
        if INTERNAL_ID.search(value):
            internal_values.append(path)
        private_use_glyphs += len(PRIVATE_USE.findall(value))
        if key not in {'text', 'evidence_text'} and CJK.search(value):
            non_evidence_cjk.append((path, value[:160]))
        if MACHINE_CODE.search(value):
            machine_code_values.append((path, value[:160]))
        if RAW_NULL.search(value):
            raw_null_values.append((path, value[:160]))
        if NON_ENGLISH_PUNCTUATION.search(value):
            punctuation_artifacts.append((path, value[:160]))
        if value == GENERIC_VISUAL_NOTE:
            generic_visual_notes.append(path)
        if GENERIC_TRANSLATION_DISCLAIMER.search(value):
            generic_translation_disclaimers.append(path)
        if key in {'text', 'evidence_text'} and CJK.search(value):
            parent = data
            try:
                for part in path[:-1]:
                    parent = parent[part]
            except (KeyError, IndexError, TypeError):
                parent = {}
            if isinstance(parent, dict) and parent.get('source_type') == 'Visual model observation':
                visual_evidence_cjk.append((path, value[:160]))
            else:
                source_evidence_cjk.append((path, value[:160]))
        if ('```' in value or EXTRACTION_CODE.search(value) or
                (key not in {'text', 'evidence_text'} and value.startswith('{') and value.endswith('}'))):
            code_artifacts.append(path)
    if internal_keys:
        errors.append(f'JSON contains {len(internal_keys)} internal ID or code keys.')
    if internal_values:
        errors.append(f'JSON contains {len(internal_values)} embedded internal ID values.')
    if whitespace_values:
        errors.append(f'JSON contains {len(whitespace_values)} whitespace-only strings.')
    if non_evidence_cjk:
        errors.append(f'JSON contains {len(non_evidence_cjk)} non-evidence Chinese strings.')
    if visual_evidence_cjk:
        errors.append(f'JSON contains {len(visual_evidence_cjk)} Chinese visual-model observation strings.')
    if source_evidence_cjk:
        errors.append(f'JSON contains {len(source_evidence_cjk)} Chinese source-evidence strings.')
    if code_artifacts:
        errors.append(f'JSON contains {len(code_artifacts)} display-code artifacts.')
    if private_use_glyphs:
        errors.append(f'JSON contains {private_use_glyphs} private-use glyphs.')
    if machine_code_values:
        errors.append(f'JSON contains {len(machine_code_values)} machine-style underscore values.')
    if raw_null_values:
        errors.append(f'JSON contains {len(raw_null_values)} visible raw null values.')
    if punctuation_artifacts:
        errors.append(f'JSON contains {len(punctuation_artifacts)} non-English punctuation artifacts.')
    if generic_visual_notes:
        errors.append(f'JSON contains {len(generic_visual_notes)} generic visual evidence placeholders.')
    if generic_translation_disclaimers:
        errors.append(f'JSON contains {len(generic_translation_disclaimers)} unrelated generic translation disclaimers.')

    invalid_material_names = [
        (index, str(item.get('material_or_equipment_name', '')))
        for index, item in enumerate(data.get('materials_and_equipment', []))
        if MATERIAL_DESCRIPTOR_NAME.fullmatch(str(item.get('material_or_equipment_name', '')).strip())
    ]
    if invalid_material_names:
        errors.append(f'JSON contains {len(invalid_material_names)} attribute or generic material names.')
    invalid_qa_items = [
        (index, str(item.get('test_or_inspection_name', '')))
        for index, item in enumerate(data.get('inspections_and_tests', []))
        if NON_QA_ITEM.search(' '.join(str(item.get(key, '')) for key in
                                   ('test_or_inspection_name', 'activity')))
    ]
    if invalid_qa_items:
        errors.append(f'JSON contains {len(invalid_qa_items)} administrative or drawing submittals as QA items.')

    wb = load_workbook(xlsx_path, read_only=False, data_only=False)
    sheet_names = set(wb.sheetnames)
    if not REQUIRED_SHEETS.issubset(sheet_names):
        errors.append('Workbook is missing a required reviewer sheet.')
    if BANNED_SHEETS & sheet_names:
        errors.append('Workbook contains a removed technical or geometry sheet.')
    if sheet_names - REQUIRED_SHEETS - {'Quantity Takeoffs'}:
        errors.append('Workbook contains an unexpected sheet.')
    material_headers = {cell.value for cell in wb['Materials & Equipment'][4]}
    inspection_headers = {cell.value for cell in wb['Inspections & Tests'][4]}
    if not {'Material or Equipment Name', 'Design Properties', 'Specification Section',
            'Quantity', 'Unit'}.issubset(material_headers):
        errors.append('Materials sheet is missing a required item, property, specification, or quantity column.')
    if not {'Test or Inspection Name', 'Specification Section', 'Performer'}.issubset(inspection_headers):
        errors.append('Inspections sheet is missing a required name, specification, or performer column.')

    formulas = hyperlinks = whitespace_cells = internal_id_cells = extraction_code_cells = 0
    machine_code_cells = raw_null_cells = punctuation_cells = private_use_cells = 0
    non_english_labels = 0
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                formulas += int(cell.data_type == 'f')
                hyperlinks += int(bool(cell.hyperlink))
                if isinstance(value, str):
                    whitespace_cells += int(bool(value) and not value.strip())
                    internal_id_cells += int(bool(INTERNAL_ID.search(value)))
                    extraction_code_cells += int(bool(EXTRACTION_CODE.search(value)))
                    private_use_cells += len(PRIVATE_USE.findall(value))
                    machine_code_cells += int(bool(MACHINE_CODE.search(value)))
                    raw_null_cells += int(bool(RAW_NULL.search(value)))
                    punctuation_cells += int(bool(NON_ENGLISH_PUNCTUATION.search(value)))
                    if cell.row <= 4:
                        non_english_labels += int(bool(CJK.search(value)))
    if formulas:
        errors.append(f'Workbook contains {formulas} formulas.')
    if hyperlinks:
        errors.append(f'Workbook contains {hyperlinks} hyperlinks.')
    if whitespace_cells:
        errors.append(f'Workbook contains {whitespace_cells} whitespace-only cells.')
    if internal_id_cells:
        errors.append(f'Workbook contains {internal_id_cells} embedded internal ID cells.')
    if extraction_code_cells:
        errors.append(f'Workbook contains {extraction_code_cells} display-code cells.')
    if private_use_cells:
        errors.append(f'Workbook contains {private_use_cells} private-use glyphs.')
    if non_english_labels:
        errors.append(f'Workbook contains {non_english_labels} non-English labels.')
    if machine_code_cells:
        errors.append(f'Workbook contains {machine_code_cells} machine-style underscore cells.')
    if raw_null_cells:
        errors.append(f'Workbook contains {raw_null_cells} visible raw null cells.')
    if punctuation_cells:
        errors.append(f'Workbook contains {punctuation_cells} non-English punctuation cells.')

    duplicated_quantity_names = []
    word_numbers = {value: key for key, value in {
        'one': '1', 'two': '2', 'three': '3', 'four': '4', 'five': '5', 'six': '6',
        'seven': '7', 'eight': '8', 'nine': '9', 'ten': '10', 'eleven': '11',
        'twelve': '12', 'thirteen': '13', 'fourteen': '14', 'fifteen': '15',
        'sixteen': '16', 'seventeen': '17', 'eighteen': '18', 'nineteen': '19', 'twenty': '20',
    }.items()}
    for index, item in enumerate(data.get('materials_and_equipment', [])):
        quantity = item.get('quantity')
        if quantity is None:
            continue
        number = str(int(quantity)) if isinstance(quantity, float) and quantity.is_integer() else str(quantity)
        word = word_numbers.get(number)
        word_tokens = [rf'{word}\s*\(\s*{re.escape(number)}\s*\)', word] if word else []
        if number == '1':
            word_tokens.append('single')
        token = '(?:' + '|'.join(word_tokens) + ')' if word_tokens else r'(?!)'
        pattern = re.compile(
            rf'^\s*(?:at\s+least\s+|minimum(?:\s+of)?\s+)?(?:{token}|\(\s*{re.escape(number)}\s*\))(?=\s)',
            re.I,
        )
        name = str(item.get('material_or_equipment_name', ''))
        if pattern.search(name) or re.search(rf'(?<![A-Za-z0-9])\(\s*{re.escape(number)}\s*\)(?![A-Za-z0-9])', name):
            duplicated_quantity_names.append(index)
    if duplicated_quantity_names:
        errors.append(f'{len(duplicated_quantity_names)} material or equipment names duplicate their separate quantity.')

    counts = {key: len(data[key]) for key in LIST_KEYS}
    report = {
        'status': 'PASS' if not errors else 'FAIL',
        'json_bytes': json_path.stat().st_size,
        'xlsx_bytes': xlsx_path.stat().st_size,
        'counts': counts,
        'sheets': wb.sheetnames,
        'sheet_rows': {sheet.title: sheet.max_row for sheet in wb.worksheets},
        'invalid_material_names': len(invalid_material_names),
        'invalid_test_or_inspection_items': len(invalid_qa_items),
        'formula_cells': formulas,
        'hyperlinks': hyperlinks,
        'whitespace_only_strings': len(whitespace_values) + whitespace_cells,
        'internal_id_or_code_keys': len(internal_keys),
        'embedded_internal_id_values': len(internal_values) + internal_id_cells,
        'generic_visual_evidence_placeholders': len(generic_visual_notes),
        'generic_translation_disclaimers': len(generic_translation_disclaimers),
        'machine_style_underscore_values': len(machine_code_values) + machine_code_cells,
        'visible_raw_null_values': len(raw_null_values) + raw_null_cells,
        'non_english_punctuation_artifacts': len(punctuation_artifacts) + punctuation_cells,
        'quantity_duplicated_in_names': len(duplicated_quantity_names),
        'non_english_output_labels': non_english_labels,
        'non_evidence_chinese_strings': len(non_evidence_cjk),
        'unique_non_evidence_chinese_strings': len({value for _, value in non_evidence_cjk}),
        'non_evidence_chinese_examples': [
            {'path': '/'.join(map(str, path)), 'text': value}
            for path, value in non_evidence_cjk[:12]
        ],
        'chinese_visual_model_observations': len(visual_evidence_cjk),
        'chinese_source_document_quotes': len(source_evidence_cjk),
        'chinese_source_document_examples': [
            {'path': '/'.join(map(str, path)), 'text': value}
            for path, value in source_evidence_cjk[:8]
        ],
        'display_code_artifacts': len(code_artifacts) + extraction_code_cells,
        'private_use_glyphs': private_use_glyphs + private_use_cells,
    }
    return report, errors


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument('json_path', type=Path)
    parser.add_argument('xlsx_path', type=Path)
    args = parser.parse_args()
    report, errors = validate(args.json_path, args.xlsx_path)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    for error in errors:
        print('[FAIL]', error)
    return int(bool(errors))


if __name__ == '__main__':
    raise SystemExit(main())
